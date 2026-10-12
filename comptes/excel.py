"""
Construction de classeurs Excel (.xlsx) pour les exports : un titre, quelques informations, un tableau mis en forme.

Utilisé par les exports de listes (élèves d'une classe, relevé de notes, absences) et par la comptabilité. La
confidentialité n'est PAS décidée ici : chaque vue choisit ce qu'elle exporte avant d'appeler ces fonctions.
"""
import datetime
import re
from decimal import Decimal

from django.conf import settings
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

TYPE_MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_TITRE = Font(bold=True, size=14, color="1F4E78")
_GRAS = Font(bold=True)
_FOND_ENTETE = PatternFill("solid", fgColor="DDEBF7")
_FILET = Side(style="thin", color="999999")
_BORDURE = Border(left=_FILET, right=_FILET, top=_FILET, bottom=_FILET)
LARGEUR_MIN, LARGEUR_MAX = 8, 45


def nom_fichier_sur(*morceaux) -> str:
    """Nom de fichier sans accents ni caractère spécial (« releve_notes_6A_2026-2027 »)."""
    texte = "_".join(str(m) for m in morceaux if m not in (None, ""))
    return re.sub(r"[^A-Za-z0-9_-]+", "_", texte).strip("_") or "export"


def nom_feuille_sur(texte) -> str:
    """Nom d'onglet valide : 31 caractères au plus, sans [ ] : * ? / \\."""
    return re.sub(r"[\[\]:*?/\\]", " ", str(texte)).strip()[:31] or "Feuille"


def nouveau_classeur() -> Workbook:
    classeur = Workbook()
    classeur.properties.creator = getattr(settings, "NOM_PLATEFORME", "Nexora Éducation")
    classeur.remove(classeur.active)  # on ajoute nos propres feuilles
    return classeur


def ecrire_valeur(feuille, ligne, colonne, valeur):
    """Écrit une cellule : nombres et dates tels quels, texte TOUJOURS en texte (jamais interprété comme formule)."""
    if valeur is None or valeur == "":
        return None
    cellule = feuille.cell(row=ligne, column=colonne, value=valeur)
    if isinstance(valeur, str):
        cellule.data_type = "s"
    elif isinstance(valeur, bool):
        cellule.value = "oui" if valeur else "non"
        cellule.data_type = "s"
    elif isinstance(valeur, Decimal):
        cellule.value = float(valeur)
    elif isinstance(valeur, datetime.datetime):
        cellule.number_format = "DD/MM/YYYY HH:MM"
    elif isinstance(valeur, datetime.date):
        cellule.number_format = "DD/MM/YYYY"
    return cellule


def ajouter_feuille_tableau(classeur, nom, titre, infos, entetes, lignes, *, total=None, formats=None):
    """
    Ajoute une feuille : titre, informations (liste de (libellé, valeur)), ligne d'en-têtes, lignes de données et
    éventuellement une ligne de total. `formats` : {numéro de colonne (1 = première): format Excel} pour les nombres.
    """
    feuille = classeur.create_sheet(nom_feuille_sur(nom))
    ligne_courante = 1
    ecrire_valeur(feuille, ligne_courante, 1, titre).font = _TITRE
    ligne_courante += 1
    for libelle, valeur in infos:
        ecrire_valeur(feuille, ligne_courante, 1, libelle)
        if feuille.cell(row=ligne_courante, column=1).value is not None:
            feuille.cell(row=ligne_courante, column=1).font = _GRAS
        ecrire_valeur(feuille, ligne_courante, 2, valeur)
        ligne_courante += 1
    ligne_courante += 1

    ligne_entete = ligne_courante
    for colonne, entete in enumerate(entetes, start=1):
        cellule = ecrire_valeur(feuille, ligne_entete, colonne, entete)
        if cellule is not None:
            cellule.font, cellule.fill, cellule.border = _GRAS, _FOND_ENTETE, _BORDURE
            cellule.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    feuille.freeze_panes = feuille.cell(row=ligne_entete + 1, column=1)

    formats = formats or {}
    largeurs = {colonne: len(str(entete)) for colonne, entete in enumerate(entetes, start=1)}
    derniere = ligne_entete
    for valeurs in lignes:
        derniere += 1
        for colonne, valeur in enumerate(valeurs, start=1):
            cellule = ecrire_valeur(feuille, derniere, colonne, valeur)
            if cellule is None:
                continue
            cellule.border = _BORDURE
            if colonne in formats and not isinstance(valeur, str):
                cellule.number_format = formats[colonne]
            largeurs[colonne] = max(largeurs.get(colonne, 0), len(str(valeur)))
    if total is not None:
        derniere += 1
        for colonne, valeur in enumerate(total, start=1):
            cellule = ecrire_valeur(feuille, derniere, colonne, valeur)
            if cellule is not None:
                cellule.font, cellule.border = _GRAS, _BORDURE
                if colonne in formats and not isinstance(valeur, str):
                    cellule.number_format = formats[colonne]
    if lignes:
        feuille.auto_filter.ref = f"A{ligne_entete}:{feuille.cell(row=ligne_entete, column=len(entetes)).column_letter}{ligne_entete + len(lignes)}"

    for colonne, longueur in largeurs.items():
        feuille.column_dimensions[feuille.cell(row=1, column=colonne).column_letter].width = min(
            max(longueur + 2, LARGEUR_MIN), LARGEUR_MAX,
        )
    feuille.page_setup.orientation = "landscape"
    feuille.page_setup.fitToWidth = 1
    feuille.page_setup.fitToHeight = 0
    feuille.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    return feuille


def reponse_xlsx(classeur, nom_fichier) -> HttpResponse:
    """Réponse de téléchargement d'un classeur (jamais mis en cache : voir le middleware des pages connectées)."""
    reponse = HttpResponse(content_type=TYPE_MIME_XLSX)
    reponse["Content-Disposition"] = f'attachment; filename="{nom_fichier}.xlsx"'
    classeur.save(reponse)
    return reponse
