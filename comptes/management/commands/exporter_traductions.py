"""
Écrit le catalogue de traduction dans un tableau Excel (Français | Arabe | Note), pour le faire relire et corriger par une personne
qui parle la langue. Le fichier corrigé se remet dans le catalogue avec `manage.py importer_traductions`.

    python manage.py exporter_traductions --sortie traductions_arabe.xlsx
"""
import csv

from django.core.management.base import BaseCommand
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from comptes.traduction import DOSSIER_CATALOGUES


class Command(BaseCommand):
    help = "Exporte le catalogue de traduction en Excel pour relecture."

    def add_arguments(self, parser):
        parser.add_argument("--langue", default="ar")
        parser.add_argument("--sortie", default="traductions_arabe.xlsx", help="Fichier Excel à écrire.")

    def handle(self, *args, **options):
        chemin = DOSSIER_CATALOGUES / f"{options['langue']}.csv"
        entrees = {}
        with open(chemin, encoding="utf-8-sig", newline="") as fichier:
            for ligne in csv.DictReader(fichier):
                entrees[ligne["francais"]] = (ligne["arabe"], ligne.get("note", ""))
        classeur = Workbook()
        feuille = classeur.active
        feuille.title = "Traductions"
        feuille.append(["Français", "Arabe (à corriger si besoin)", "Note"])
        for cellule in feuille[1]:
            cellule.font = Font(bold=True, color="FFFFFF")
            cellule.fill = PatternFill("solid", fgColor="12203D")
        for francais, (arabe, note) in entrees.items():
            feuille.append([francais, arabe, note])
        feuille.column_dimensions["A"].width = 70
        feuille.column_dimensions["B"].width = 70
        feuille.column_dimensions["C"].width = 55
        for rangee in feuille.iter_rows(min_row=2):
            rangee[0].alignment = Alignment(wrap_text=True, vertical="top")
            rangee[1].alignment = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)
            rangee[2].alignment = Alignment(wrap_text=True, vertical="top")
        feuille.freeze_panes = "A2"
        classeur.save(options["sortie"])
        self.stdout.write(f"{len(entrees)} traductions écrites dans {options['sortie']}.")
        self.stdout.write("Les mots entre accolades, par exemple {0} dans « {0} élèves inscrits », sont remplacés par une valeur : ils doivent rester dans la traduction.")
