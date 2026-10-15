"""
Remet dans le catalogue de traduction (comptes/traductions/ar.csv) le tableau Excel relu et corrigé
(voir `manage.py exporter_traductions`). Refuse tout le fichier si une ligne est incohérente : texte vide, valeur entre
accolades ({0}) perdue ou ajoutée, texte sans lettre arabe.

    python manage.py importer_traductions traductions_arabe.xlsx --simulation
    python manage.py importer_traductions traductions_arabe.xlsx
"""
import csv
import re

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from comptes.traduction import DOSSIER_CATALOGUES, TROU, normaliser, vider_le_cache

LETTRE_ARABE = re.compile(r"[؀-ۿ]")


class Command(BaseCommand):
    help = "Importe un tableau Excel relu dans le catalogue de traduction."

    def add_arguments(self, parser):
        parser.add_argument("fichier", help="Fichier Excel (colonnes : Français, Arabe, Note).")
        parser.add_argument("--langue", default="ar")
        parser.add_argument("--simulation", action="store_true", help="Contrôle et compte les changements sans écrire.")

    def handle(self, *args, **options):
        feuille = load_workbook(options["fichier"], read_only=True).active
        lignes = [rangee for rangee in feuille.iter_rows(min_row=2, values_only=True) if rangee and rangee[0]]
        chemin = DOSSIER_CATALOGUES / f"{options['langue']}.csv"
        with open(chemin, encoding="utf-8-sig", newline="") as fichier:
            anciens = {normaliser(l["francais"]): (l["arabe"], l.get("note", "")) for l in csv.DictReader(fichier)}
        problemes, nouveaux = [], {}
        for numero, rangee in enumerate(lignes, start=2):
            francais = str(rangee[0]).strip()
            arabe = str(rangee[1] or "").strip()
            note = str(rangee[2] or "").strip() if len(rangee) > 2 else ""
            if not arabe:
                problemes.append(f"ligne {numero} : traduction vide pour « {francais} »")
                continue
            if sorted(TROU.findall(francais)) != sorted(TROU.findall(arabe)):
                problemes.append(f"ligne {numero} : les valeurs entre accolades ne correspondent pas ({francais} / {arabe})")
            if not LETTRE_ARABE.search(arabe) and normaliser(arabe) != normaliser(francais) and re.search(r"[^\W\d_]", arabe):
                problemes.append(f"ligne {numero} : aucune lettre arabe dans « {arabe} »")
            nouveaux[normaliser(francais)] = (francais, arabe, note)
        if problemes:
            raise CommandError("Fichier refusé :\n  " + "\n  ".join(problemes[:30]) + (f"\n  ... et {len(problemes) - 30} autres" if len(problemes) > 30 else ""))
        modifies = sum(1 for cle, (_, arabe, _) in nouveaux.items() if cle in anciens and anciens[cle][0] != arabe)
        ajoutes = sum(1 for cle in nouveaux if cle not in anciens)
        retires = sum(1 for cle in anciens if cle not in nouveaux)
        self.stdout.write(f"{len(nouveaux)} traductions lues : {modifies} corrigée(s), {ajoutes} ajoutée(s), {retires} retirée(s).")
        if options["simulation"]:
            self.stdout.write("Simulation : rien n'a été écrit.")
            return
        with open(chemin, "w", encoding="utf-8-sig", newline="") as fichier:
            ecrivain = csv.writer(fichier)
            ecrivain.writerow(["francais", "arabe", "note"])
            for francais, arabe, note in nouveaux.values():
                ecrivain.writerow([francais, arabe, note])
        vider_le_cache()
        self.stdout.write(f"Catalogue écrit : {chemin}")
