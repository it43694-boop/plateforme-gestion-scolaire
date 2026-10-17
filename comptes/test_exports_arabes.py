"""Les exports (Excel, CSV) d'une page arabe sont écrits en arabe : textes, noms, chiffres des textes ; feuilles de droite à gauche."""
import csv
import io

from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook

from comptes.arabisation import invalider
from comptes.langues import NOM_COOKIE_LANGUE
from comptes.models import Utilisateur
from comptes.roles import Role
from scolarite.models import Inscription
from statistiques.test_education import EducationBase, creer_compte


class ExportsEnArabeTests(EducationBase):
    def setUp(self):
        self.fondateur = creer_compte(Role.FONDATEUR, "fondateur@exports-ar.test", self.ecole)
        Utilisateur.objects.filter(pk=self.e1.pk).update(prenom="Salem", nom="Keita", prenom_ar="سالم", nom_ar="كيتا")
        Inscription.objects.get_or_create(eleve=self.e1, classe=self.c6a, defaults={"statut": Inscription.Statut.EN_COURS})
        invalider(self.ecole.pk)

    def client_en(self, langue):
        client = Client()
        client.force_login(self.fondateur)
        client.cookies[NOM_COOKIE_LANGUE] = langue
        return client

    def liste_classe(self, langue):
        reponse = self.client_en(langue).get(reverse("scolarite:exporter_liste_classe_xlsx", args=[self.c6a.pk]))
        self.assertEqual(reponse.status_code, 200)
        return load_workbook(io.BytesIO(reponse.content))

    @staticmethod
    def textes(classeur):
        return [c.value for feuille in classeur.worksheets for rangee in feuille.iter_rows() for c in rangee if isinstance(c.value, str)]

    def test_la_liste_de_classe_est_en_arabe_et_de_droite_a_gauche(self):
        classeur = self.liste_classe("ar")
        textes = self.textes(classeur)
        self.assertIn("سالم", textes)  # le prénom et le nom sont deux colonnes
        self.assertIn("كيتا", textes)
        self.assertFalse(any("Salem" in texte for texte in textes), textes)
        self.assertTrue(all(feuille.sheet_view.rightToLeft for feuille in classeur.worksheets))

    def test_la_liste_de_classe_en_francais_ne_change_pas(self):
        classeur = self.liste_classe("fr")
        textes = " ".join(self.textes(classeur))
        self.assertIn("Salem", textes)
        self.assertFalse(any(feuille.sheet_view.rightToLeft for feuille in classeur.worksheets))

    def test_les_nombres_restent_des_nombres(self):
        classeur = self.load_fiche("ar")
        nombres = [c.value for feuille in classeur.worksheets for rangee in feuille.iter_rows() for c in rangee if isinstance(c.value, (int, float))]
        self.assertTrue(nombres)

    def load_fiche(self, langue):
        reponse = self.client_en(langue).get(reverse("statistiques:exporter_fiche_annuelle_xlsx"))
        self.assertEqual(reponse.status_code, 200)
        return load_workbook(io.BytesIO(reponse.content))

    def test_le_csv_de_la_fiche_est_en_arabe(self):
        reponse = self.client_en("ar").get(reverse("statistiques:exporter_fiche_annuelle_csv"))
        self.assertEqual(reponse.status_code, 200)
        lignes = list(csv.reader(io.StringIO(reponse.content.decode("utf-8-sig")), delimiter=";"))
        titre = lignes[0][0]
        self.assertTrue(any("؀" <= lettre <= "ۿ" for lettre in titre), titre)

    def test_le_csv_des_effectifs_a_un_en_tete_arabe(self):
        reponse = self.client_en("ar").get(reverse("statistiques:exporter_effectifs_csv"))
        premiere = reponse.content.decode("utf-8-sig").split("\n")[0]
        self.assertIn("القسم", premiere)
