"""
Statistiques de l'éducation : fiche annuelle, résultats aux examens, infrastructures, préscolaire,
situation de l'école, champs du personnel et des élèves.

Les chiffres attendus sont calculés à la main à partir du jeu de données de EducationBase (voir les
commentaires) : un test qui échoue signale un vrai écart de calcul, pas un détail d'affichage.
"""
import datetime
import io
import zipfile

from django.test import TestCase
from django.urls import reverse

from comptes.models import JournalAudit, Utilisateur, generer_matricule
from comptes.roles import Role, StatutCompte
from etablissement.models import Etablissement
from openpyxl import load_workbook
from permissions_matrix.models import PermissionMatrix
from scolarite.models import (
    Affectation, AnneeScolaire, Classe, Cycle, FicheInfrastructure, Inscription, ResultatExamen,
)
from statistiques.fiche import age_au, cellule_csv, construire_fiche, lignes_csv

D = datetime.date


def creer_compte(role, email, ecole=None, **extra):
    compte = Utilisateur(
        email=email, prenom="Test", nom=email.split("@")[0], role=role, etablissement=ecole,
        statut=StatutCompte.ACTIF, is_active=True, email_verifie=True, **extra,
    )
    if role == Role.ELEVE:
        compte.matricule = generer_matricule()
    compte.set_unusable_password()
    compte.save()
    return compte


class EducationBase(TestCase):
    """
    École A, année 2026-2027 (début 1er octobre 2026), année précédente 2025-2026.

    Classes : 6A et 6B (1er cycle, niveau 6), 9A (2ème cycle), TSE (lycée, série TSE), CAP-E (CAP industriel).
    Élèves (sexe, naissance -> âge au 1er octobre 2026) :
      6A : e1 F 2014-05-01 -> 12 (redouble, déjà inscrite), e2 M 2013-03-10 -> 13 (déjà inscrit), e3 F sans date (nouvelle)
      6B : e4 M 2014-07-01 -> 12 (nouveau)
      9A : e5 F 2011-09-01 -> 15, e6 M 2011-12-01 -> 14
      TSE: e7 F 2008-01-01 -> 18, e8 M 2008-11-01 -> 17
      CAP: e9 M 2009-02-02 -> 17
    Total 9 : 5 garçons, 4 filles, 5 classes. Enseignants : t1, t2, t3 (3). e2 handicap moteur, e6 visuel, e9 moteur ;
    e3 déplacée interne, e4 réfugié.
    """

    @classmethod
    def setUpTestData(cls):
        cls.ecole = Etablissement.objects.create(
            nom="École statistiques A", code_etablissement="ML-001", region="Koulikoro", cercle="Kati", commune="Kati",
        )
        cls.ecole_b = Etablissement.objects.create(nom="École statistiques B")
        for ecole in (cls.ecole, cls.ecole_b):
            PermissionMatrix.seed_pour(ecole)
        cls.precedente = AnneeScolaire.objects.create(
            etablissement=cls.ecole, libelle="2025-2026", date_debut=D(2025, 10, 1), date_fin=D(2026, 7, 31),
        )
        cls.annee = AnneeScolaire.objects.create(
            etablissement=cls.ecole, libelle="2026-2027", date_debut=D(2026, 10, 1), date_fin=D(2027, 7, 31),
            est_active=True,
        )
        cls.annee_b = AnneeScolaire.objects.create(
            etablissement=cls.ecole_b, libelle="2026-2027", date_debut=D(2026, 10, 1), date_fin=D(2027, 7, 31),
            est_active=True,
        )
        cls.classe_5a = Classe.objects.create(
            nom="5A", cycle=Cycle.PREMIER_CYCLE, niveau="annee_5", annee_scolaire=cls.precedente,
        )
        cls.c6a = Classe.objects.create(nom="6A", cycle=Cycle.PREMIER_CYCLE, niveau="annee_6", annee_scolaire=cls.annee)
        cls.c6b = Classe.objects.create(nom="6B", cycle=Cycle.PREMIER_CYCLE, niveau="annee_6", annee_scolaire=cls.annee)
        cls.c9a = Classe.objects.create(nom="9A", cycle=Cycle.DEUXIEME_CYCLE, niveau="annee_9", annee_scolaire=cls.annee)
        cls.ctse = Classe.objects.create(
            nom="TSE", cycle=Cycle.LYCEE, niveau="annee_12", serie="tse", annee_scolaire=cls.annee,
        )
        cls.ccap = Classe.objects.create(
            nom="CAP-E", cycle=Cycle.PROFESSIONNEL, filiere_professionnelle="cap",
            domaine_professionnel="industriel", annee_scolaire=cls.annee,
        )
        cls.classe_b = Classe.objects.create(
            nom="6A", cycle=Cycle.PREMIER_CYCLE, niveau="annee_6", annee_scolaire=cls.annee_b,
        )

        def eleve(numero, sexe, naissance, **extra):
            return creer_compte(Role.ELEVE, f"eleve{numero}@stat.test", cls.ecole, sexe=sexe, date_naissance=naissance, **extra)

        cls.e1 = eleve(1, "F", D(2014, 5, 1))
        cls.e2 = eleve(2, "M", D(2013, 3, 10), handicap="moteur")
        cls.e3 = eleve(3, "F", None, statut_deplacement="deplace_interne")
        cls.e4 = eleve(4, "M", D(2014, 7, 1), statut_deplacement="refugie")
        cls.e5 = eleve(5, "F", D(2011, 9, 1))
        cls.e6 = eleve(6, "M", D(2011, 12, 1), handicap="visuel")
        cls.e7 = eleve(7, "F", D(2008, 1, 1))
        cls.e8 = eleve(8, "M", D(2008, 11, 1))
        cls.e9 = eleve(9, "M", D(2009, 2, 2), handicap="moteur")

        # année précédente : e1 a redoublé, e2 a été admis
        Inscription.objects.create(eleve=cls.e1, classe=cls.classe_5a, statut=Inscription.Statut.REDOUBLE)
        Inscription.objects.create(eleve=cls.e2, classe=cls.classe_5a, statut=Inscription.Statut.ADMIS)
        cls.i = {}
        for nom, eleve_, classe in [
            ("e1", cls.e1, cls.c6a), ("e2", cls.e2, cls.c6a), ("e3", cls.e3, cls.c6a), ("e4", cls.e4, cls.c6b),
            ("e5", cls.e5, cls.c9a), ("e6", cls.e6, cls.c9a), ("e7", cls.e7, cls.ctse), ("e8", cls.e8, cls.ctse),
            ("e9", cls.e9, cls.ccap),
        ]:
            cls.i[nom] = Inscription.objects.create(eleve=eleve_, classe=classe)

        cls.t1 = creer_compte(Role.ENSEIGNANT, "t1@stat.test", cls.ecole, sexe="F", statut_personnel="fonctionnaire",
                              formation_pedagogique="forme")
        cls.t2 = creer_compte(Role.ENSEIGNANT, "t2@stat.test", cls.ecole, sexe="M",
                              statut_personnel="fonctionnaire_collectivite", formation_pedagogique="non_forme")
        cls.t3 = creer_compte(Role.ENSEIGNANT, "t3@stat.test", cls.ecole)
        cls.t_ancien = creer_compte(Role.ENSEIGNANT, "tancien@stat.test", cls.ecole, sexe="M")
        Affectation.objects.create(enseignant=cls.t1, classe=cls.c6a, matiere="Maths")
        Affectation.objects.create(enseignant=cls.t2, classe=cls.c6b, matiere="Maths")
        Affectation.objects.create(enseignant=cls.t2, classe=cls.c9a, matiere="Maths")
        Affectation.objects.create(enseignant=cls.t3, classe=cls.ctse, matiere="Maths")
        Affectation.objects.create(enseignant=cls.t_ancien, classe=cls.classe_5a, matiere="Maths")

        Resultat = ResultatExamen.Resultat
        for nom, examen, resultat in [
            ("e1", "entree_7e", Resultat.ADMIS), ("e2", "entree_7e", Resultat.AJOURNE), ("e3", "entree_7e", Resultat.ABSENT),
            ("e5", "def", Resultat.ADMIS), ("e6", "def", Resultat.AJOURNE),
            ("e7", "bac", Resultat.ADMIS), ("e8", "bac", Resultat.ABSENT), ("e9", "cap", Resultat.ADMIS),
        ]:
            ResultatExamen.objects.create(inscription=cls.i[nom], examen=examen, resultat=resultat)

        FicheInfrastructure.objects.create(
            etablissement=cls.ecole, annee_scolaire=cls.annee, salles_classe=10, salles_bon_etat=8, salles_provisoires=2,
            places_assises=300, points_eau=2, latrines=6, latrines_fonctionnelles=4, manuels_lecture=18, manuels_calcul=9,
            cantine=True, cantine_beneficiaires=3,
        )

        cls.fondateur = creer_compte(Role.FONDATEUR, "fondateur@stat.test", cls.ecole)
        cls.dir_1er = creer_compte(Role.DIRECTEUR_1ER_CYCLE, "dir1@stat.test", cls.ecole)
        cls.comptable = creer_compte(Role.COMPTABLE, "comptable@stat.test", cls.ecole)
        cls.secretaire = creer_compte(Role.SECRETAIRE, "secretaire@stat.test", cls.ecole)
        cls.censeur = creer_compte(Role.CENSEUR, "censeur@stat.test", cls.ecole)  # voit les dossiers du lycée, pas l'état civil
        cls.parent = creer_compte(Role.PARENT, "parent@stat.test", cls.ecole, telephone="70999001")
        cls.e1.parents_lies.add(cls.parent)
        cls.fondateur_b = creer_compte(Role.FONDATEUR, "fondateur-b@stat.test", cls.ecole_b)

    def fiche_pour(self, utilisateur):
        return construire_fiche(utilisateur, self.annee)

    def ligne(self, fiche, niveau):
        return next(ligne for ligne in fiche["effectifs"] if ligne["niveau"] == niveau)


class FicheAnnuelleCalculsTests(EducationBase):
    def test_effectifs_par_niveau_et_sexe(self):
        fiche = self.fiche_pour(self.fondateur)
        niveau_6 = self.ligne(fiche, "annee_6")
        self.assertEqual((niveau_6["classes"], niveau_6["garcons"], niveau_6["filles"], niveau_6["total"]), (2, 2, 2, 4))
        self.assertEqual(niveau_6["parite"], 100.0)
        self.assertEqual(niveau_6["eleves_par_classe"], 2.0)
        totaux = fiche["totaux"]
        self.assertEqual((totaux["classes"], totaux["garcons"], totaux["filles"], totaux["total"]), (5, 5, 4, 9))
        self.assertEqual(totaux["parite"], 80.0)
        self.assertEqual(totaux["eleves_par_classe"], 1.8)

    def test_ordre_des_niveaux_par_cycle(self):
        fiche = self.fiche_pour(self.fondateur)
        self.assertEqual(
            [ligne["niveau"] for ligne in fiche["effectifs"]], ["annee_6", "annee_9", "annee_12", ""],
        )  # la classe CAP n'a pas de niveau : « non précisé », en dernier

    def test_redoublants_et_nouveaux_inscrits(self):
        fiche = self.fiche_pour(self.fondateur)
        self.assertEqual(fiche["annee_precedente"], self.precedente)
        niveau_6 = self.ligne(fiche, "annee_6")
        self.assertEqual(niveau_6["redoublants"], 1)  # e1
        self.assertEqual(niveau_6["nouveaux"], 2)     # e3 et e4 (e1 et e2 étaient déjà inscrits)
        self.assertEqual(fiche["totaux"]["redoublants"], 1)
        self.assertEqual(fiche["totaux"]["nouveaux"], 7)

    def test_sans_annee_precedente_redoublants_et_nouveaux_non_calcules(self):
        fiche = construire_fiche(self.fondateur_b, self.annee_b)
        self.assertIsNone(fiche["annee_precedente"])
        self.assertIsNone(fiche["totaux"]["redoublants"])
        self.assertIsNone(fiche["totaux"]["nouveaux"])
        self.assertIsNone(fiche["effectifs"][0]["redoublants"])  # classe vide : effectif 0, pas d'erreur

    def test_ages_au_debut_de_l_annee(self):
        fiche = self.fiche_pour(self.fondateur)
        par_age = {ligne["age"]: (ligne["garcons"], ligne["filles"]) for ligne in fiche["ages"]}
        self.assertEqual(par_age, {12: (1, 1), 13: (1, 0), 14: (1, 0), 15: (0, 1), 17: (2, 0), 18: (0, 1)})
        self.assertEqual(fiche["ages_inconnus"]["total"], 1)  # e3
        self.assertEqual(age_au(D(2011, 12, 1), D(2026, 10, 1)), 14)  # pas encore 15 ans
        self.assertEqual(age_au(D(2011, 10, 1), D(2026, 10, 1)), 15)  # anniversaire le jour même

    def test_enseignants_par_situation_formation_et_ratio(self):
        enseignants = self.fiche_pour(self.fondateur)["enseignants"]
        self.assertEqual(enseignants["total"], 3)  # t1, t2 (deux classes, compté une fois), t3 ; pas t_ancien
        self.assertEqual(enseignants["qualifies"], 1)
        self.assertEqual(enseignants["eleves_par_enseignant"], 3.0)
        statuts = {ligne["statut"]: ligne["total"] for ligne in enseignants["par_statut"]}
        self.assertEqual(statuts, {
            "Fonctionnaire de l'État": 1, "Fonctionnaire des collectivités territoriales": 1,
            "Situation non renseignée": 1,
        })

    def test_resultats_aux_examens(self):
        examens = {(ligne["examen"], ligne["sous_groupe"]): ligne for ligne in self.fiche_pour(self.fondateur)["examens"]}
        entree = examens[("entree_7e", "")]
        self.assertEqual((entree["presentes"], entree["admis"], entree["ajournes"], entree["absents"], entree["taux"]),
                         (2, 1, 1, 1, 50.0))
        definitif = examens[("def", "")]
        self.assertEqual((definitif["taux"], definitif["taux_garcons"], definitif["taux_filles"]), (50.0, 0.0, 100.0))
        bac = examens[("bac", "Sciences Exactes (TSE)")]
        self.assertEqual((bac["presentes"], bac["absents"], bac["taux"]), (1, 1, 100.0))  # l'absent ne compte pas
        self.assertEqual(examens[("cap", "Industriel")]["taux"], 100.0)

    def test_infrastructures_et_ratios_pour_la_direction_generale(self):
        infra = self.fiche_pour(self.fondateur)["infrastructures"]
        self.assertEqual(infra["fiche"].salles_classe, 10)
        self.assertEqual(infra["ratios"]["eleves_par_salle"], 0.9)
        self.assertEqual(infra["ratios"]["manuels_lecture_par_eleve"], 2.0)
        self.assertEqual(infra["ratios"]["part_cantine"], 33.3)

    def test_aucune_fiche_d_infrastructures(self):
        self.assertIsNone(construire_fiche(self.fondateur_b, self.annee_b)["infrastructures"])


class FicheAnnuelleCloisonnementTests(EducationBase):
    def test_directeur_de_cycle_ne_voit_que_son_cycle(self):
        fiche = self.fiche_pour(self.dir_1er)
        self.assertEqual(fiche["totaux"]["total"], 4)  # e1..e4 uniquement
        self.assertEqual([ligne["niveau"] for ligne in fiche["effectifs"]], ["annee_6"])
        self.assertIn("1er cycle", fiche["portee"])
        self.assertEqual(fiche["enseignants"]["total"], 2)  # t1 et t2 (t2 enseigne aussi en 9A, mais seul 6B compte)
        self.assertEqual([ligne["examen"] for ligne in fiche["examens"]], ["entree_7e"])
        self.assertIsNone(fiche["infrastructures"]["ratios"])  # ratios sur toute l'école : masqués

    def test_chiffres_sensibles_reserves_a_l_administration(self):
        sensibles = self.fiche_pour(self.fondateur)["sensibles"]
        self.assertEqual(sensibles["total_handicap"], 3)
        self.assertEqual({ligne["libelle"]: ligne["total"] for ligne in sensibles["handicap"]},
                         {"Handicap moteur": 2, "Handicap visuel": 1})
        self.assertEqual(sensibles["total_deplacement"], 2)
        self.assertIsNone(self.fiche_pour(self.comptable)["sensibles"])
        cycle = self.fiche_pour(self.dir_1er)["sensibles"]
        self.assertEqual((cycle["total_handicap"], cycle["total_deplacement"]), (1, 2))

    def test_autre_ecole_ne_voit_rien_de_l_ecole_a(self):
        fiche = construire_fiche(self.fondateur_b, self.annee_b)
        self.assertEqual(fiche["totaux"]["total"], 0)
        self.assertEqual(fiche["enseignants"]["total"], 0)
        self.assertEqual(fiche["examens"], [])

    def test_page_de_l_autre_ecole_introuvable(self):
        self.client.force_login(self.fondateur)
        self.assertEqual(self.client.get(reverse("statistiques:fiche_annuelle"), {"annee": self.annee_b.id}).status_code, 404)
        self.assertEqual(self.client.get(reverse("statistiques:infrastructures"), {"annee": self.annee_b.id}).status_code, 404)
        for nom in ("exporter_fiche_annuelle_csv", "exporter_fiche_annuelle_xlsx"):
            with self.subTest(export=nom):
                self.assertEqual(
                    self.client.get(reverse(f"statistiques:{nom}"), {"annee": self.annee_b.id}).status_code, 404,
                )

    def test_identifiant_d_annee_invalide_est_une_page_introuvable(self):
        self.client.force_login(self.fondateur)
        for nom in ("fiche_annuelle", "infrastructures", "resultats_examens", "vue_ensemble"):
            with self.subTest(page=nom):
                self.assertEqual(self.client.get(reverse(f"statistiques:{nom}"), {"annee": "abc"}).status_code, 404)

    def test_pages_fermees_aux_roles_sans_le_module_statistiques(self):
        for compte in (self.secretaire, self.parent, self.t1):
            self.client.force_login(compte)
            for nom in (
                "fiche_annuelle", "exporter_fiche_annuelle_csv", "exporter_fiche_annuelle_xlsx", "resultats_examens",
                "infrastructures",
            ):
                with self.subTest(role=compte.role, page=nom):
                    self.assertEqual(self.client.get(reverse(f"statistiques:{nom}")).status_code, 403)

    def test_visiteur_anonyme_redirige(self):
        for nom in (
            "fiche_annuelle", "exporter_fiche_annuelle_csv", "exporter_fiche_annuelle_xlsx", "resultats_examens",
            "infrastructures",
        ):
            self.assertEqual(self.client.get(reverse(f"statistiques:{nom}")).status_code, 302)


class FicheAnnuellePagesEtCsvTests(EducationBase):
    def test_page_affiche_les_rubriques(self):
        self.client.force_login(self.fondateur)
        page = self.client.get(reverse("statistiques:fiche_annuelle"), {"annee": self.annee.id})
        self.assertEqual(page.status_code, 200)
        for texte in ("Effectifs par niveau", "Âge des élèves", "Enseignants", "Résultats aux examens", "Infrastructures",
                      "École statistiques A", "ML-001", "Fiche de travail", "Sciences Exactes (TSE)",
                      "Exporter en Excel", "Exporter en CSV", reverse("statistiques:exporter_fiche_annuelle_xlsx")):
            self.assertContains(page, texte)

    def test_export_csv(self):
        self.client.force_login(self.fondateur)
        reponse = self.client.get(reverse("statistiques:exporter_fiche_annuelle_csv"), {"annee": self.annee.id})
        self.assertEqual(reponse.status_code, 200)
        self.assertIn("text/csv", reponse["Content-Type"])
        self.assertIn("fiche_statistique_2026-2027.csv", reponse["Content-Disposition"])
        contenu = reponse.content.decode("utf-8")
        self.assertTrue(contenu.startswith("﻿"))
        self.assertIn("EFFECTIFS PAR NIVEAU", contenu)
        self.assertIn("1er cycle;6ème année;2;2;2;0;4;100.0;2.0;1;2", contenu)
        self.assertIn("TOTAL;;5;5;4;0;9;80.0;1.8;1;7", contenu)
        self.assertTrue(JournalAudit.objects.filter(action="export_fiche_statistique_annuelle", cible="2026-2027").exists())

    def test_export_csv_du_directeur_de_cycle_ne_contient_pas_les_autres_cycles(self):
        self.client.force_login(self.dir_1er)
        contenu = self.client.get(
            reverse("statistiques:exporter_fiche_annuelle_csv"), {"annee": self.annee.id},
        ).content.decode("utf-8")
        self.assertNotIn("Baccalauréat", contenu)
        self.assertNotIn("2ème cycle", contenu)
        self.assertNotIn("INFRASTRUCTURES\nFiche", contenu)

    def test_une_formule_dans_un_nom_n_est_pas_interpretee_par_le_tableur(self):
        self.assertEqual(cellule_csv("=CMD()"), "'=CMD()")
        self.assertEqual(cellule_csv("+1"), "'+1")
        self.assertEqual(cellule_csv("@x"), "'@x")
        self.assertEqual(cellule_csv("École normale"), "École normale")
        self.assertEqual(cellule_csv(None), "")
        self.assertEqual(cellule_csv(12.5), 12.5)
        self.ecole.nom = "=HYPERLINK(\"http://x\")"
        self.ecole.save()
        self.client.force_login(self.fondateur)
        contenu = self.client.get(
            reverse("statistiques:exporter_fiche_annuelle_csv"), {"annee": self.annee.id},
        ).content.decode("utf-8")
        self.assertNotIn(";=HYPERLINK", contenu)
        self.assertIn("'=HYPERLINK", contenu)

    def test_lignes_csv_sans_infrastructures(self):
        lignes = lignes_csv(construire_fiche(self.fondateur_b, self.annee_b))
        self.assertIn(["Fiche d'infrastructures non renseignée pour cette année"], lignes)


def lignes_du_classeur(contenu):
    """Lignes d'un .xlsx telles que l'utilisateur les voit dans Excel (cellules vides de fin retirées)."""
    feuille = load_workbook(io.BytesIO(contenu)).active
    lignes = []
    for rangee in feuille.iter_rows():
        valeurs = [cellule.value for cellule in rangee]
        while valeurs and valeurs[-1] is None:
            valeurs.pop()
        lignes.append(valeurs)
    return lignes


def lignes_attendues(fiche):
    """Les lignes du CSV, avec « vide » = cellule vide."""
    lignes = []
    for ligne in lignes_csv(fiche):
        valeurs = [None if valeur in ("", None) else valeur for valeur in ligne]
        while valeurs and valeurs[-1] is None:
            valeurs.pop()
        lignes.append(valeurs)
    return lignes


class FicheAnnuelleExcelTests(EducationBase):
    def exporter(self, compte, annee=None):
        self.client.force_login(compte)
        return self.client.get(reverse("statistiques:exporter_fiche_annuelle_xlsx"), {"annee": (annee or self.annee).id})

    def test_export_excel(self):
        reponse = self.exporter(self.fondateur)
        self.assertEqual(reponse.status_code, 200)
        self.assertEqual(
            reponse["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn('filename="fiche_statistique_2026-2027.xlsx"', reponse["Content-Disposition"])
        self.assertIn("no-store", reponse["Cache-Control"])
        classeur = load_workbook(io.BytesIO(reponse.content))
        self.assertEqual(classeur.sheetnames, ["Fiche statistique"])
        self.assertTrue(JournalAudit.objects.filter(
            action="export_fiche_statistique_annuelle", cible="2026-2027", details={"format": "xlsx"},
        ).exists())

    def test_meme_contenu_que_le_csv_et_chiffres_en_nombres(self):
        reponse = self.exporter(self.fondateur)
        self.assertEqual(
            lignes_du_classeur(reponse.content), lignes_attendues(construire_fiche(self.fondateur, self.annee)),
        )
        lignes = lignes_du_classeur(reponse.content)
        # mêmes chiffres que le test du CSV (« TOTAL;;5;5;4;0;9;80.0;1.8;1;7 »), mais de vrais nombres : sommables
        self.assertIn(["TOTAL", None, 5, 5, 4, 0, 9, 80, 1.8, 1, 7], lignes)
        total = next(ligne for ligne in lignes if ligne[:1] == ["TOTAL"])
        self.assertTrue(all(isinstance(valeur, (int, float)) for valeur in total[2:]))

    def test_les_rubriques_sont_mises_en_forme(self):
        feuille = load_workbook(io.BytesIO(self.exporter(self.fondateur).content)).active
        titre = next(c for c in feuille["A"] if c.value == "EFFECTIFS PAR NIVEAU")
        self.assertTrue(titre.font.bold)
        entete = feuille.cell(row=titre.row + 1, column=1)
        self.assertEqual(entete.value, "Cycle")
        self.assertTrue(entete.font.bold)
        self.assertEqual(feuille.page_setup.orientation, "landscape")

    def test_un_nom_commencant_par_egal_reste_du_texte(self):
        self.ecole.nom = '=HYPERLINK("http://x")'
        self.ecole.save()
        contenu = self.exporter(self.fondateur).content
        self.assertIn(['Établissement', '=HYPERLINK("http://x")'], lignes_du_classeur(contenu))
        with zipfile.ZipFile(io.BytesIO(contenu)) as archive:
            for nom in archive.namelist():
                if nom.startswith("xl/worksheets/"):
                    self.assertNotIn("<f>", archive.read(nom).decode("utf-8"), "aucune formule ne doit être écrite")

    def test_le_directeur_de_cycle_n_obtient_que_son_cycle(self):
        lignes = lignes_du_classeur(self.exporter(self.dir_1er).content)
        texte = " ".join(str(valeur) for ligne in lignes for valeur in ligne)
        self.assertIn("1er cycle", texte)
        self.assertNotIn("Baccalauréat", texte)
        self.assertNotIn("2ème cycle", texte)
        self.assertEqual(lignes, lignes_attendues(construire_fiche(self.dir_1er, self.annee)))

    def test_l_export_ne_contient_aucun_nom_ni_matricule(self):
        contenu = self.exporter(self.fondateur).content
        texte = " ".join(str(valeur) for ligne in lignes_du_classeur(contenu) for valeur in ligne)
        for eleve in Utilisateur.objects.filter(role=Role.ELEVE, etablissement=self.ecole):
            self.assertNotIn(eleve.matricule, texte)
            self.assertNotIn(eleve.email, texte)
        with zipfile.ZipFile(io.BytesIO(contenu)) as archive:  # ni dans les propriétés du fichier
            proprietes = archive.read("docProps/core.xml").decode("utf-8")
        self.assertNotIn(self.fondateur.email, proprietes)
        self.assertNotIn("@", proprietes)


class InfrastructuresTests(EducationBase):
    URL = "statistiques:infrastructures"

    def donnees(self, **modifications):
        valeurs = {
            "salles_classe": 12, "salles_bon_etat": 10, "salles_provisoires": 1, "places_assises": 400,
            "points_eau": 3, "latrines": 8, "latrines_fonctionnelles": 6, "manuels_lecture": 50, "manuels_calcul": 40,
            "cantine_beneficiaires": 0,
        }
        valeurs.update(modifications)
        return valeurs

    def test_la_direction_generale_met_a_jour_la_fiche_existante(self):
        self.client.force_login(self.fondateur)
        reponse = self.client.post(f"{reverse(self.URL)}?annee={self.annee.id}", self.donnees())
        self.assertEqual(reponse.status_code, 302)
        fiche = FicheInfrastructure.objects.get(etablissement=self.ecole, annee_scolaire=self.annee)
        self.assertEqual((fiche.salles_classe, fiche.manuels_lecture, fiche.modifie_par), (12, 50, self.fondateur))
        self.assertFalse(fiche.cantine)  # case décochée dans l'envoi
        self.assertEqual(FicheInfrastructure.objects.filter(etablissement=self.ecole).count(), 1)  # pas de doublon
        self.assertTrue(JournalAudit.objects.filter(action="modification_fiche_infrastructures").exists())

    def test_creation_pour_une_annee_sans_fiche(self):
        self.client.force_login(self.fondateur)
        self.client.post(f"{reverse(self.URL)}?annee={self.precedente.id}", self.donnees())
        self.assertTrue(FicheInfrastructure.objects.filter(etablissement=self.ecole, annee_scolaire=self.precedente).exists())

    def test_valeurs_incoherentes_refusees_et_rien_n_est_enregistre(self):
        self.client.force_login(self.fondateur)
        for modification, message in [
            ({"salles_bon_etat": 99}, "total de salles"),
            ({"latrines_fonctionnelles": 99}, "nombre de latrines"),
            ({"cantine_beneficiaires": 5}, "cantine"),
            ({"salles_classe": -1}, ""),
        ]:
            with self.subTest(modification=modification):
                reponse = self.client.post(f"{reverse(self.URL)}?annee={self.annee.id}", self.donnees(**modification))
                self.assertEqual(reponse.status_code, 200)
                if message:
                    self.assertContains(reponse, message)
        self.assertEqual(FicheInfrastructure.objects.get(etablissement=self.ecole, annee_scolaire=self.annee).salles_classe, 10)

    def test_le_directeur_de_cycle_consulte_sans_modifier(self):
        self.client.force_login(self.dir_1er)
        page = self.client.get(f"{reverse(self.URL)}?annee={self.annee.id}")
        self.assertEqual(page.status_code, 200)
        self.assertNotContains(page, "Enregistrer")
        self.assertNotContains(page, "Indicateurs")  # ratios sur toute l'école : masqués
        self.assertEqual(self.client.post(f"{reverse(self.URL)}?annee={self.annee.id}", self.donnees()).status_code, 403)
        self.assertEqual(FicheInfrastructure.objects.get(etablissement=self.ecole, annee_scolaire=self.annee).salles_classe, 10)

    def test_la_direction_generale_voit_les_indicateurs(self):
        self.client.force_login(self.fondateur)
        page = self.client.get(f"{reverse(self.URL)}?annee={self.annee.id}")
        self.assertContains(page, "Indicateurs")
        self.assertContains(page, "Enregistrer")

    def test_une_autre_ecole_ne_modifie_pas_la_fiche(self):
        self.client.force_login(self.fondateur_b)
        reponse = self.client.post(f"{reverse(self.URL)}?annee={self.annee.id}", self.donnees())
        self.assertEqual(reponse.status_code, 404)
        self.assertEqual(FicheInfrastructure.objects.get(etablissement=self.ecole, annee_scolaire=self.annee).salles_classe, 10)

    def test_validation_du_modele(self):
        from django.core.exceptions import ValidationError
        fiche = FicheInfrastructure(
            etablissement=self.ecole_b, annee_scolaire=self.annee, salles_classe=1,
        )
        with self.assertRaises(ValidationError):  # année d'une autre école
            fiche.full_clean(exclude=["modifie_par"])


class ResultatsExamenTests(EducationBase):
    def url(self, classe):
        return reverse("scolarite:saisir_resultats_examen", args=[classe.id])

    def test_coherence_examen_et_classe(self):
        from django.core.exceptions import ValidationError

        def resultat(nom, examen):
            return ResultatExamen(
                inscription=self.i[nom], examen=examen, resultat="admis", enregistre_par=self.fondateur,
            )

        with self.assertRaisesMessage(ValidationError, "concerne les classes du cycle"):
            resultat("e1", "def").full_clean()  # une 6ème année ne passe pas le DEF
        with self.assertRaisesMessage(ValidationError, "filière de la classe"):
            resultat("e9", "bt").full_clean()  # une classe CAP ne passe pas le BT
        resultat("e4", "entree_7e").full_clean()  # ne doit pas lever d'exception

    def test_un_seul_resultat_par_inscription_et_par_examen(self):
        from django.db import IntegrityError, transaction
        with self.assertRaises(IntegrityError), transaction.atomic():
            ResultatExamen.objects.create(inscription=self.i["e5"], examen="def", resultat="admis")

    def test_saisie_ajoute_modifie_et_efface(self):
        self.client.force_login(self.secretaire)
        reponse = self.client.post(self.url(self.c9a), {
            "examen": "def", f"resultat_{self.i['e5'].id}": "ajourne", f"resultat_{self.i['e6'].id}": "",
        })
        self.assertEqual(reponse.status_code, 302)
        self.assertEqual(ResultatExamen.objects.get(inscription=self.i["e5"], examen="def").resultat, "ajourne")
        self.assertFalse(ResultatExamen.objects.filter(inscription=self.i["e6"], examen="def").exists())
        entree = JournalAudit.objects.get(action="resultats_examen_saisis")
        self.assertEqual(entree.details["modifies"], 1)
        self.assertEqual(entree.details["supprimes"], 1)
        self.assertNotIn(self.e5.nom, f"{entree.cible} {entree.details}")  # aucun nom d'élève dans le journal

    def test_un_champ_absent_de_l_envoi_ne_touche_pas_au_resultat(self):
        self.client.force_login(self.secretaire)
        self.client.post(self.url(self.c9a), {"examen": "def"})
        self.assertEqual(ResultatExamen.objects.filter(inscription__classe=self.c9a).count(), 2)

    def test_valeur_invalide_ignoree(self):
        self.client.force_login(self.secretaire)
        reponse = self.client.post(self.url(self.c9a), {"examen": "def", f"resultat_{self.i['e5'].id}": "mention_tres_bien"}, follow=True)
        self.assertEqual(ResultatExamen.objects.get(inscription=self.i["e5"], examen="def").resultat, "admis")
        self.assertTrue(any("invalide" in str(m) for m in reponse.context["messages"]))

    def test_les_inscriptions_d_une_autre_classe_ne_sont_pas_modifiables(self):
        self.client.force_login(self.secretaire)
        self.client.post(self.url(self.c9a), {"examen": "def", f"resultat_{self.i['e7'].id}": "ajourne"})  # e7 est en TSE
        self.assertEqual(ResultatExamen.objects.get(inscription=self.i["e7"], examen="bac").resultat, "admis")
        self.assertFalse(ResultatExamen.objects.filter(inscription=self.i["e7"], examen="def").exists())

    def test_examen_qui_ne_concerne_pas_la_classe(self):
        self.client.force_login(self.secretaire)
        reponse = self.client.post(self.url(self.c9a), {"examen": "bac", f"resultat_{self.i['e5'].id}": "admis"})
        self.assertEqual(reponse.status_code, 302)
        self.assertFalse(ResultatExamen.objects.filter(inscription=self.i["e5"], examen="bac").exists())

    def test_les_eleves_partis_ne_sont_pas_candidats(self):
        Inscription.objects.filter(pk=self.i["e6"].pk).update(statut=Inscription.Statut.ABANDON)
        self.client.force_login(self.secretaire)
        page = self.client.get(self.url(self.c9a))
        self.assertContains(page, self.e5.matricule)
        self.assertNotContains(page, self.e6.matricule)

    def test_get_affiche_la_liste_et_les_resultats_existants(self):
        self.client.force_login(self.fondateur)
        page = self.client.get(self.url(self.c9a))
        self.assertContains(page, "Diplôme d&#x27;études fondamentales")
        self.assertContains(page, self.e5.matricule)
        self.assertEqual(
            [ligne["resultat"] for ligne in page.context["lignes"]],
            [("admis" if ligne["inscription"].eleve_id == self.e5.id else "ajourne") for ligne in page.context["lignes"]],
        )

    def test_roles_sans_droit_et_autres_ecoles_refuses(self):
        for compte in (self.t1, self.parent, self.comptable, self.e5):
            with self.subTest(role=compte.role):
                self.client.force_login(compte)
                self.assertEqual(self.client.post(self.url(self.c9a), {"examen": "def"}).status_code, 403)
        self.client.force_login(self.fondateur_b)
        self.assertEqual(self.client.get(self.url(self.c9a)).status_code, 404)

    def test_directeur_de_cycle_limite_a_son_cycle(self):
        self.client.force_login(self.dir_1er)
        self.assertEqual(self.client.get(self.url(self.c6a)).status_code, 200)
        self.assertEqual(self.client.get(self.url(self.c9a)).status_code, 404)

    def test_classe_sans_examen_officiel(self):
        prescolaire = Classe.objects.create(nom="PS", cycle=Cycle.PRESCOLAIRE, annee_scolaire=self.annee)
        self.client.force_login(self.fondateur)
        self.assertRedirects(
            self.client.get(self.url(prescolaire)), reverse("scolarite:liste_classes"), fetch_redirect_response=False,
        )

    def test_page_de_synthese_et_acces_a_la_saisie(self):
        self.client.force_login(self.fondateur)
        page = self.client.get(reverse("statistiques:resultats_examens"), {"annee": self.annee.id})
        self.assertContains(page, "Taux de réussite")
        self.assertContains(page, self.url(self.c9a))
        self.client.force_login(self.dir_1er)
        page = self.client.get(reverse("statistiques:resultats_examens"), {"annee": self.annee.id})
        self.assertContains(page, self.url(self.c6a))
        self.assertNotContains(page, self.url(self.c9a))


class DossierEleveEducationTests(EducationBase):
    def dossier(self, eleve):
        return reverse("scolarite:dossier_eleve", args=[eleve.matricule])

    def test_handicap_et_deplacement_visibles_de_l_administration_et_du_parent(self):
        for compte in (self.fondateur, self.parent):
            with self.subTest(role=compte.role):
                self.client.force_login(compte)
                eleve = self.e2 if compte == self.fondateur else self.e1
                if compte == self.parent:  # e1 n'a ni handicap ni déplacement : on lui en donne pour le test
                    Utilisateur.objects.filter(pk=self.e1.pk).update(handicap="auditif")
                page = self.client.get(self.dossier(eleve))
                self.assertContains(page, "Situation de handicap")

    def test_handicap_masque_aux_roles_sans_acces_a_l_etat_civil(self):
        Utilisateur.objects.filter(pk=self.e7.pk).update(handicap="auditif", statut_deplacement="refugie")
        self.client.force_login(self.censeur)
        page = self.client.get(self.dossier(self.e7))
        self.assertEqual(page.status_code, 200)  # le Censeur voit le dossier d'un élève du lycée...
        self.assertNotContains(page, "Situation de handicap")  # ...mais pas ces renseignements sensibles
        self.assertNotContains(page, "Réfugié")
        self.client.force_login(self.fondateur)
        self.assertContains(self.client.get(self.dossier(self.e7)), "Situation de handicap")

    def test_resultats_d_examen_dans_le_dossier(self):
        self.client.force_login(self.fondateur)
        page = self.client.get(self.dossier(self.e7))
        self.assertContains(page, "Résultats aux examens")
        self.assertContains(page, "Baccalauréat")
        self.client.force_login(self.censeur)
        page = self.client.get(self.dossier(self.e7))
        self.assertEqual(page.status_code, 200)
        self.assertNotContains(page, "Résultats aux examens")


class ChampsPersonnelEtElevesTests(EducationBase):
    """Fiche du personnel et renseignements des élèves (formation, prise de service, handicap, déplacement)."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.dev = creer_compte(Role.DEVELOPPEUR, "dev@stat.test", cls.ecole)

    def modifier(self, compte, **donnees):
        self.client.force_login(self.dev)
        base = {"role": compte.role, "statut": "actif"}
        base.update(donnees)
        return self.client.post(reverse("espace_developpeur:modifier_compte", args=[compte.id]), base, follow=True)

    def test_anciens_statuts_toujours_valides_et_statut_collectivite_ajoute(self):
        from comptes.models import StatutPersonnel
        for valeur in ("fonctionnaire", "contractuel_etat", "contractuel_collectivite", "personnel_ecole", "volontaire"):
            self.assertIn(valeur, StatutPersonnel.values)
        self.assertIn("fonctionnaire_collectivite", StatutPersonnel.values)

    def test_fiche_du_personnel_enregistree(self):
        self.modifier(
            self.t3, statut_personnel="contractuel_etat", diplome="Licence", discipline="Maths",
            formation_pedagogique="forme", date_prise_service="2019-10-01", heures_hebdomadaires="24",
        )
        self.t3.refresh_from_db()
        self.assertEqual(
            (self.t3.statut_personnel, self.t3.formation_pedagogique, self.t3.date_prise_service, self.t3.heures_hebdomadaires),
            ("contractuel_etat", "forme", D(2019, 10, 1), 24),
        )

    def test_valeurs_invalides_refusees_sans_rien_modifier(self):
        for donnees in ({"heures_hebdomadaires": "61"}, {"date_prise_service": "pas-une-date"}, {"formation_pedagogique": "expert"}):
            with self.subTest(donnees=donnees):
                reponse = self.modifier(self.t3, **donnees)
                self.assertEqual(reponse.status_code, 200)
                self.t3.refresh_from_db()
                self.assertIsNone(self.t3.heures_hebdomadaires)
                self.assertIsNone(self.t3.date_prise_service)
                self.assertEqual(self.t3.formation_pedagogique, "")

    def test_champs_vides_restent_vides(self):
        self.modifier(self.t3, statut_personnel="", heures_hebdomadaires="", date_prise_service="")
        self.t3.refresh_from_db()
        self.assertEqual((self.t3.heures_hebdomadaires, self.t3.date_prise_service), (None, None))

    def test_changement_de_role_efface_les_nouveaux_champs_et_previent(self):
        Utilisateur.objects.filter(pk=self.t3.pk).update(
            formation_pedagogique="forme", date_prise_service=D(2019, 10, 1), heures_hebdomadaires=20,
        )
        reponse = self.modifier(self.t3, role="parent")
        self.t3.refresh_from_db()
        self.assertEqual((self.t3.formation_pedagogique, self.t3.date_prise_service, self.t3.heures_hebdomadaires), ("", None, None))
        textes = " ".join(str(m) for m in reponse.context["messages"])
        self.assertIn("date de prise de service", textes)
        self.assertIn("heures d'enseignement par semaine", textes)

    def test_renseignements_eleve_via_l_espace_developpeur(self):
        self.modifier(self.e1, handicap="moteur", statut_deplacement="deplace_interne", date_prise_service="2019-10-01")
        self.e1.refresh_from_db()
        self.assertEqual((self.e1.handicap, self.e1.statut_deplacement), ("moteur", "deplace_interne"))
        self.assertIsNone(self.e1.date_prise_service)  # champ du personnel : sans objet pour un élève

    def test_valeur_de_handicap_inconnue_refusee(self):
        reponse = self.modifier(self.e1, handicap="inconnu")
        self.assertEqual(reponse.status_code, 200)
        self.e1.refresh_from_db()
        self.assertEqual(self.e1.handicap, "")

    def test_inscription_enregistre_handicap_et_deplacement(self):
        self.client.force_login(self.secretaire)
        donnees = {
            "prenom": "Awa", "nom": "Zqxdiallo", "sexe": "F", "classe": self.c6a.id,
            "parent_email_1": "parent-insc-stat@stat.test", "handicap": "visuel", "statut_deplacement": "refugie",
        }
        self.assertEqual(self.client.post(reverse("scolarite:inscrire_eleve"), donnees).status_code, 302)
        eleve = Utilisateur.objects.get(nom="Zqxdiallo")
        self.assertEqual((eleve.handicap, eleve.statut_deplacement), ("visuel", "refugie"))

    def test_inscription_possible_sans_ces_renseignements_et_valeur_invalide_refusee(self):
        self.client.force_login(self.secretaire)
        donnees = {"prenom": "Awa", "nom": "Zqxkeita", "sexe": "F", "classe": self.c6a.id, "parent_email_1": "p-k@stat.test"}
        self.assertEqual(self.client.post(reverse("scolarite:inscrire_eleve"), donnees).status_code, 302)
        self.assertEqual(Utilisateur.objects.get(nom="Zqxkeita").handicap, "")
        mauvaises = {**donnees, "nom": "Zqxrefuse", "parent_email_1": "p-r@stat.test", "handicap": "inconnu"}
        self.assertEqual(self.client.post(reverse("scolarite:inscrire_eleve"), mauvaises).status_code, 200)
        self.assertFalse(Utilisateur.objects.filter(nom="Zqxrefuse").exists())

    def test_le_journal_ne_contient_pas_le_handicap_ni_le_deplacement(self):
        self.client.force_login(self.secretaire)
        self.client.post(reverse("scolarite:inscrire_eleve"), {
            "prenom": "Awa", "nom": "Zqxjournal", "sexe": "F", "classe": self.c6a.id,
            "parent_email_1": "p-j@stat.test", "handicap": "intellectuel", "statut_deplacement": "refugie",
        })
        for entree in JournalAudit.objects.all():
            texte = f"{entree.cible} {entree.details}"
            self.assertNotIn("intellectuel", texte)
            self.assertNotIn("refugie", texte)


class PrescolaireTests(EducationBase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.proprietaire = creer_compte(Role.DEVELOPPEUR, "owner@stat.test", None, is_superuser=True, is_staff=True)
        cls.dev = creer_compte(Role.DEVELOPPEUR, "dev2@stat.test", cls.ecole)

    def classe_prescolaire(self, **extra):
        return Classe(
            nom="PS", cycle=Cycle.PRESCOLAIRE, annee_scolaire=AnneeScolaire.objects.get(pk=self.annee.pk), **extra,
        )

    def test_absent_de_toutes_les_formules_existantes(self):
        from etablissement.models import PlanEtablissement
        from scolarite.models import cycles_autorises_pour
        for plan in PlanEtablissement.values:
            if plan == "prescolaire_seul":
                continue
            with self.subTest(plan=plan):
                self.ecole.plan = plan
                self.assertNotIn("prescolaire", cycles_autorises_pour(self.ecole))

    def test_option_cumulable_et_formule_prescolaire_seul(self):
        from scolarite.models import cycles_autorises_pour
        self.ecole.plan, self.ecole.inclut_prescolaire = "lycee_seul", True
        self.assertEqual(cycles_autorises_pour(self.ecole), {"lycee", "prescolaire"})
        self.ecole.plan, self.ecole.inclut_prescolaire = "prescolaire_seul", False
        self.assertEqual(cycles_autorises_pour(self.ecole), {"prescolaire"})

    def test_classe_refusee_sans_l_option_acceptee_avec(self):
        from django.core.exceptions import ValidationError
        with self.assertRaises(ValidationError):
            self.classe_prescolaire().clean()
        Etablissement.objects.filter(pk=self.ecole.pk).update(inclut_prescolaire=True)
        self.classe_prescolaire().clean()  # ne doit pas lever d'exception

    def test_niveaux_du_prescolaire(self):
        from django.core.exceptions import ValidationError
        Etablissement.objects.filter(pk=self.ecole.pk).update(inclut_prescolaire=True)
        self.classe_prescolaire(niveau="presco_gs").clean()
        with self.assertRaises(ValidationError):
            self.classe_prescolaire(niveau="annee_3").clean()
        with self.assertRaises(ValidationError):  # une section de maternelle dans une classe de 1er cycle
            Classe(nom="X", cycle=Cycle.PREMIER_CYCLE, niveau="presco_ps", annee_scolaire=self.annee).clean()

    def test_formulaire_de_creation_de_classe(self):
        self.client.force_login(self.fondateur)
        self.assertNotContains(self.client.get(reverse("scolarite:creer_classe")), '<option value="prescolaire">')
        Etablissement.objects.filter(pk=self.ecole.pk).update(inclut_prescolaire=True)
        self.assertContains(self.client.get(reverse("scolarite:creer_classe")), '<option value="prescolaire">')

    def test_seul_le_proprietaire_active_l_option(self):
        url = reverse("espace_plateforme:modifier_plan", args=[self.ecole.id])
        self.client.force_login(self.dev)
        self.client.post(url, {"plan": "tous_cycles", "inclut_prescolaire": "True"})
        self.ecole.refresh_from_db()
        self.assertFalse(self.ecole.inclut_prescolaire)
        self.client.force_login(self.proprietaire)
        self.client.post(url, {"plan": "tous_cycles", "inclut_prescolaire": "True"})
        self.ecole.refresh_from_db()
        self.assertTrue(self.ecole.inclut_prescolaire)
        entree = JournalAudit.objects.get(action="modification_plan_etablissement")
        self.assertIn("inclut_prescolaire", entree.details)

    def test_les_statistiques_comptent_le_prescolaire(self):
        Etablissement.objects.filter(pk=self.ecole.pk).update(inclut_prescolaire=True)
        classe = Classe.objects.create(
            nom="PS", cycle=Cycle.PRESCOLAIRE, niveau="presco_ps", annee_scolaire=self.annee,
        )
        eleve = creer_compte(Role.ELEVE, "petit@stat.test", self.ecole, sexe="F", date_naissance=D(2022, 3, 3))
        Inscription.objects.create(eleve=eleve, classe=classe)
        fiche = self.fiche_pour(self.fondateur)
        premiere = fiche["effectifs"][0]
        self.assertEqual((premiere["cycle"], premiere["niveau_libelle"], premiere["total"]), ("prescolaire", "Petite section", 1))
        self.assertEqual(fiche["totaux"]["total"], 10)


class SituationEcoleEtDomaineTests(EducationBase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.dev = creer_compte(Role.DEVELOPPEUR, "dev3@stat.test", cls.ecole)
        cls.fondateur.set_password("MotDePasse#2026")
        cls.fondateur.save()

    def parametres(self, **extra):
        donnees = {
            "nom": "École statistiques A", "devise": "", "code_devise": "FCFA", "pas_montant": 5000,
            "taux_inps_salarie": "3.6", "taux_inps_employeur": "20.0", "taux_amo_salarie": "3.06",
            "taux_amo_employeur": "3.5",
        }
        donnees.update(extra)
        return donnees

    def enregistrer(self, **extra):
        self.client.force_login(self.dev)
        return self.client.post(reverse("espace_developpeur:parametres_etablissement"), self.parametres(**extra))

    def test_une_ecole_est_ouverte_par_defaut(self):
        self.assertEqual(Etablissement.objects.get(pk=self.ecole.pk).situation_fonctionnement, "ouvert")

    def test_fermeture_enregistree_avec_motif_et_date(self):
        reponse = self.enregistrer(situation_fonctionnement="ferme", motif_fermeture="insecurite", date_situation="2026-09-15")
        self.assertEqual(reponse.status_code, 302)
        self.ecole.refresh_from_db()
        self.assertEqual((self.ecole.situation_fonctionnement, self.ecole.motif_fermeture, self.ecole.date_situation),
                         ("ferme", "insecurite", D(2026, 9, 15)))

    def test_reouverture_efface_motif_et_date(self):
        self.enregistrer(situation_fonctionnement="ferme", motif_fermeture="insecurite", date_situation="2026-09-15")
        self.enregistrer(situation_fonctionnement="ouvert", motif_fermeture="insecurite", date_situation="2026-09-15")
        self.ecole.refresh_from_db()
        self.assertEqual((self.ecole.situation_fonctionnement, self.ecole.motif_fermeture, self.ecole.date_situation),
                         ("ouvert", "", None))

    def test_un_ancien_envoi_sans_ce_champ_conserve_la_situation(self):
        self.enregistrer(situation_fonctionnement="relocalise", motif_fermeture="insecurite")
        self.enregistrer()  # formulaire d'avant l'ajout du champ
        self.ecole.refresh_from_db()
        self.assertEqual(self.ecole.situation_fonctionnement, "relocalise")

    def test_valeur_inconnue_refusee(self):
        self.assertEqual(self.enregistrer(situation_fonctionnement="detruite").status_code, 200)
        self.ecole.refresh_from_db()
        self.assertEqual(self.ecole.situation_fonctionnement, "ouvert")

    def test_une_ecole_fermee_garde_ses_connexions(self):
        """La situation est descriptive : seule l'activation par la plateforme bloque la connexion."""
        from django.core.cache import cache
        cache.clear()
        self.addCleanup(cache.clear)
        Etablissement.objects.filter(pk=self.ecole.pk).update(situation_fonctionnement="ferme")
        reponse = self.client.post(reverse("comptes:connexion"), {"email": "fondateur@stat.test", "mot_de_passe": "MotDePasse#2026"})
        self.assertEqual(reponse.status_code, 302)

    def test_la_fiche_annuelle_signale_la_fermeture(self):
        Etablissement.objects.filter(pk=self.ecole.pk).update(situation_fonctionnement="ferme", motif_fermeture="insecurite")
        self.client.force_login(self.fondateur)
        page = self.client.get(reverse("statistiques:fiche_annuelle"), {"annee": self.annee.id})
        self.assertContains(page, "fermé temporairement")
        self.assertContains(page, "insécurité")

    # ---- domaine du CAP / BT ----
    def test_domaine_reserve_aux_classes_professionnelles(self):
        from django.core.exceptions import ValidationError
        with self.assertRaisesMessage(ValidationError, "domaine ne s'applique"):
            Classe(nom="L", cycle=Cycle.LYCEE, domaine_professionnel="industriel", annee_scolaire=self.annee).clean()
        Etablissement.objects.filter(pk=self.ecole.pk).update(inclut_professionnel=True)
        Classe(
            nom="P", cycle=Cycle.PROFESSIONNEL, filiere_professionnelle="bt", domaine_professionnel="tertiaire",
            annee_scolaire=AnneeScolaire.objects.get(pk=self.annee.pk),
        ).clean()  # ne doit pas lever d'exception

    def test_formulaire_de_classe_domaine(self):
        from scolarite.forms import CreerClasseForm
        from scolarite.models import Periodicite
        Etablissement.objects.filter(pk=self.ecole.pk).update(inclut_professionnel=True)
        ecole = Etablissement.objects.get(pk=self.ecole.pk)
        periodicite = next(valeur for valeur in Periodicite.values if valeur != Periodicite.MENSUEL)

        def formulaire(**extra):
            donnees = {
                "nom": "Nouvelle classe", "cycle": "professionnel", "filiere_professionnelle": "cap",
                "annee_scolaire": self.annee.id, "montant_inscription": 0, "periodicite": periodicite, "montant_periode": 0,
            }
            donnees.update(extra)
            return CreerClasseForm(donnees, etablissement=ecole)

        pro = formulaire(domaine_professionnel="tertiaire")
        self.assertTrue(pro.is_valid(), pro.errors)
        self.assertEqual(pro.save().domaine_professionnel, "tertiaire")
        sans_domaine = formulaire(nom="Autre")
        self.assertTrue(sans_domaine.is_valid(), sans_domaine.errors)  # facultatif : les classes existantes restent valides
        hors_pro = formulaire(nom="Lycée", cycle="1er_cycle", filiere_professionnelle="", domaine_professionnel="industriel")
        self.assertFalse(hors_pro.is_valid())
        self.assertIn("domaine_professionnel", hors_pro.errors)


class InterfaceEducationTests(EducationBase):
    """Cases du propriétaire, liens de la liste des classes et de la barre latérale des statistiques."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.proprietaire = creer_compte(Role.DEVELOPPEUR, "owner2@stat.test", None, is_superuser=True, is_staff=True)

    def test_le_proprietaire_voit_les_trois_reglages_de_cycle(self):
        self.client.force_login(self.proprietaire)
        page = self.client.get(reverse("espace_plateforme:liste_etablissements"))
        self.assertContains(page, "<th>Préscolaire</th>")
        self.assertContains(page, 'name="inclut_prescolaire"')
        self.assertContains(page, 'name="inclut_professionnel"')

    def test_les_deux_formulaires_de_formule_conservent_les_autres_reglages(self):
        """Changer la formule ou une case ne doit jamais réinitialiser l'autre case (champs cachés)."""
        Etablissement.objects.filter(pk=self.ecole.pk).update(inclut_professionnel=True, inclut_prescolaire=True)
        self.client.force_login(self.proprietaire)
        page = self.client.get(reverse("espace_plateforme:liste_etablissements")).content.decode()
        # Pour une école qui a les deux options, chaque mini-formulaire renvoie l'état de l'autre option.
        self.assertIn('type="hidden" name="inclut_prescolaire" value="True"', page)
        self.assertIn('type="hidden" name="inclut_professionnel" value="True"', page)
        url = reverse("espace_plateforme:modifier_plan", args=[self.ecole.id])
        self.client.post(url, {"plan": "lycee_seul", "inclut_professionnel": "True", "inclut_prescolaire": "True"})
        self.ecole.refresh_from_db()
        self.assertEqual((self.ecole.plan, self.ecole.inclut_professionnel, self.ecole.inclut_prescolaire), ("lycee_seul", True, True))

    def test_formule_prescolaire_seul_proposee_et_case_remplacee(self):
        Etablissement.objects.filter(pk=self.ecole.pk).update(plan="prescolaire_seul")
        self.client.force_login(self.proprietaire)
        page = self.client.get(reverse("espace_plateforme:liste_etablissements"))
        self.assertContains(page, "Préscolaire seul")
        self.assertContains(page, "Inclus dans la formule")

    def test_ecole_fermee_signalee_au_proprietaire(self):
        Etablissement.objects.filter(pk=self.ecole.pk).update(situation_fonctionnement="relocalise")
        self.client.force_login(self.proprietaire)
        self.assertContains(self.client.get(reverse("espace_plateforme:liste_etablissements")), "École relocalisé")

    def test_liste_des_classes_propose_la_saisie_des_resultats(self):
        Classe.objects.create(nom="Jardin", cycle=Cycle.PRESCOLAIRE, annee_scolaire=self.annee)
        self.client.force_login(self.secretaire)
        page = self.client.get(reverse("scolarite:liste_classes"))
        self.assertContains(page, reverse("scolarite:saisir_resultats_examen", args=[self.c9a.id]))
        self.assertContains(page, "(industriel)")  # domaine de la classe CAP
        # 6 classes préparent un examen (6A, 6B, 9A, TSE, CAP-E et la 5A de l'année précédente) ; pas la maternelle
        self.assertEqual(page.content.decode().count("Résultats d'examen"), 6)

    def test_barre_laterale_des_statistiques(self):
        self.client.force_login(self.fondateur)
        page = self.client.get(reverse("statistiques:vue_ensemble"), {"annee": self.annee.id})
        for lien in ("fiche_annuelle", "resultats_examens", "infrastructures"):
            self.assertContains(page, reverse(f"statistiques:{lien}"))


class SexeDuPersonnelTests(EducationBase):
    """Le sexe des enseignants alimente les statistiques par sexe ; il ne dépend pas du rôle."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.dev = creer_compte(Role.DEVELOPPEUR, "dev4@stat.test", cls.ecole)

    def modifier(self, compte, **donnees):
        self.client.force_login(self.dev)
        base = {"role": compte.role, "statut": "actif"}
        base.update(donnees)
        return self.client.post(reverse("espace_developpeur:modifier_compte", args=[compte.id]), base, follow=True)

    def test_sexe_du_personnel_enregistre_puis_efface(self):
        self.modifier(self.t3, sexe="F")
        self.t3.refresh_from_db()
        self.assertEqual(self.t3.sexe, "F")
        self.assertEqual(self.fiche_pour(self.fondateur)["enseignants"]["par_statut"][-1]["filles"], 1)  # t3 : situation non renseignée
        self.modifier(self.t3, sexe="")
        self.t3.refresh_from_db()
        self.assertIsNone(self.t3.sexe)  # NULL, comme les élèves sans sexe renseigné

    def test_valeur_inconnue_refusee(self):
        reponse = self.modifier(self.t3, sexe="X")
        self.assertEqual(reponse.status_code, 200)
        self.t3.refresh_from_db()
        self.assertIsNone(self.t3.sexe)

    def test_le_sexe_n_est_pas_efface_par_un_changement_de_role(self):
        self.modifier(self.t1, role="parent")  # t1 : enseignante, sexe F, formée
        self.t1.refresh_from_db()
        self.assertEqual(self.t1.sexe, "F")
        self.assertEqual(self.t1.formation_pedagogique, "")  # la fiche du personnel, elle, est effacée

    def test_le_formulaire_d_un_eleve_ne_propose_pas_le_sexe(self):
        self.client.force_login(self.dev)
        page = self.client.get(reverse("espace_developpeur:modifier_compte", args=[self.e1.id]))
        self.assertNotContains(page, "Sexe (personnel)")
        self.modifier(self.e1, sexe="M")  # envoi forcé : ignoré (le sexe d'un élève se saisit à l'inscription)
        self.e1.refresh_from_db()
        self.assertEqual(self.e1.sexe, "F")

    def test_page_des_examens_sans_calcul_de_la_fiche_entiere(self):
        self.client.force_login(self.fondateur)
        page = self.client.get(reverse("statistiques:resultats_examens"), {"annee": self.annee.id})
        self.assertContains(page, "Certificat d&#x27;entrée en 7ème année")
        self.assertContains(page, "Sciences Exactes (TSE)")
