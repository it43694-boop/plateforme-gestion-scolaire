"""
Export Excel (.xlsx) de la fiche statistique annuelle.

Les lignes sont EXACTEMENT celles du CSV (`fiche.lignes_csv`) : les deux exports ne peuvent donc pas diverger, et la
confidentialité (rubriques réservées à l'administration, cycle limité d'un directeur) est décidée une seule fois, à la
construction de la fiche. Ici on ne fait que la mise en forme : titres de rubriques, en-têtes de tableaux, largeurs de
colonnes, impression en paysage sur la largeur d'une page.

Les rubriques sont séparées par une ligne vide ; la première ligne d'une rubrique est son titre, la suivante ses
en-têtes quand c'est un tableau.
"""
from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

from statistiques.fiche import lignes_csv

TYPE_MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_TITRE = Font(bold=True, size=12, color="FFFFFF")
_FOND_TITRE = PatternFill("solid", fgColor="1F4E78")
_GRAS = Font(bold=True)
_FOND_ENTETE = PatternFill("solid", fgColor="DDEBF7")
_FILET = Side(style="thin", color="999999")
_BORDURE = Border(left=_FILET, right=_FILET, top=_FILET, bottom=_FILET)
_LARGEUR_MIN, _LARGEUR_MAX = 10, 48
_COLONNES_TITRE = 6  # le bandeau de titre couvre au moins la largeur d'un tableau d'effectifs


def _est_entete(ligne):
    """Ligne d'en-têtes d'un tableau : au moins 3 cellules, toutes du texte."""
    return len(ligne) >= 3 and all(isinstance(cellule, str) and cellule for cellule in ligne)


def _ecrire(feuille, numero, ligne):
    for colonne, valeur in enumerate(ligne, start=1):
        if valeur is None or valeur == "":
            continue
        cellule = feuille.cell(row=numero, column=colonne, value=valeur)
        if isinstance(valeur, str):
            # Toujours du texte, jamais une formule : openpyxl traiterait sinon un nom d'école ou de classe commençant
            # par « = » comme une formule que le tableur exécuterait.
            cellule.data_type = "s"


def _mettre_en_forme(feuille, numero, ligne, rang, premiere_rubrique):
    """Style d'une ligne selon son rôle ; renvoie True si c'est une ligne de titre (exclue du calcul des largeurs)."""
    if rang == 0:
        for colonne in range(1, max(len(ligne), _COLONNES_TITRE) + 1):
            cellule = feuille.cell(row=numero, column=colonne)
            cellule.font, cellule.fill = _TITRE, _FOND_TITRE
        return True
    if rang == 1 and _est_entete(ligne):
        for colonne in range(1, len(ligne) + 1):
            cellule = feuille.cell(row=numero, column=colonne)
            cellule.font, cellule.fill, cellule.border = _GRAS, _FOND_ENTETE, _BORDURE
            cellule.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        return False
    if ligne[0] == "TOTAL":
        for colonne in range(1, len(ligne) + 1):
            cellule = feuille.cell(row=numero, column=colonne)
            cellule.font, cellule.border = _GRAS, _BORDURE
        return False
    if premiere_rubrique:  # identification de l'établissement : « libellé ; valeur »
        feuille.cell(row=numero, column=1).font = _GRAS
    # Les libellés de la première colonne, parfois longs, passent à la ligne au lieu de déborder sur les chiffres.
    feuille.cell(row=numero, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    return False


def classeur_fiche(fiche) -> Workbook:
    classeur = Workbook()
    classeur.properties.creator = getattr(settings, "NOM_PLATEFORME", "Nexora Éducation")
    feuille = classeur.active
    feuille.title = "Fiche statistique"

    largeurs = {}
    rang = None  # rang de la ligne dans sa rubrique ; None = on attend le début d'une rubrique
    premiere_rubrique = True
    for numero, ligne in enumerate(lignes_csv(fiche), start=1):
        if not ligne:
            if rang is not None:
                premiere_rubrique = False
            rang = None
            continue
        rang = 0 if rang is None else rang + 1
        _ecrire(feuille, numero, ligne)
        est_titre = _mettre_en_forme(feuille, numero, ligne, rang, premiere_rubrique)
        if not est_titre:
            for colonne, valeur in enumerate(ligne, start=1):
                if valeur not in (None, ""):
                    largeurs[colonne] = max(largeurs.get(colonne, 0), len(str(valeur)))

    for colonne, longueur in largeurs.items():
        lettre = feuille.cell(row=1, column=colonne).column_letter
        feuille.column_dimensions[lettre].width = min(max(longueur + 2, _LARGEUR_MIN), _LARGEUR_MAX)

    feuille.page_setup.orientation = "landscape"
    feuille.page_setup.fitToWidth = 1
    feuille.page_setup.fitToHeight = 0
    feuille.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    return classeur
