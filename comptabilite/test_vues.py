"""
Pages de la comptabilité : accès par rôle, cloisonnement entre écoles, saisie, journal, états, clôture, export Excel.
"""
import io

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from comptabilite import etats
from comptabilite.models import CompteComptable, Ecriture, Exercice, Journal
from comptabilite.services import creer_compte, creer_ecriture, creer_exercice, enregistrer_operation
from comptabilite.test_services import ANNEE, ComptaBase, compte_utilisateur, jour
from comptes.models import JournalAudit
from comptes.roles import Role
from permissions_matrix.models import PermissionMatrix

URLS_LECTURE = (
    "comptabilite:accueil", "comptabilite:exercices", "comptabilite:plan_comptable", "comptabilite:journal",
    "comptabilite:grand_livre", "comptabilite:balance", "comptabilite:compte_de_resultat", "comptabilite:bilan",
    "comptabilite:exporter_etats",
)
URLS_ECRITURE = ("comptabilite:saisir_ecriture", "comptabilite:operation_simple")


class VuesBase(ComptaBase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        for ecole in (cls.ecole, cls.ecole_b):
            PermissionMatrix.seed_pour(ecole)
        cls.fondateur = compte_utilisateur("fond-v@compta.test", Role.FONDATEUR, cls.ecole)
        cls.fondateur_b = compte_utilisateur("fond-vb@compta.test", Role.FONDATEUR, cls.ecole_b)
        cls.comptable_b = compte_utilisateur("comptable-vb@compta.test", Role.COMPTABLE, cls.ecole_b)
        cls.secretaire = compte_utilisateur("secr-v@compta.test", Role.SECRETAIRE, cls.ecole)
        cls.enseignant = compte_utilisateur("prof-v@compta.test", Role.ENSEIGNANT, cls.ecole)
        cls.parent = compte_utilisateur("parent-v@compta.test", Role.PARENT, cls.ecole, telephone="70888001")
        cls.super_admin = compte_utilisateur("super-v@compta.test", Role.SUPER_ADMINISTRATEUR, cls.ecole)
        cls.proprietaire = compte_utilisateur("owner-v@compta.test", Role.DEVELOPPEUR, None, is_superuser=True)

    def connecte(self, compte):
        self.client.force_login(compte)
        return self.client


class AccesTests(VuesBase):
    def test_direction_et_comptable_lisent_et_ecrivent(self):
        for compte in (self.fondateur, self.comptable):
            self.connecte(compte)
            for nom in URLS_LECTURE + URLS_ECRITURE:
                with self.subTest(role=compte.role, page=nom):
                    self.assertEqual(self.client.get(reverse(nom)).status_code, 200)

    def test_roles_sans_le_module_caisse_refuses(self):
        for compte in (self.secretaire, self.enseignant, self.parent):
            self.connecte(compte)
            for nom in URLS_LECTURE + URLS_ECRITURE:
                with self.subTest(role=compte.role, page=nom):
                    self.assertEqual(self.client.get(reverse(nom)).status_code, 403)

    def test_un_role_en_lecture_seule_ne_passe_pas_d_ecriture(self):
        """Le super administrateur a le module Caisse (consultation) mais n'écrit pas."""
        self.connecte(self.super_admin)
        for nom in URLS_LECTURE:
            self.assertEqual(self.client.get(reverse(nom)).status_code, 200, nom)
        for nom in URLS_ECRITURE:
            self.assertEqual(self.client.get(reverse(nom)).status_code, 403, nom)
        for nom, donnees in [("comptabilite:exercices", {"libelle": "X", "date_debut": "2040-01-01", "date_fin": "2040-12-31"}),
                             ("comptabilite:plan_comptable", {"numero": "52101", "libelle": "Banque"})]:
            self.assertEqual(self.client.post(reverse(nom), donnees).status_code, 403, nom)
        self.assertFalse(CompteComptable.objects.filter(numero="52101", etablissement=self.ecole).exists())
        self.assertFalse(Exercice.objects.filter(libelle="X").exists())
        # et ne voit pas les boutons de saisie
        self.assertNotContains(self.client.get(reverse("comptabilite:accueil")), reverse("comptabilite:saisir_ecriture"))

    def test_anonyme_redirige(self):
        self.client.logout()
        for nom in URLS_LECTURE + URLS_ECRITURE:
            self.assertEqual(self.client.get(reverse(nom)).status_code, 302, nom)

    def test_le_proprietaire_de_la_plateforme_sans_etablissement_est_refuse(self):
        self.connecte(self.proprietaire)
        for nom in URLS_LECTURE + URLS_ECRITURE:
            with self.subTest(page=nom):
                reponse = self.client.get(reverse(nom))
                # renvoyé vers son espace plateforme (cloisonnement existant) ou refusé : jamais servi
                self.assertIn(reponse.status_code, (302, 403))
                if reponse.status_code == 302:
                    self.assertEqual(reponse["Location"], "/plateforme/")

    def test_les_actions_d_ecriture_exigent_le_post(self):
        self.connecte(self.fondateur)
        ecriture = self.ecrire([("622", 100, 0), ("401", 0, 100)])
        for nom, args in [("comptabilite:synchroniser", []), ("comptabilite:cloturer", [self.exercice.id]),
                          ("comptabilite:reporter", [self.exercice.id]), ("comptabilite:contrepasser", [ecriture.id])]:
            self.assertEqual(self.client.get(reverse(nom, args=args)).status_code, 405, nom)


class CloisonnementTests(VuesBase):
    def test_les_donnees_d_une_ecole_ne_fuient_pas_vers_une_autre(self):
        creer_ecriture(
            etablissement=self.ecole, journal=Journal.OPERATIONS_DIVERSES, date_ecriture=jour(3, 1), piece="PIECE-ZQXA",
            libelle="Libellé secret zqxjka", lignes=[("521", 1000, 0), ("101", 0, 1000)],
        )
        self.connecte(self.fondateur_b)
        for nom in URLS_LECTURE:
            if nom == "comptabilite:exporter_etats":
                continue
            reponse = self.client.get(reverse(nom), {"exercice": self.exercice_b.id})
            texte = reponse.content.decode("utf-8")
            with self.subTest(page=nom):
                self.assertEqual(reponse.status_code, 200)
                self.assertNotIn("zqxjka", texte)
                self.assertNotIn("PIECE-ZQXA", texte)

    def test_exercice_ecriture_et_compte_d_une_autre_ecole_introuvables(self):
        ecriture = self.ecrire([("622", 100, 0), ("401", 0, 100)])
        self.connecte(self.fondateur_b)
        attendu = 404
        for nom in URLS_LECTURE:
            self.assertEqual(self.client.get(reverse(nom), {"exercice": self.exercice.id}).status_code, attendu, nom)
        for nom, args in [("comptabilite:cloturer", [self.exercice.id]), ("comptabilite:reporter", [self.exercice.id]),
                          ("comptabilite:contrepasser", [ecriture.id])]:
            self.assertEqual(self.client.post(reverse(nom, args=args)).status_code, attendu, nom)
        creer_compte(etablissement=self.ecole, numero="52101", libelle="Compte propre à A")
        self.assertEqual(self.client.get(reverse("comptabilite:grand_livre"), {"compte": "52101"}).status_code, 404)  # compte de A
        # rien n'a bougé dans l'école A
        self.assertFalse(Exercice.objects.get(pk=self.exercice.pk).cloture)
        self.assertFalse(hasattr(Ecriture.objects.get(pk=ecriture.pk), "contrepassation"))

    def test_meme_reponse_pour_un_identifiant_inexistant(self):
        self.connecte(self.fondateur_b)
        ecriture = self.ecrire([("622", 100, 0), ("401", 0, 100)])
        for nom, args in [("comptabilite:cloturer", [self.exercice.id]), ("comptabilite:cloturer", [999999]),
                          ("comptabilite:contrepasser", [ecriture.id]), ("comptabilite:contrepasser", [999999])]:
            self.assertEqual(self.client.post(reverse(nom, args=args)).status_code, 404)

    def test_identifiants_invalides(self):
        self.connecte(self.fondateur)
        for nom in URLS_LECTURE:
            self.assertEqual(self.client.get(reverse(nom), {"exercice": "abc"}).status_code, 404, nom)
        for parametres in ({"du": "pas-une-date"}, {"au": "31/12/2026"}, {"journal": "ZZZ"}):
            self.assertEqual(self.client.get(reverse("comptabilite:journal"), parametres).status_code, 404, parametres)
        self.assertEqual(self.client.get(reverse("comptabilite:balance"), {"jusqu_au": "x"}).status_code, 404)

    def test_le_plan_comptable_est_propre_a_chaque_ecole(self):
        self.connecte(self.fondateur)
        self.client.post(reverse("comptabilite:plan_comptable"), {"numero": "52101", "libelle": "Banque BDM zqxjka"})
        self.connecte(self.fondateur_b)
        self.assertNotContains(self.client.get(reverse("comptabilite:plan_comptable")), "zqxjka")


class SaisieTests(VuesBase):
    def test_operation_simple_depuis_la_page(self):
        self.connecte(self.comptable)
        compte = CompteComptable.objects.get(etablissement=self.ecole, numero="605")
        reponse = self.client.post(reverse("comptabilite:operation_simple"), {
            "type_operation": "depense", "date_ecriture": jour(3, 10).isoformat(), "montant": 50000, "compte": compte.pk,
            "mode": "especes", "libelle": "Facture d'électricité", "piece": "F-1",
        })
        self.assertRedirects(reponse, f"{reverse('comptabilite:journal')}?exercice={self.exercice.id}")
        ecriture = Ecriture.objects.get()
        self.assertEqual((ecriture.journal, ecriture.piece, ecriture.cree_par), (Journal.CAISSE, "F-1", self.comptable))
        self.assertEqual({l.compte.numero: (l.debit, l.credit) for l in ecriture.lignes.select_related("compte")},
                         {"605": (50000, 0), "571": (0, 50000)})
        entree = JournalAudit.objects.get(action="ecriture_comptable")
        self.assertEqual((entree.cible, entree.details["montant"], entree.details["type"]), (ecriture.numero, 50000, "depense"))

    def test_operation_simple_les_comptes_de_tresorerie_ne_sont_pas_proposes(self):
        self.connecte(self.comptable)
        formulaire = self.client.get(reverse("comptabilite:operation_simple")).context["formulaire"]
        numeros = {c.numero for c in formulaire.fields["compte"].queryset}
        self.assertIn("605", numeros)
        self.assertFalse({n for n in numeros if n[0] in "58"})
        autre_ecole = CompteComptable.objects.get(etablissement=self.ecole_b, numero="605")
        self.assertNotIn(autre_ecole, formulaire.fields["compte"].queryset)

    def test_ecriture_libre_depuis_la_page(self):
        self.connecte(self.fondateur)
        donnees = {
            "journal": "OD", "date_ecriture": jour(2, 1).isoformat(), "libelle": "Achat de mobilier", "piece": "BC-7",
            "compte_1": "244", "debit_1": "400 000", "compte_2": "521 - Banques locales", "credit_2": "400000", "libelle_2": "Chèque",
        }
        reponse = self.client.post(reverse("comptabilite:saisir_ecriture"), donnees)
        self.assertRedirects(reponse, f"{reverse('comptabilite:journal')}?exercice={self.exercice.id}")
        ecriture = Ecriture.objects.get()
        self.assertEqual({l.compte.numero: (l.debit, l.credit) for l in ecriture.lignes.select_related("compte")},
                         {"244": (400000, 0), "521": (0, 400000)})

    def test_ecriture_libre_desequilibree_ou_invalide_est_refusee_sans_rien_creer(self):
        self.connecte(self.fondateur)
        base = {"journal": "OD", "date_ecriture": jour(2, 1).isoformat(), "libelle": "Test"}
        for attendu, lignes in [
            ("déséquilibrée", {"compte_1": "244", "debit_1": "100", "compte_2": "521", "credit_2": "90"}),
            ("existe pas dans votre plan", {"compte_1": "244", "debit_1": "100", "compte_2": "999999", "credit_2": "100"}),
            ("Montant invalide", {"compte_1": "244", "debit_1": "abc", "compte_2": "521", "credit_2": "100"}),
            ("Indiquez le compte", {"debit_1": "100", "compte_2": "521", "credit_2": "100"}),
            ("au moins deux lignes", {"compte_1": "244", "debit_1": "100"}),
        ]:
            reponse = self.client.post(reverse("comptabilite:saisir_ecriture"), {**base, **lignes})
            with self.subTest(attendu=attendu):
                self.assertEqual(reponse.status_code, 200)
                self.assertContains(reponse, attendu)
        self.assertFalse(Ecriture.objects.exists())

    def test_date_hors_exercice_refusee(self):
        self.connecte(self.fondateur)
        reponse = self.client.post(reverse("comptabilite:saisir_ecriture"), {
            "journal": "OD", "date_ecriture": jour(1, 1, ANNEE + 5).isoformat(), "libelle": "Trop tard",
            "compte_1": "244", "debit_1": "100", "compte_2": "521", "credit_2": "100",
        })
        self.assertContains(reponse, "Aucun exercice")
        self.assertFalse(Ecriture.objects.exists())

    def test_contre_passation_depuis_le_journal(self):
        ecriture = self.ecrire([("622", 80000, 0), ("401", 0, 80000)], libelle="Loyer")
        self.connecte(self.comptable)
        page = self.client.get(reverse("comptabilite:journal"), {"exercice": self.exercice.id})
        self.assertContains(page, "Contre-passer")
        self.client.post(reverse("comptabilite:contrepasser", args=[ecriture.id]), {"motif": "Doublon"})
        self.assertEqual(Ecriture.objects.count(), 2)
        page = self.client.get(reverse("comptabilite:journal"), {"exercice": self.exercice.id})
        self.assertContains(page, "Annulée par")
        self.assertContains(page, "Annule une écriture")
        self.client.post(reverse("comptabilite:contrepasser", args=[ecriture.id]), {"motif": "Encore"})  # déjà contre-passée
        self.assertEqual(Ecriture.objects.count(), 2)
        self.assertTrue(JournalAudit.objects.filter(action="contrepassation_comptable").exists())

    def test_le_journal_n_offre_pas_la_contre_passation_des_ecritures_automatiques(self):
        creer_ecriture(etablissement=self.ecole, journal=Journal.CAISSE, date_ecriture=jour(3, 1), libelle="Auto",
                       lignes=[("571", 10, 0), ("70602", 0, 10)], origine=Ecriture.Origine.AUTOMATIQUE)
        self.connecte(self.fondateur)
        self.assertNotContains(self.client.get(reverse("comptabilite:journal"), {"exercice": self.exercice.id}), "Contre-passer")


class ExercicesEtPlanTests(VuesBase):
    def test_creation_d_un_exercice_et_chevauchement(self):
        self.connecte(self.fondateur)
        reponse = self.client.post(reverse("comptabilite:exercices"), {
            "libelle": f"Exercice {ANNEE + 1}", "date_debut": jour(1, 1, ANNEE + 1).isoformat(), "date_fin": jour(12, 31, ANNEE + 1).isoformat(),
        })
        self.assertRedirects(reponse, reverse("comptabilite:exercices"))
        self.assertTrue(Exercice.objects.filter(etablissement=self.ecole, libelle=f"Exercice {ANNEE + 1}").exists())
        reponse = self.client.post(reverse("comptabilite:exercices"), {
            "libelle": "Chevauche", "date_debut": jour(6, 1).isoformat(), "date_fin": jour(6, 1, ANNEE + 1).isoformat(),
        })
        self.assertContains(reponse, "chevauche")
        self.assertTrue(JournalAudit.objects.filter(action="creation_exercice_comptable").exists())

    def test_cloture_et_report_depuis_la_page(self):
        creer_exercice(etablissement=self.ecole, libelle=f"Exercice {ANNEE + 1}", date_debut=jour(1, 1, ANNEE + 1), date_fin=jour(12, 31, ANNEE + 1))
        enregistrer_operation(etablissement=self.ecole, type_operation="recette", date_ecriture=jour(3, 5), montant=300000,
                              compte="70602", mode="especes", libelle="Scolarité")
        self.connecte(self.fondateur)
        page = self.client.get(reverse("comptabilite:exercices"))
        self.assertContains(page, "Clôturer")
        self.client.post(reverse("comptabilite:cloturer", args=[self.exercice.id]))
        self.assertTrue(Exercice.objects.get(pk=self.exercice.pk).cloture)
        self.assertTrue(Ecriture.objects.filter(origine=Ecriture.Origine.CLOTURE).exists())
        self.assertContains(self.client.get(reverse("comptabilite:exercices")), "Reporter les soldes")
        self.client.post(reverse("comptabilite:reporter", args=[self.exercice.id]))
        self.assertTrue(Ecriture.objects.filter(origine=Ecriture.Origine.REPORT).exists())
        self.assertContains(self.client.get(reverse("comptabilite:exercices")), "Soldes reportés")
        for action in ("cloture_exercice_comptable", "report_a_nouveau_comptable"):
            self.assertTrue(JournalAudit.objects.filter(action=action).exists(), action)
        # un exercice clôturé n'offre plus les boutons de saisie
        page = self.client.get(reverse("comptabilite:accueil"), {"exercice": self.exercice.id})
        self.assertNotContains(page, reverse("comptabilite:operation_simple"))
        self.assertContains(page, "Clôturé")

    def test_plan_comptable_recherche_et_ajout(self):
        self.connecte(self.comptable)
        page = self.client.get(reverse("comptabilite:plan_comptable"), {"q": "caisse"})
        self.assertContains(page, "571")
        self.assertNotContains(page, "Capital social")
        self.client.post(reverse("comptabilite:plan_comptable"), {"numero": "52101", "libelle": "Banque BDM"})
        self.assertTrue(CompteComptable.objects.filter(etablissement=self.ecole, numero="52101", systeme=False).exists())
        reponse = self.client.post(reverse("comptabilite:plan_comptable"), {"numero": "52101", "libelle": "Doublon"})
        self.assertContains(reponse, "existe déjà")
        self.assertContains(self.client.get(reverse("comptabilite:plan_comptable")), "Ajouté par l")

    def test_accueil_sans_exercice(self):
        Exercice.objects.filter(etablissement=self.ecole).delete()
        self.connecte(self.fondateur)
        page = self.client.get(reverse("comptabilite:accueil"))
        self.assertContains(page, "Aucun exercice comptable")
        for nom in URLS_LECTURE:
            if nom != "comptabilite:exporter_etats":
                self.assertEqual(self.client.get(reverse(nom)).status_code, 200, nom)
        self.assertEqual(self.client.get(reverse("comptabilite:exporter_etats")).status_code, 404)


class EtatsEtExportTests(VuesBase):
    def setUp(self):
        self.ecrire([("521", 1000000, 0), ("101", 0, 1000000)], date=jour(1, 5))
        self.ecrire([("244", 400000, 0), ("521", 0, 400000)], date=jour(2, 1))
        enregistrer_operation(etablissement=self.ecole, type_operation="recette", date_ecriture=jour(3, 5), montant=300000,
                              compte="70602", mode="especes", libelle="Scolarité")
        enregistrer_operation(etablissement=self.ecole, type_operation="depense", date_ecriture=jour(3, 10), montant=50000,
                              compte="605", mode="especes", libelle="Électricité")

    def test_pages_d_etats(self):
        self.connecte(self.comptable)
        balance = self.client.get(reverse("comptabilite:balance"))
        self.assertContains(balance, "Équilibrée")
        contenu = balance.content.decode("utf-8")
        # total des débits, groupé par la locale (espaces insécables) : on compare sans espaces
        self.assertIn("1750000", "".join(c for c in contenu if not c.isspace()))
        resultat = self.client.get(reverse("comptabilite:compte_de_resultat"))
        self.assertContains(resultat, "Résultat net de l")
        self.assertContains(resultat, "bénéfice")
        bilan = self.client.get(reverse("comptabilite:bilan"))
        self.assertContains(bilan, "Actif et passif sont égaux")
        grand_livre = self.client.get(reverse("comptabilite:grand_livre"), {"compte": "571"})
        self.assertContains(grand_livre, "Scolarité")
        self.assertContains(self.client.get(reverse("comptabilite:journal")), "OD-")
        self.assertEqual(self.client.get(reverse("comptabilite:grand_livre"), {"compte": "999999"}).status_code, 404)

    def test_export_excel_de_tous_les_etats(self):
        self.connecte(self.comptable)
        reponse = self.client.get(reverse("comptabilite:exporter_etats"), {"exercice": self.exercice.id})
        self.assertEqual(reponse.status_code, 200)
        self.assertIn("spreadsheetml.sheet", reponse["Content-Type"])
        self.assertIn("etats_comptables_Exercice_", reponse["Content-Disposition"])
        classeur = load_workbook(io.BytesIO(reponse.content))
        self.assertEqual(classeur.sheetnames, ["Journal", "Grand livre", "Balance", "Compte de résultat", "Bilan"])

        def lignes(nom):
            return [[c.value for c in rangee] for rangee in classeur[nom].iter_rows()]

        journal = [l for l in lignes("Journal") if l[1] and str(l[1]).startswith(("OD-", "CAI-"))]
        self.assertEqual(len(journal), 8)  # 4 écritures de 2 lignes
        total = next(l for l in lignes("Journal") if l[4] == "Total")
        self.assertEqual((total[7], total[8]), (1750000, 1750000))
        balance = {l[0]: l for l in lignes("Balance") if l[0] and str(l[0]).isdigit()}
        self.assertEqual((balance["521"][4], balance["571"][4], balance["101"][5]), (600000, 250000, 1000000))
        resultat = lignes("Compte de résultat")
        self.assertEqual(next(l for l in resultat if l[0] and str(l[0]).startswith("Résultat net"))[3], 250000)
        bilan = lignes("Bilan")
        self.assertEqual({l[1]: l[3] for l in bilan if l[1] in ("TOTAL ACTIF", "TOTAL PASSIF")}, {"TOTAL ACTIF": 1250000, "TOTAL PASSIF": 1250000})
        self.assertTrue(any(l[0] == "Équilibre actif / passif" and l[1] == "Équilibré" for l in bilan))
        self.assertTrue(JournalAudit.objects.filter(action="export_etats_comptables").exists())

    def test_l_export_n_est_pas_ouvert_aux_roles_sans_caisse(self):
        self.connecte(self.secretaire)
        self.assertEqual(self.client.get(reverse("comptabilite:exporter_etats")).status_code, 403)

    def test_la_navigation_propose_la_comptabilite_aux_roles_autorises(self):
        self.connecte(self.comptable)
        self.assertContains(self.client.get(reverse("comptabilite:accueil")), reverse("comptabilite:accueil"))
        self.connecte(self.secretaire)
        self.assertNotContains(self.client.get(reverse("scolarite:liste_eleves")), reverse("comptabilite:accueil"))

    def test_alerte_quand_la_caisse_n_est_pas_comptabilisee(self):
        from finances.models import TypeTranche, enregistrer_paiement
        from scolarite.models import AnneeScolaire, Classe, Cycle, Inscription
        annee = AnneeScolaire.objects.create(etablissement=self.ecole, libelle="2026-2027", date_debut=jour(10, 1),
                                             date_fin=jour(7, 31, ANNEE + 1), est_active=True)
        classe = Classe.objects.create(nom="6A", cycle=Cycle.PREMIER_CYCLE, annee_scolaire=annee)
        eleve = compte_utilisateur("eleve-alerte@compta.test", Role.ELEVE, self.ecole, matricule="202699777")
        inscription = Inscription.objects.create(eleve=eleve, classe=classe)
        enregistrer_paiement(eleve=eleve, inscription=inscription, tranche=TypeTranche.VERSEMENT, montant=20000,
                             mode_paiement="especes", enregistre_par=self.comptable)
        self.connecte(self.comptable)
        page = self.client.get(reverse("comptabilite:bilan"))
        self.assertContains(page, "1 mouvement de caisse")
        self.assertContains(page, "Mettre à jour la comptabilité")
        reponse = self.client.post(reverse("comptabilite:synchroniser"), {"retour": reverse("comptabilite:bilan")}, follow=True)
        self.assertContains(reponse, "1 écriture passée depuis la caisse")
        self.assertNotContains(reponse, "Mettre à jour la comptabilité")
        reponse = self.client.post(reverse("comptabilite:synchroniser"), follow=True)
        self.assertContains(reponse, "déjà à jour")
        # le lien de retour ne peut pas rediriger vers un autre site
        reponse = self.client.post(reverse("comptabilite:synchroniser"), {"retour": "https://pirate.example/"})
        self.assertRedirects(reponse, reverse("comptabilite:accueil"))
        reponse = self.client.post(reverse("comptabilite:synchroniser"), {"retour": "//pirate.example/"})
        self.assertRedirects(reponse, reverse("comptabilite:accueil"))
        self.assertTrue(JournalAudit.objects.filter(action="synchronisation_comptable").exists())

    def test_pas_de_cache(self):
        self.connecte(self.comptable)
        for nom in ("comptabilite:accueil", "comptabilite:balance", "comptabilite:exporter_etats"):
            self.assertIn("no-store", self.client.get(reverse(nom))["Cache-Control"], nom)
