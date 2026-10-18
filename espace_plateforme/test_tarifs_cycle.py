"""Tarifs par cycle du propriétaire de la plateforme : par élève et par an, ou forfait par an, selon l'école ; le montant proposé pour la facture en découle."""
import io

from django.urls import reverse
from openpyxl import load_workbook

from comptes.models import JournalAudit
from comptes.outils_test_arabe import client_en, restes_hors_textes_saisis
from comptes.roles import Role
from espace_plateforme.effectifs import effectifs_de_l_ecole, effectifs_de_toutes_les_ecoles, total_par_cycle
from espace_plateforme.tests import creer_utilisateur_actif
from etablissement.models import Etablissement, PlanEtablissement, TarifCycle
from scolarite.models import Cycle
from statistiques.test_education import EducationBase, creer_compte


def par_cycle(effectifs):
    return {ligne.cycle: ligne for ligne in effectifs.lignes}


class CalculDuMontantTests(EducationBase):
    def tarif(self, cycle, mode, montant, ecole=None):
        return TarifCycle.objects.create(etablissement=ecole or self.ecole, cycle=cycle, mode=mode, montant=montant)

    def test_par_eleve_le_montant_est_multiplie_par_les_eleves_inscrits(self):
        self.tarif(Cycle.PREMIER_CYCLE, TarifCycle.Mode.PAR_ELEVE, 2500)
        ligne = par_cycle(effectifs_de_l_ecole(self.ecole))[Cycle.PREMIER_CYCLE]
        self.assertEqual(ligne.eleves, 4)
        self.assertEqual(ligne.montant_estime, 4 * 2500)

    def test_au_forfait_le_montant_s_applique_une_fois_pour_un_cycle_ouvert(self):
        self.tarif(Cycle.DEUXIEME_CYCLE, TarifCycle.Mode.FORFAIT, 150000)
        self.assertEqual(par_cycle(effectifs_de_l_ecole(self.ecole))[Cycle.DEUXIEME_CYCLE].montant_estime, 150000)  # 2 élèves ou 200 : le même forfait

    def test_un_forfait_ne_s_applique_pas_a_un_cycle_ferme(self):
        # le préscolaire n'est ni dans la formule ni habité : il n'est pas ouvert
        self.tarif(Cycle.PRESCOLAIRE, TarifCycle.Mode.FORFAIT, 90000)
        ligne = par_cycle(effectifs_de_l_ecole(self.ecole))[Cycle.PRESCOLAIRE]
        self.assertFalse(ligne.cycle_ouvert)
        self.assertEqual(ligne.montant_estime, 0)
        self.assertIn("Préscolaire", [l.libelle for l in effectifs_de_l_ecole(self.ecole).lignes_utiles])  # mais le tarif saisi reste visible

    def test_un_cycle_hors_formule_qui_a_des_comptes_est_ouvert(self):
        self.ecole.plan = PlanEtablissement.PREMIER_CYCLE
        self.ecole.save()
        self.tarif(Cycle.LYCEE, TarifCycle.Mode.FORFAIT, 120000)
        ligne = par_cycle(effectifs_de_l_ecole(self.ecole))[Cycle.LYCEE]
        self.assertTrue(ligne.hors_formule)
        self.assertEqual(ligne.montant_estime, 120000)

    def test_sans_tarif_pas_de_montant(self):
        effectifs = effectifs_de_l_ecole(self.ecole)
        self.assertEqual(effectifs.montant_estime, 0)
        self.assertFalse(effectifs.a_des_tarifs)

    def test_chaque_ecole_a_son_mode_et_son_montant(self):
        self.tarif(Cycle.PREMIER_CYCLE, TarifCycle.Mode.PAR_ELEVE, 2000)
        self.tarif(Cycle.PREMIER_CYCLE, TarifCycle.Mode.FORFAIT, 300000, ecole=self.ecole_b)
        self.assertEqual(effectifs_de_l_ecole(self.ecole).montant_estime, 4 * 2000)
        self.assertEqual(effectifs_de_l_ecole(self.ecole_b).montant_estime, 300000)  # un forfait : le cycle de la formule est ouvert, même sans élève

    def test_le_total_de_l_ecole_et_celui_de_la_plateforme(self):
        self.tarif(Cycle.PREMIER_CYCLE, TarifCycle.Mode.PAR_ELEVE, 2500)
        self.tarif(Cycle.LYCEE, TarifCycle.Mode.PAR_ELEVE, 4000)
        self.tarif(Cycle.PREMIER_CYCLE, TarifCycle.Mode.FORFAIT, 100000, ecole=self.ecole_b)
        ecoles = effectifs_de_toutes_les_ecoles(Etablissement.objects.order_by("nom"))
        self.assertEqual(ecoles[0].montant_estime, 4 * 2500 + 2 * 4000)
        totaux = {ligne.cycle: ligne for ligne in total_par_cycle(ecoles)}
        self.assertEqual(totaux[Cycle.PREMIER_CYCLE].montant_cumule, 4 * 2500 + 100000)

    def test_un_seul_tarif_par_ecole_et_par_cycle(self):
        from django.db import IntegrityError, transaction

        self.tarif(Cycle.PREMIER_CYCLE, TarifCycle.Mode.PAR_ELEVE, 2500)
        with self.assertRaises(IntegrityError), transaction.atomic():
            self.tarif(Cycle.PREMIER_CYCLE, TarifCycle.Mode.FORFAIT, 1)

    def test_un_cycle_inconnu_est_refuse(self):
        from django.core.exceptions import ValidationError

        with self.assertRaises(ValidationError):
            TarifCycle(etablissement=self.ecole, cycle="universite", mode=TarifCycle.Mode.FORFAIT, montant=10).full_clean()


class PageDesTarifsTests(EducationBase):
    def setUp(self):
        self.proprietaire = creer_utilisateur_actif("tarifs-effectifs@example.com", Role.DEVELOPPEUR, is_superuser=True)
        self.client.force_login(self.proprietaire)
        self.adresse = reverse("espace_plateforme:tarifs_cycle", args=[self.ecole.pk])

    def donnees(self, **autres):
        base = {f"mode_{cycle}": "" for cycle, _ in Cycle.choices}
        base.update({f"montant_{cycle}": "" for cycle, _ in Cycle.choices})
        base.update(autres)
        return base

    def test_la_page_propose_un_tarif_par_cycle(self):
        page = self.client.get(self.adresse)
        self.assertEqual(page.status_code, 200)
        for attendu in ("Tarifs par cycle", "École statistiques A", "Par élève et par an", "Forfait par an", "Pas de tarif", "Enregistrer les tarifs", "1er cycle", "Lycée"):
            self.assertContains(page, attendu)

    def test_enregistrer_des_tarifs(self):
        reponse = self.client.post(self.adresse, self.donnees(mode_1er_cycle="par_eleve", montant_1er_cycle="2500", mode_lycee="forfait", montant_lycee="200000"))
        self.assertRedirects(reponse, reverse("espace_plateforme:effectifs"), fetch_redirect_response=False)
        tarifs = {t.cycle: (t.mode, t.montant) for t in TarifCycle.objects.filter(etablissement=self.ecole)}
        self.assertEqual(tarifs, {"1er_cycle": ("par_eleve", 2500), "lycee": ("forfait", 200000)})
        journal = JournalAudit.objects.get(action="tarifs_cycle", acteur=self.proprietaire)
        self.assertEqual(journal.details, {"enregistres": 2, "supprimes": 0})

    def test_modifier_puis_retirer_un_tarif(self):
        TarifCycle.objects.create(etablissement=self.ecole, cycle="1er_cycle", mode="par_eleve", montant=2500)
        self.client.post(self.adresse, self.donnees(mode_1er_cycle="forfait", montant_1er_cycle="90000"))
        self.assertEqual(TarifCycle.objects.get(etablissement=self.ecole).mode, "forfait")
        self.client.post(self.adresse, self.donnees())  # tout vide : le tarif disparaît
        self.assertFalse(TarifCycle.objects.filter(etablissement=self.ecole).exists())

    def test_la_page_reprend_les_tarifs_existants(self):
        TarifCycle.objects.create(etablissement=self.ecole, cycle="1er_cycle", mode="forfait", montant=90000)
        page = self.client.get(self.adresse)
        self.assertContains(page, 'value="90000"')
        self.assertContains(page, '<option value="forfait" selected>')

    def test_un_mode_sans_montant_ou_un_montant_sans_mode_est_refuse(self):
        reponse = self.client.post(self.adresse, self.donnees(mode_1er_cycle="par_eleve", montant_lycee="1000"))
        self.assertEqual(reponse.status_code, 200)
        self.assertContains(reponse, "Indiquez le montant.")
        self.assertContains(reponse, "Choisissez le mode de facturation.")
        self.assertFalse(TarifCycle.objects.exists())

    def test_un_montant_nul_negatif_ou_trop_grand_est_refuse(self):
        for montant in ("0", "-5", "2000000001", "abc"):
            with self.subTest(montant=montant):
                reponse = self.client.post(self.adresse, self.donnees(mode_1er_cycle="forfait", montant_1er_cycle=montant))
                self.assertEqual(reponse.status_code, 200)
                self.assertFalse(TarifCycle.objects.exists())

    def test_un_mode_inconnu_est_refuse(self):
        reponse = self.client.post(self.adresse, self.donnees(mode_1er_cycle="gratuit", montant_1er_cycle="10"))
        self.assertEqual(reponse.status_code, 200)
        self.assertFalse(TarifCycle.objects.exists())

    def test_les_tarifs_d_une_ecole_ne_touchent_pas_ceux_d_une_autre(self):
        TarifCycle.objects.create(etablissement=self.ecole_b, cycle="1er_cycle", mode="forfait", montant=77000)
        self.client.post(self.adresse, self.donnees(mode_1er_cycle="par_eleve", montant_1er_cycle="2500"))
        self.assertEqual(TarifCycle.objects.get(etablissement=self.ecole_b).montant, 77000)
        self.client.post(self.adresse, self.donnees())  # vider l'école A ne vide pas l'école B
        self.assertEqual(TarifCycle.objects.get(etablissement=self.ecole_b).montant, 77000)

    def test_une_ecole_inconnue_repond_404(self):
        self.assertEqual(self.client.get(reverse("espace_plateforme:tarifs_cycle", args=[999999])).status_code, 404)

    def test_seul_le_proprietaire_modifie_les_tarifs(self):
        for compte in (self.fondateur, self.comptable, self.secretaire):
            self.client.force_login(compte)
            self.assertEqual(self.client.get(self.adresse).status_code, 302)
            self.assertEqual(self.client.post(self.adresse, self.donnees(mode_1er_cycle="forfait", montant_1er_cycle="1")).status_code, 302)
        self.client.logout()
        self.assertEqual(self.client.post(self.adresse, self.donnees(mode_1er_cycle="forfait", montant_1er_cycle="1")).status_code, 302)
        self.assertFalse(TarifCycle.objects.exists())


class MontantEtFactureTests(EducationBase):
    def setUp(self):
        self.proprietaire = creer_utilisateur_actif("montant-effectifs@example.com", Role.DEVELOPPEUR, is_superuser=True)
        self.client.force_login(self.proprietaire)
        TarifCycle.objects.create(etablissement=self.ecole, cycle="1er_cycle", mode="par_eleve", montant=2500)
        TarifCycle.objects.create(etablissement=self.ecole, cycle="lycee", mode="forfait", montant=200000)

    def test_la_page_montre_le_tarif_et_le_montant_annuel_estime(self):
        page = self.client.get(reverse("espace_plateforme:effectifs"))
        total = 4 * 2500 + 200000
        for attendu in ("Montant annuel estimé", "Par élève et par an", "Forfait par an", "Émettre la facture"):
            self.assertContains(page, attendu)
        contenu = page.content.decode().replace(" ", " ")
        self.assertIn("10 000 FCFA", contenu)
        self.assertIn("200 000 FCFA", contenu)
        ecole = next(e for e in page.context["ecoles"] if e.etablissement == self.ecole)
        self.assertEqual(ecole.montant_estime, total)
        self.assertEqual(page.context["total_montant"], total)

    def test_une_ecole_sans_montant_n_a_pas_de_bouton_facture(self):
        page = self.client.get(reverse("espace_plateforme:effectifs"))
        ecole_b = next(e for e in page.context["ecoles"] if e.etablissement == self.ecole_b)
        self.assertEqual(ecole_b.lien_facture, "")

    def test_le_bouton_facture_prerempli_le_formulaire_d_emission(self):
        page = self.client.get(reverse("espace_plateforme:effectifs"))
        ecole = next(e for e in page.context["ecoles"] if e.etablissement == self.ecole)
        formulaire = self.client.get(ecole.lien_facture)
        self.assertEqual(formulaire.status_code, 200)
        champs = formulaire.context["formulaire"].initial
        self.assertEqual(str(champs["etablissement"]), str(self.ecole.pk))
        self.assertEqual(str(champs["montant"]), str(4 * 2500 + 200000))
        self.assertEqual(champs["periode_debut"], self.annee.date_debut)
        self.assertEqual(champs["periode_fin"], self.annee.date_fin)
        self.assertIn("2026-2027", champs["libelle"])
        self.assertContains(formulaire, f'value="{4 * 2500 + 200000}"')  # le montant est dans la case, modifiable

    def test_le_prerempli_ignore_les_valeurs_qui_ne_sont_pas_valables(self):
        adresse = reverse("espace_plateforme:emettre_facture") + "?montant=abc&periode_debut=pas-une-date&libelle=" + "x" * 500
        formulaire = self.client.get(adresse).context["formulaire"]
        self.assertNotIn("montant", formulaire.initial)
        self.assertNotIn("periode_debut", formulaire.initial)
        self.assertEqual(len(formulaire.initial["libelle"]), 200)

    def test_l_export_contient_les_tarifs_et_le_montant(self):
        reponse = self.client.get(reverse("espace_plateforme:exporter_effectifs_xlsx"))
        lignes = [[c for c in ligne if c is not None] for ligne in load_workbook(io.BytesIO(reponse.content)).active.iter_rows(values_only=True)]
        premier_cycle = next(ligne for ligne in lignes if "École statistiques A" in ligne and "1er cycle" in ligne)
        self.assertIn("Par élève et par an", premier_cycle)
        self.assertIn(2500, premier_cycle)
        self.assertIn(10000, premier_cycle)
        total = next(ligne for ligne in lignes if ligne[:1] == ["Total"])  # pas l'en-tête, dont une colonne s'appelle aussi « Total »
        self.assertIn(4 * 2500 + 200000, total)

    def test_la_page_des_tarifs_et_des_montants_en_arabe_pur(self):
        from django.core.cache import cache

        from comptes.arabisation import invalider

        cache.clear()
        for ecole in (self.ecole, self.ecole_b):
            invalider(ecole.pk)
        client = client_en(self.proprietaire, "ar")
        saisis = {"Français", "École statistiques A", "École statistiques B"}
        for adresse in (reverse("espace_plateforme:effectifs"), reverse("espace_plateforme:tarifs_cycle", args=[self.ecole.pk])):
            with self.subTest(adresse=adresse):
                page = client.get(adresse)
                self.assertEqual(page.status_code, 200)
                self.assertEqual(restes_hors_textes_saisis(page.content.decode(), saisis), [])
        erreurs = client.post(reverse("espace_plateforme:tarifs_cycle", args=[self.ecole.pk]), {"mode_1er_cycle": "forfait", "montant_lycee": "5"})
        self.assertEqual(restes_hors_textes_saisis(erreurs.content.decode(), saisis), [])
