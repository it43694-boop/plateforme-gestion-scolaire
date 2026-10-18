"""Effectifs par cycle du propriétaire de la plateforme : élèves et personnel de chaque école, par cycle, pour facturer chaque cycle ; jamais un nom, jamais les parents."""
import io

from django.urls import reverse
from openpyxl import load_workbook

from comptes.models import JournalAudit, Utilisateur
from comptes.outils_test_arabe import client_en, restes_hors_textes_saisis
from comptes.roles import Role, StatutCompte
from espace_plateforme.effectifs import effectifs_de_l_ecole, effectifs_de_toutes_les_ecoles, total_par_cycle
from espace_plateforme.tests import creer_utilisateur_actif
from etablissement.models import Etablissement, PlanEtablissement
from scolarite.models import Cycle, Inscription
from statistiques.test_education import EducationBase, creer_compte


def par_cycle(effectifs):
    return {ligne.cycle: ligne for ligne in effectifs.lignes}


class CalculDesEffectifsTests(EducationBase):
    def test_les_eleves_les_enseignants_et_la_direction_de_chaque_cycle(self):
        lignes = par_cycle(effectifs_de_l_ecole(self.ecole))
        # élèves en cours de l'année active : 6A (e1, e2, e3) et 6B (e4) ; 9A (e5, e6) ; TSE (e7, e8) ; CAP-E (e9)
        self.assertEqual([lignes[c].eleves for c in (Cycle.PRESCOLAIRE, Cycle.PREMIER_CYCLE, Cycle.DEUXIEME_CYCLE, Cycle.LYCEE, Cycle.PROFESSIONNEL)], [0, 4, 2, 2, 1])
        # enseignants : t1 (6A), t2 (6B et 9A : compté dans les deux cycles), t3 (TSE) ; l'ancien n'enseigne que l'année précédente
        self.assertEqual([lignes[c].enseignants for c in (Cycle.PRESCOLAIRE, Cycle.PREMIER_CYCLE, Cycle.DEUXIEME_CYCLE, Cycle.LYCEE, Cycle.PROFESSIONNEL)], [0, 2, 1, 1, 0])
        # direction : le directeur du 1er cycle ; pour le lycée, le censeur
        self.assertEqual([lignes[c].direction for c in (Cycle.PREMIER_CYCLE, Cycle.DEUXIEME_CYCLE, Cycle.LYCEE)], [1, 0, 1])
        self.assertEqual(lignes[Cycle.PREMIER_CYCLE].total, 4 + 2 + 1)

    def test_le_reste_du_personnel_est_sans_cycle(self):
        effectifs = effectifs_de_l_ecole(self.ecole)
        personnel_actif = Utilisateur.objects.filter(etablissement=self.ecole, is_active=True, statut=StatutCompte.ACTIF).exclude(role__in=[Role.PARENT, Role.ELEVE])
        rattaches = {self.t1.pk, self.t2.pk, self.t3.pk, self.dir_1er.pk, self.censeur.pk}
        self.assertEqual(effectifs.sans_cycle, personnel_actif.exclude(pk__in=rattaches).count())
        self.assertGreaterEqual(effectifs.sans_cycle, 4)  # le fondateur, le comptable, la secrétaire et l'enseignant sans classe cette année
        self.assertEqual(effectifs.personnel, personnel_actif.count() + 1)  # t2, compté deux fois (une fois dans chacun de ses cycles)

    def test_les_parents_ne_sont_pas_comptes(self):
        avant = effectifs_de_l_ecole(self.ecole).total
        for numero in range(3):
            creer_compte(Role.PARENT, f"parent-en-plus-{numero}@stat.test", self.ecole, telephone=f"7099910{numero}")
        self.assertEqual(effectifs_de_l_ecole(self.ecole).total, avant)

    def test_un_compte_inactif_ou_en_attente_n_est_pas_compte(self):
        avant = effectifs_de_l_ecole(self.ecole)
        Utilisateur.objects.filter(pk=self.t1.pk).update(statut=StatutCompte.EN_ATTENTE_VALIDATION)
        Utilisateur.objects.filter(pk=self.censeur.pk).update(is_active=False)
        apres = effectifs_de_l_ecole(self.ecole)
        self.assertEqual(par_cycle(apres)[Cycle.PREMIER_CYCLE].enseignants, par_cycle(avant)[Cycle.PREMIER_CYCLE].enseignants - 1)
        self.assertEqual(par_cycle(apres)[Cycle.LYCEE].direction, 0)

    def test_seuls_les_eleves_inscrits_en_cours_de_l_annee_active_comptent(self):
        Inscription.objects.filter(pk=self.i["e4"].pk).update(statut=Inscription.Statut.ABANDON)
        self.assertEqual(par_cycle(effectifs_de_l_ecole(self.ecole))[Cycle.PREMIER_CYCLE].eleves, 3)
        # l'année précédente (e1 et e2 en 5A) ne compte pas
        self.assertEqual(Inscription.objects.filter(classe=self.classe_5a).count(), 2)

    def test_une_ecole_sans_annee_active_ne_compte_aucun_eleve_ni_enseignant(self):
        self.annee.est_active = False
        self.annee.save()
        effectifs = effectifs_de_l_ecole(self.ecole)
        self.assertIsNone(effectifs.annee)
        self.assertEqual(effectifs.eleves, 0)
        self.assertEqual(sum(ligne.enseignants for ligne in effectifs.lignes), 0)

    def test_les_cycles_hors_formule_sont_signales(self):
        self.ecole.plan = PlanEtablissement.PREMIER_CYCLE
        self.ecole.save()
        lignes = par_cycle(effectifs_de_l_ecole(self.ecole))
        self.assertTrue(lignes[Cycle.PREMIER_CYCLE].dans_la_formule)
        self.assertFalse(lignes[Cycle.PREMIER_CYCLE].hors_formule)
        for cycle in (Cycle.DEUXIEME_CYCLE, Cycle.LYCEE, Cycle.PROFESSIONNEL):
            self.assertFalse(lignes[cycle].dans_la_formule)
            self.assertTrue(lignes[cycle].hors_formule, cycle)
        self.assertTrue(effectifs_de_l_ecole(self.ecole).a_regulariser)
        self.ecole.inclut_professionnel = True  # une option de la formule : le professionnel n'est plus « hors formule »
        self.ecole.save()
        self.assertFalse(par_cycle(effectifs_de_l_ecole(self.ecole))[Cycle.PROFESSIONNEL].hors_formule)

    def test_les_lignes_utiles_masquent_les_cycles_vides_hors_formule(self):
        for ecole in (self.ecole, self.ecole_b):
            ecole.plan = PlanEtablissement.PREMIER_CYCLE
            ecole.save()
        libelles = [ligne.libelle for ligne in effectifs_de_l_ecole(self.ecole_b).lignes_utiles]
        self.assertEqual(libelles, ["1er cycle"])  # l'école B n'a personne : seul le cycle de sa formule est affiché
        # l'école A a des comptes dans d'autres cycles que sa formule : ils ressortent
        self.assertEqual(len(effectifs_de_l_ecole(self.ecole).lignes_utiles), 4)

    def test_chaque_ecole_garde_ses_propres_chiffres(self):
        creer_compte(Role.ELEVE, "eleve-b-1@stat.test", self.ecole_b)
        self.assertEqual(par_cycle(effectifs_de_l_ecole(self.ecole_b))[Cycle.PREMIER_CYCLE].eleves, 0)  # un compte élève sans inscription ne compte pas
        Inscription.objects.create(eleve=Utilisateur.objects.get(email="eleve-b-1@stat.test"), classe=self.classe_b)
        self.assertEqual(par_cycle(effectifs_de_l_ecole(self.ecole_b))[Cycle.PREMIER_CYCLE].eleves, 1)
        self.assertEqual(par_cycle(effectifs_de_l_ecole(self.ecole))[Cycle.PREMIER_CYCLE].eleves, 4)

    def test_les_totaux_de_la_plateforme(self):
        ecoles = effectifs_de_toutes_les_ecoles(Etablissement.objects.order_by("nom"))
        totaux = {ligne.cycle: ligne for ligne in total_par_cycle(ecoles)}
        self.assertEqual(totaux[Cycle.PREMIER_CYCLE].eleves, sum(par_cycle(e)[Cycle.PREMIER_CYCLE].eleves for e in ecoles))
        self.assertEqual(totaux[Cycle.PREMIER_CYCLE].eleves, 4)


class PageDesEffectifsTests(EducationBase):
    def setUp(self):
        self.proprietaire = creer_utilisateur_actif("proprietaire-effectifs@example.com", Role.DEVELOPPEUR, is_superuser=True)
        self.client.force_login(self.proprietaire)
        self.adresse = reverse("espace_plateforme:effectifs")

    def test_le_proprietaire_voit_les_ecoles_et_les_chiffres(self):
        page = self.client.get(self.adresse)
        self.assertEqual(page.status_code, 200)
        for attendu in ("Effectifs par cycle", "École statistiques A", "École statistiques B", "1er cycle", "2ème cycle", "Lycée", "Sans cycle (administration et services)", "Exporter en Excel"):
            self.assertContains(page, attendu)
        self.assertEqual(page.context["total_general"], sum(e.total for e in page.context["ecoles"]))
        self.assertEqual({e.etablissement.nom for e in page.context["ecoles"]}, {"École statistiques A", "École statistiques B"})

    def test_la_page_n_affiche_aucun_nom_ni_adresse_de_personne(self):
        Utilisateur.objects.filter(pk=self.e1.pk).update(prenom="Prénom-Confidentiel", nom="Nom-Confidentiel")
        Utilisateur.objects.filter(pk=self.t1.pk).update(prenom="Enseignant-Confidentiel")
        contenu = self.client.get(self.adresse).content.decode()
        for interdit in ("Confidentiel", self.e1.email, self.t1.email, self.e1.matricule, self.parent.email, self.fondateur.email):
            self.assertNotIn(interdit, contenu)

    def test_une_ecole_hors_formule_est_signalee_et_une_ecole_suspendue_aussi(self):
        self.ecole.plan = PlanEtablissement.PREMIER_CYCLE
        self.ecole.save()
        Etablissement.objects.filter(pk=self.ecole_b.pk).update(actif=False)
        page = self.client.get(self.adresse)
        self.assertContains(page, "Hors formule")
        self.assertContains(page, "Des comptes hors formule")
        self.assertContains(page, "Suspendue")

    def test_une_ecole_sans_annee_active_l_explique(self):
        self.annee_b.est_active = False
        self.annee_b.save()
        self.assertContains(self.client.get(self.adresse), "Aucune année scolaire active")

    def test_le_lien_est_dans_le_menu_et_sur_la_liste_des_ecoles(self):
        self.assertContains(self.client.get(reverse("espace_plateforme:liste_etablissements")), self.adresse)
        self.assertContains(self.client.get(self.adresse), f'href="{self.adresse}"')  # le menu du propriétaire

    def test_seul_le_proprietaire_y_acces(self):
        for compte in (self.fondateur, self.secretaire, self.comptable, self.parent):
            self.client.force_login(compte)
            for adresse in (self.adresse, reverse("espace_plateforme:exporter_effectifs_xlsx")):
                with self.subTest(compte=compte.email, adresse=adresse):
                    self.assertEqual(self.client.get(adresse).status_code, 302)  # redirigé vers la connexion
        developpeur = creer_utilisateur_actif("dev-effectifs@example.com", Role.DEVELOPPEUR, etablissement=self.ecole)
        self.client.force_login(developpeur)
        self.assertEqual(self.client.get(self.adresse).status_code, 302)  # un développeur d'école n'est pas le propriétaire
        self.client.logout()
        self.assertEqual(self.client.get(self.adresse).status_code, 302)


class ExportDesEffectifsTests(EducationBase):
    def setUp(self):
        self.proprietaire = creer_utilisateur_actif("export-effectifs@example.com", Role.DEVELOPPEUR, is_superuser=True)
        self.client.force_login(self.proprietaire)

    def feuille(self, client=None):
        reponse = (client or self.client).get(reverse("espace_plateforme:exporter_effectifs_xlsx"))
        self.assertEqual(reponse.status_code, 200)
        self.assertIn("spreadsheetml", reponse["Content-Type"])
        return list((load_workbook(io.BytesIO(reponse.content)).active).iter_rows(values_only=True))

    def test_le_fichier_contient_une_ligne_par_ecole_et_par_cycle(self):
        lignes = self.feuille()
        textes = [[c for c in ligne if c is not None] for ligne in lignes]
        a_plat = [ligne for ligne in textes if "École statistiques A" in ligne]
        cycles = [ligne for ligne in a_plat if "1er cycle" in ligne]
        self.assertEqual(len(cycles), 1)
        self.assertIn(4, cycles[0])  # les quatre élèves du 1er cycle
        self.assertTrue(any("Sans cycle (administration et services)" in ligne for ligne in a_plat))
        self.assertTrue(any("Total" in ligne for ligne in textes))

    def test_le_fichier_ne_contient_aucun_nom_de_personne(self):
        Utilisateur.objects.filter(pk=self.e1.pk).update(prenom="Prénom-Confidentiel")
        reponse = self.client.get(reverse("espace_plateforme:exporter_effectifs_xlsx"))
        contenu = " ".join(str(c) for ligne in load_workbook(io.BytesIO(reponse.content)).active.iter_rows(values_only=True) for c in ligne if c is not None)
        for interdit in ("Confidentiel", self.e1.email, self.t1.email, self.e1.matricule):
            self.assertNotIn(interdit, contenu)

    def test_l_export_est_journalise(self):
        self.feuille()
        self.assertTrue(JournalAudit.objects.filter(action="export_effectifs_par_cycle", acteur=self.proprietaire).exists())


class EffectifsEnArabeTests(EducationBase):
    def test_la_page_et_l_export_sont_en_arabe_pur(self):
        from django.core.cache import cache

        from comptes.arabisation import invalider

        cache.clear()
        for ecole in (self.ecole, self.ecole_b):
            invalider(ecole.pk)
        proprietaire = creer_utilisateur_actif("effectifs-ar@example.com", Role.DEVELOPPEUR, is_superuser=True)
        client = client_en(proprietaire, "ar")
        page = client.get(reverse("espace_plateforme:effectifs"))
        self.assertEqual(page.status_code, 200)
        self.assertEqual(restes_hors_textes_saisis(page.content.decode(), {"Français", "École statistiques A", "École statistiques B"}), [])
        self.assertIn("المرحلة الأولى", page.content.decode())
        export = client.get(reverse("espace_plateforme:exporter_effectifs_xlsx"))
        self.assertEqual(export.status_code, 200)
