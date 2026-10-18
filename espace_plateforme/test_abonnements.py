"""
Facturation de l'abonnement des écoles (propriétaire de la plateforme) : numérotation, chevauchements, règlement,
annulation, situation d'un abonnement, pages, PDF et export Excel.
"""
import datetime
import io

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from comptes.models import JournalAudit
from comptes.roles import Role
from espace_plateforme.tests import creer_utilisateur_actif
from etablissement.abonnements import (
    annuler_facture, emettre_facture, enregistrer_reglement, situation_abonnement, synthese_abonnements,
)
from etablissement.models import Etablissement, FactureAbonnement, PlanEtablissement

JOUR = datetime.timedelta(days=1)
STATUT = FactureAbonnement.Statut


def aujourdhui():
    return timezone.localdate()


class AbonnementsBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.alpha = Etablissement.objects.create(nom="École Alpha", plan=PlanEtablissement.TOUS_CYCLES, code_etablissement="ML-A")
        cls.beta = Etablissement.objects.create(nom="École Beta", plan=PlanEtablissement.PREMIER_CYCLE)
        cls.proprietaire = creer_utilisateur_actif("owner-abo@example.com", Role.DEVELOPPEUR, is_superuser=True)
        cls.dev_alpha = creer_utilisateur_actif("dev-abo@example.com", Role.DEVELOPPEUR, cls.alpha)
        cls.fondateur_alpha = creer_utilisateur_actif("fond-abo@example.com", Role.FONDATEUR, cls.alpha)

    def facturer(self, ecole=None, debut=None, fin=None, montant=120000, **extra):
        debut = debut or datetime.date(aujourdhui().year, 1, 1)
        fin = fin or datetime.date(aujourdhui().year, 12, 31)
        return emettre_facture(
            etablissement=ecole or self.alpha, periode_debut=debut, periode_fin=fin, montant=montant,
            cree_par=self.proprietaire, **extra,
        )


class EmissionTests(AbonnementsBase):
    def test_numerotation_continue_par_annee(self):
        a = self.facturer(self.alpha)
        b = self.facturer(self.beta)
        self.assertEqual((a.numero, b.numero), (f"ABO-{aujourdhui().year}-0001", f"ABO-{aujourdhui().year}-0002"))
        ancienne = self.facturer(self.alpha, datetime.date(2024, 1, 1), datetime.date(2024, 12, 31),
                                 date_emission=datetime.date(2024, 1, 5))
        self.assertEqual(ancienne.numero, "ABO-2024-0001")

    def test_valeurs_par_defaut_et_formule_conservee(self):
        facture = self.facturer(self.alpha)
        self.assertEqual(facture.statut, STATUT.EMISE)
        self.assertEqual(facture.date_emission, aujourdhui())
        self.assertEqual(facture.date_echeance, aujourdhui() + 15 * JOUR)
        self.assertEqual(facture.plan, PlanEtablissement.TOUS_CYCLES)
        self.assertIn("Tous les cycles", facture.libelle)
        self.assertIn("01/01/", facture.libelle)
        self.alpha.plan = PlanEtablissement.LYCEE_SEUL
        self.alpha.save()
        facture.refresh_from_db()
        self.assertEqual(facture.plan, PlanEtablissement.TOUS_CYCLES)  # la formule facturée ne change pas après coup

    def test_libelle_personnalise(self):
        self.assertEqual(self.facturer(libelle="  Abonnement spécial  ").libelle, "Abonnement spécial")

    def test_validations(self):
        debut = datetime.date(2030, 1, 1)
        for attendu, valeurs in [
            ("postérieure", dict(debut=debut, fin=debut)),
            ("postérieure", dict(debut=debut, fin=debut - JOUR)),
            ("strictement positif", dict(debut=debut, fin=debut + 30 * JOUR, montant=0)),
            ("précéder", dict(debut=debut, fin=debut + 30 * JOUR, date_emission=aujourdhui(), date_echeance=aujourdhui() - JOUR)),
        ]:
            with self.subTest(attendu=attendu, valeurs=valeurs), self.assertRaisesMessage(ValidationError, attendu):
                self.facturer(**valeurs)
        self.assertEqual(FactureAbonnement.objects.count(), 0)

    def test_une_periode_ne_se_facture_qu_une_fois(self):
        premiere = self.facturer(self.alpha, datetime.date(2030, 1, 1), datetime.date(2030, 12, 31))
        with self.assertRaisesMessage(ValidationError, premiere.numero):
            self.facturer(self.alpha, datetime.date(2030, 6, 1), datetime.date(2031, 5, 31))
        with self.assertRaises(ValidationError):  # même période exactement
            self.facturer(self.alpha, datetime.date(2030, 1, 1), datetime.date(2030, 12, 31))
        self.facturer(self.alpha, datetime.date(2031, 1, 1), datetime.date(2031, 12, 31))  # période suivante : acceptée
        self.facturer(self.beta, datetime.date(2030, 1, 1), datetime.date(2030, 12, 31))   # autre école : acceptée

    def test_une_facture_annulee_ne_bloque_plus_sa_periode(self):
        premiere = self.facturer(self.alpha, datetime.date(2030, 1, 1), datetime.date(2030, 12, 31))
        annuler_facture(facture=premiere, motif="Montant erroné")
        self.facturer(self.alpha, datetime.date(2030, 1, 1), datetime.date(2030, 12, 31), montant=100000)

    def test_contraintes_de_base_de_donnees(self):
        facture = self.facturer()
        with self.assertRaises(IntegrityError), transaction.atomic():
            FactureAbonnement.objects.filter(pk=facture.pk).update(statut=STATUT.PAYEE, date_paiement=None)
        with self.assertRaises(IntegrityError), transaction.atomic():
            FactureAbonnement.objects.filter(pk=facture.pk).update(montant=0)
        with self.assertRaises(IntegrityError), transaction.atomic():
            FactureAbonnement.objects.filter(pk=facture.pk).update(periode_fin=facture.periode_debut)


class ReglementEtAnnulationTests(AbonnementsBase):
    def regler(self, facture, **extra):
        valeurs = dict(date_paiement=aujourdhui(), mode_reglement="mobile_money", reference="OM-123", acteur=self.proprietaire)
        valeurs.update(extra)
        return enregistrer_reglement(facture=facture, **valeurs)

    def test_reglement(self):
        facture = self.regler(self.facturer())
        self.assertEqual((facture.statut, facture.date_paiement, facture.mode_reglement, facture.reference_reglement),
                         (STATUT.PAYEE, aujourdhui(), "mobile_money", "OM-123"))
        facture.refresh_from_db()
        self.assertEqual(facture.statut, STATUT.PAYEE)

    def test_reglements_refuses(self):
        facture = self.facturer()
        for attendu, valeurs in [
            ("futur", dict(date_paiement=aujourdhui() + JOUR)),
            ("précéder", dict(date_paiement=facture.date_emission - JOUR)),
            ("invalide", dict(mode_reglement="troc")),
        ]:
            with self.subTest(attendu=attendu), self.assertRaisesMessage(ValidationError, attendu):
                self.regler(facture, **valeurs)
        facture.refresh_from_db()
        self.assertEqual(facture.statut, STATUT.EMISE)

    def test_pas_deux_reglements_ni_reglement_d_une_facture_annulee(self):
        payee = self.regler(self.facturer())
        with self.assertRaisesMessage(ValidationError, "déjà payée"):
            self.regler(payee)
        annulee = annuler_facture(facture=self.facturer(self.beta), motif="Doublon")
        with self.assertRaisesMessage(ValidationError, "annulée"):
            self.regler(annulee)

    def test_annulation(self):
        facture = annuler_facture(facture=self.facturer(), motif="  Erreur de période ")
        self.assertEqual((facture.statut, facture.motif_annulation), (STATUT.ANNULEE, "Erreur de période"))

    def test_annulations_refusees(self):
        with self.assertRaisesMessage(ValidationError, "motif"):
            annuler_facture(facture=self.facturer(), motif="   ")
        payee = self.regler(self.facturer(self.beta))
        with self.assertRaisesMessage(ValidationError, "payée"):
            annuler_facture(facture=payee, motif="Trop tard")
        annulee = annuler_facture(facture=self.facturer(self.beta, datetime.date(2035, 1, 1), datetime.date(2035, 12, 31)), motif="x")
        with self.assertRaisesMessage(ValidationError, "déjà annulée"):
            annuler_facture(facture=annulee, motif="encore")


class SituationTests(AbonnementsBase):
    def test_aucune_facture(self):
        self.assertEqual(situation_abonnement(self.alpha).code, "aucune")

    def test_facture_a_payer_dans_les_delais(self):
        self.facturer(self.alpha)
        situation = situation_abonnement(self.alpha)
        self.assertEqual((situation.code, situation.nombre_en_retard, situation.montant_a_payer), ("en_attente", 0, 120000))

    def test_en_retard_apres_l_echeance(self):
        self.facturer(self.alpha, date_emission=aujourdhui() - 40 * JOUR, date_echeance=aujourdhui() - 10 * JOUR)
        situation = situation_abonnement(self.alpha)
        self.assertEqual((situation.code, situation.nombre_en_retard, situation.montant_en_retard), ("en_retard", 1, 120000))
        self.assertEqual(situation.libelle, "En retard de paiement")

    def test_a_jour_quand_une_facture_payee_couvre_aujourd_hui(self):
        facture = self.facturer(self.alpha)
        enregistrer_reglement(facture=facture, date_paiement=aujourdhui(), mode_reglement="virement")
        situation = situation_abonnement(self.alpha)
        self.assertEqual(situation.code, "a_jour")
        self.assertEqual(situation.couvert_jusqu_au, datetime.date(aujourdhui().year, 12, 31))

    def test_expire_quand_la_periode_payee_est_passee(self):
        facture = self.facturer(self.alpha, datetime.date(2020, 1, 1), datetime.date(2020, 12, 31), date_emission=datetime.date(2020, 1, 2))
        enregistrer_reglement(facture=facture, date_paiement=datetime.date(2020, 1, 10), mode_reglement="especes")
        self.assertEqual(situation_abonnement(self.alpha).code, "expire")

    def test_le_retard_prime_sur_un_abonnement_encore_couvert(self):
        payee = self.facturer(self.alpha)
        enregistrer_reglement(facture=payee, date_paiement=aujourdhui(), mode_reglement="virement")
        self.facturer(self.alpha, datetime.date(aujourdhui().year + 1, 1, 1), datetime.date(aujourdhui().year + 1, 12, 31),
                      date_emission=aujourdhui() - 30 * JOUR, date_echeance=aujourdhui() - JOUR)
        self.assertEqual(situation_abonnement(self.alpha).code, "en_retard")

    def test_une_facture_annulee_ne_compte_pas(self):
        annuler_facture(facture=self.facturer(self.alpha), motif="Erreur")
        self.assertEqual(situation_abonnement(self.alpha).code, "aucune")

    def test_synthese_et_totaux(self):
        payee = self.facturer(self.alpha, montant=100000)
        enregistrer_reglement(facture=payee, date_paiement=aujourdhui(), mode_reglement="virement")
        self.facturer(self.beta, montant=50000)                                       # dans les délais
        self.facturer(self.beta, datetime.date(aujourdhui().year - 1, 1, 1), datetime.date(aujourdhui().year - 1, 12, 31),
                      montant=30000, date_emission=aujourdhui() - 60 * JOUR, date_echeance=aujourdhui() - 30 * JOUR)  # en retard
        annuler_facture(facture=self.facturer(self.alpha, datetime.date(2040, 1, 1), datetime.date(2040, 12, 31), montant=999), motif="x")
        lignes, totaux = synthese_abonnements()
        par_ecole = {etablissement.nom: situation.code for etablissement, situation in lignes}
        self.assertEqual(par_ecole["École Alpha"], "a_jour")
        self.assertEqual(par_ecole["École Beta"], "en_retard")
        self.assertEqual((totaux["encaisse_annee"], totaux["a_encaisser"], totaux["en_retard"]), (100000, 50000, 30000))
        self.assertEqual((totaux["ecoles_a_jour"], totaux["ecoles_en_retard"]), (1, 1))


class PagesTests(AbonnementsBase):
    ROUTES_SANS_PARAMETRE = ("abonnements", "liste_factures", "emettre_facture", "exporter_factures_xlsx")

    def test_reservees_au_proprietaire_de_la_plateforme(self):
        facture = self.facturer()
        adresses = [reverse(f"espace_plateforme:{nom}") for nom in self.ROUTES_SANS_PARAMETRE] + [
            reverse("espace_plateforme:payer_facture", args=[facture.id]),
            reverse("espace_plateforme:facture_pdf", args=[facture.id]),
        ]
        for compte in (self.dev_alpha, self.fondateur_alpha):
            self.client.force_login(compte)
            for adresse in adresses:
                with self.subTest(role=compte.role, adresse=adresse):
                    self.assertEqual(self.client.get(adresse).status_code, 302)  # renvoyé vers la connexion
            self.assertEqual(self.client.post(reverse("espace_plateforme:annuler_facture", args=[facture.id]),
                                              {"motif": "x"}).status_code, 302)
        self.client.logout()
        for adresse in adresses:
            self.assertEqual(self.client.get(adresse).status_code, 302)
        facture.refresh_from_db()
        self.assertEqual(facture.statut, STATUT.EMISE, "un compte d'école ne doit pas pouvoir annuler une facture")
        self.client.force_login(self.proprietaire)
        for adresse in adresses:
            with self.subTest(role="proprietaire", adresse=adresse):
                self.assertEqual(self.client.get(adresse).status_code, 200)

    def test_vue_d_ensemble(self):
        self.facturer(self.alpha, date_emission=aujourdhui() - 40 * JOUR, date_echeance=aujourdhui() - 10 * JOUR)
        self.client.force_login(self.proprietaire)
        page = self.client.get(reverse("espace_plateforme:abonnements"))
        for texte in ("École Alpha", "École Beta", "En retard de paiement", "Aucune facture", "1 facture en retard (120000 FCFA)",
                      reverse("espace_plateforme:abonnements"), "n'est pas désactivée automatiquement"):
            self.assertContains(page, texte)

    def test_liste_et_filtres(self):
        retard = self.facturer(self.alpha, datetime.date(2020, 1, 1), datetime.date(2020, 12, 31),
                               date_emission=aujourdhui() - 60 * JOUR, date_echeance=aujourdhui() - 30 * JOUR)
        payee = self.facturer(self.beta)
        enregistrer_reglement(facture=payee, date_paiement=aujourdhui(), mode_reglement="especes")
        a_payer = self.facturer(self.alpha, datetime.date(2031, 1, 1), datetime.date(2031, 12, 31))
        self.client.force_login(self.proprietaire)
        url = reverse("espace_plateforme:liste_factures")

        def numeros(**parametres):
            return {f.numero for f in self.client.get(url, parametres).context["factures"]}

        self.assertEqual(numeros(), {retard.numero, payee.numero, a_payer.numero})
        self.assertEqual(numeros(statut="en_retard"), {retard.numero})
        self.assertEqual(numeros(statut="payee"), {payee.numero})
        self.assertEqual(numeros(statut="emise"), {retard.numero, a_payer.numero})
        self.assertEqual(numeros(etablissement=self.beta.id), {payee.numero})
        self.assertEqual(numeros(etablissement=self.alpha.id, statut="emise"), {retard.numero, a_payer.numero})
        self.assertEqual(self.client.get(url, {"etablissement": "abc"}).status_code, 404)
        self.assertContains(self.client.get(url), "En retard")

    def test_emission_depuis_la_page(self):
        self.client.force_login(self.proprietaire)
        adresse = reverse("espace_plateforme:emettre_facture")
        self.assertEqual(self.client.get(adresse, {"etablissement": self.beta.id}).context["formulaire"].initial["etablissement"], str(self.beta.id))
        reponse = self.client.post(adresse, {
            "etablissement": self.alpha.id, "periode_debut": "2031-01-01", "periode_fin": "2031-12-31", "montant": 150000,
            "date_emission": aujourdhui().isoformat(), "date_echeance": "", "libelle": "",
        })
        self.assertRedirects(reponse, reverse("espace_plateforme:liste_factures"))
        facture = FactureAbonnement.objects.get()
        self.assertEqual((facture.etablissement, facture.montant, facture.cree_par), (self.alpha, 150000, self.proprietaire))
        entree = JournalAudit.objects.get(action="emission_facture_abonnement")
        self.assertEqual((entree.cible, entree.details["montant"]), (facture.numero, 150000))

    def test_emission_refusee_en_cas_de_chevauchement(self):
        premiere = self.facturer(self.alpha, datetime.date(2031, 1, 1), datetime.date(2031, 12, 31))
        self.client.force_login(self.proprietaire)
        reponse = self.client.post(reverse("espace_plateforme:emettre_facture"), {
            "etablissement": self.alpha.id, "periode_debut": "2031-06-01", "periode_fin": "2032-05-31", "montant": 1000,
            "date_emission": aujourdhui().isoformat(),
        })
        self.assertEqual(reponse.status_code, 200)
        self.assertContains(reponse, premiere.numero)
        self.assertEqual(FactureAbonnement.objects.count(), 1)

    def test_emission_valeurs_invalides(self):
        self.client.force_login(self.proprietaire)
        reponse = self.client.post(reverse("espace_plateforme:emettre_facture"), {
            "etablissement": self.alpha.id, "periode_debut": "pas-une-date", "periode_fin": "2031-12-31", "montant": -5,
            "date_emission": aujourdhui().isoformat(),
        })
        self.assertEqual(reponse.status_code, 200)
        self.assertFalse(FactureAbonnement.objects.exists())

    def test_reglement_depuis_la_page(self):
        facture = self.facturer()
        Etablissement.objects.filter(pk=self.alpha.pk).update(actif=False)
        self.client.force_login(self.proprietaire)
        adresse = reverse("espace_plateforme:payer_facture", args=[facture.id])
        self.assertContains(self.client.get(adresse), facture.numero)
        reponse = self.client.post(adresse, {"date_paiement": aujourdhui().isoformat(), "mode_reglement": "mobile_money", "reference": "OM-9"},
                                   follow=True)
        facture.refresh_from_db()
        self.assertEqual((facture.statut, facture.reference_reglement), (STATUT.PAYEE, "OM-9"))
        self.assertContains(reponse, "est toujours désactivé")  # rappel : la réactivation reste une décision du propriétaire
        self.assertTrue(JournalAudit.objects.filter(action="reglement_facture_abonnement", cible=facture.numero).exists())
        # une facture déjà payée n'ouvre plus le formulaire
        self.assertRedirects(self.client.get(adresse), reverse("espace_plateforme:liste_factures"))

    def test_reglement_avec_mode_invalide(self):
        facture = self.facturer()
        self.client.force_login(self.proprietaire)
        reponse = self.client.post(reverse("espace_plateforme:payer_facture", args=[facture.id]),
                                   {"date_paiement": aujourdhui().isoformat(), "mode_reglement": "troc"})
        self.assertEqual(reponse.status_code, 200)
        facture.refresh_from_db()
        self.assertEqual(facture.statut, STATUT.EMISE)

    def test_annulation_depuis_la_page(self):
        facture = self.facturer()
        self.client.force_login(self.proprietaire)
        adresse = reverse("espace_plateforme:annuler_facture", args=[facture.id])
        self.client.post(adresse, {"motif": ""})
        facture.refresh_from_db()
        self.assertEqual(facture.statut, STATUT.EMISE)
        self.client.post(adresse, {"motif": "Facture en double"})
        facture.refresh_from_db()
        self.assertEqual((facture.statut, facture.motif_annulation), (STATUT.ANNULEE, "Facture en double"))
        self.assertEqual(self.client.get(adresse).status_code, 405)  # l'annulation ne se fait jamais par un simple lien
        self.assertTrue(JournalAudit.objects.filter(action="annulation_facture_abonnement").exists())

    def test_facture_inexistante_introuvable(self):
        self.client.force_login(self.proprietaire)
        for nom in ("payer_facture", "facture_pdf"):
            self.assertEqual(self.client.get(reverse(f"espace_plateforme:{nom}", args=[999999])).status_code, 404)
        self.assertEqual(self.client.post(reverse("espace_plateforme:annuler_facture", args=[999999])).status_code, 404)

    @override_settings(EDITEUR_NOM="Yelenova Test SARL", ABONNEMENT_INSTRUCTIONS_PAIEMENT="Orange Money : 70 00 00 00")
    def test_facture_pdf(self):
        facture = self.facturer()
        self.client.force_login(self.proprietaire)
        for etat in ("emise", "payee", "annulee"):
            if etat == "payee":
                enregistrer_reglement(facture=facture, date_paiement=aujourdhui(), mode_reglement="virement", reference="VIR-1")
            elif etat == "annulee":
                facture = self.facturer(self.beta)
                annuler_facture(facture=facture, motif="Erreur")
            with self.subTest(etat=etat):
                reponse = self.client.get(reverse("espace_plateforme:facture_pdf", args=[facture.id]))
                self.assertEqual(reponse.status_code, 200)
                self.assertEqual(reponse["Content-Type"], "application/pdf")
                self.assertTrue(reponse.content.startswith(b"%PDF"))
                self.assertIn(f'facture_{facture.numero}.pdf', reponse["Content-Disposition"])

    def test_export_excel(self):
        payee = self.facturer(self.alpha, montant=100000)
        enregistrer_reglement(facture=payee, date_paiement=aujourdhui(), mode_reglement="mobile_money", reference="OM-5")
        retard = self.facturer(self.beta, datetime.date(2020, 1, 1), datetime.date(2020, 12, 31), montant=40000,
                               date_emission=aujourdhui() - 60 * JOUR, date_echeance=aujourdhui() - 30 * JOUR)
        annulee = self.facturer(self.beta, datetime.date(2035, 1, 1), datetime.date(2035, 12, 31), montant=999999)
        annuler_facture(facture=annulee, motif="x")
        self.client.force_login(self.proprietaire)
        reponse = self.client.get(reverse("espace_plateforme:exporter_factures_xlsx"))
        self.assertEqual(reponse.status_code, 200)
        self.assertIn("spreadsheetml.sheet", reponse["Content-Type"])
        feuille = load_workbook(io.BytesIO(reponse.content))["Factures"]
        lignes = [[c.value for c in rangee] for rangee in feuille.iter_rows()]
        entete = next(ligne for ligne in lignes if ligne[0] == "Numéro")
        self.assertEqual(entete[:6], ["Numéro", "Établissement", "Libellé", "Début", "Fin", "Montant (FCFA)"])
        donnees = {ligne[0]: ligne for ligne in lignes[lignes.index(entete) + 1:] if ligne[0] not in (None, "Total")}
        self.assertEqual(set(donnees), {payee.numero, retard.numero, annulee.numero})
        self.assertEqual((donnees[payee.numero][8], donnees[payee.numero][10], donnees[payee.numero][11]), ("Payée", "Mobile Money", "OM-5"))
        self.assertEqual(donnees[retard.numero][8], "En retard")
        self.assertEqual(donnees[annulee.numero][8], "Annulée")
        total = next(ligne for ligne in lignes if ligne[0] == "Total")
        self.assertEqual(total[5], 140000)  # sans la facture annulée
        self.assertTrue(JournalAudit.objects.filter(action="export_factures_abonnement").exists())
        # le filtre s'applique aussi à l'export
        filtre = self.client.get(reverse("espace_plateforme:exporter_factures_xlsx"), {"statut": "payee"})
        lignes = [[c.value for c in rangee] for rangee in load_workbook(io.BytesIO(filtre.content))["Factures"].iter_rows()]
        self.assertEqual({ligne[0] for ligne in lignes if ligne[0] and ligne[0].startswith("ABO-")}, {payee.numero})

    def test_lien_dans_la_navigation_du_proprietaire(self):
        self.client.force_login(self.proprietaire)
        self.assertContains(self.client.get(reverse("espace_plateforme:liste_etablissements")), reverse("espace_plateforme:abonnements"))
