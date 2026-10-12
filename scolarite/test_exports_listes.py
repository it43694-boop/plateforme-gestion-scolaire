"""
Exports Excel d'une classe : liste des élèves, relevé de notes, absences.

Jeu de données (classe 6A de EducationBase : e1, e2, e3 ; matières Maths coef 1 et Français coef 2) :
  1er trimestre   Maths (classe/40 compo)      Français
    e1            12 / 24 -> 12.00             15 / 30 -> 15.00      moyenne (12*1 + 15*2) / 3 = 14.00
    e2            10 / 20 -> 10.00             8 / 16 -> 8.00        moyenne (10*1 + 8*2) / 3  =  8.67
    e3            14 / 28 -> 14.00             -                     moyenne 14.00 (ex æquo avec e1 : rang 1, 1, 3)
  2ème trimestre  e1 Maths 16 (classe seule) ; e2 Maths 6
  Moyennes annuelles (toutes les notes) : e1 (12 + 30 + 16) / 4 = 14.5 ; e2 (10 + 16 + 6) / 4 = 8.0 ; e3 14.0
  Absences        e1 : 5 oct non justifiée, 6 oct justifiée (« Paludisme ») ; e2 : 7 oct non justifiée
"""
import datetime
import io
import zipfile

from django.urls import reverse
from openpyxl import load_workbook

from comptes.models import JournalAudit, Utilisateur
from comptes.roles import Role
from pedagogie.models import Absence, Trimestre, saisir_note
from pedagogie.views import _calculer_bulletin, _calculer_classement
from scolarite.models import Affectation, Inscription
from statistiques.test_education import EducationBase, creer_compte

D = datetime.date


def feuilles(contenu):
    """{nom de feuille: lignes (sans cellules vides de fin)} d'un classeur .xlsx."""
    classeur = load_workbook(io.BytesIO(contenu))
    resultat = {}
    for feuille in classeur.worksheets:
        lignes = []
        for rangee in feuille.iter_rows(values_only=True):
            valeurs = list(rangee)
            while valeurs and valeurs[-1] is None:
                valeurs.pop()
            lignes.append(valeurs)
        resultat[feuille.title] = lignes
    return resultat


def donnees(lignes, premiere_colonne):
    """Lignes d'un tableau : celles qui commencent après la ligne d'en-têtes (dont la première cellule est `premiere_colonne`)."""
    debut = next(i for i, ligne in enumerate(lignes) if ligne and ligne[0] == premiere_colonne)
    return lignes[debut + 1:]


def texte_du_classeur(contenu):
    return " ".join(str(v) for lignes in feuilles(contenu).values() for ligne in lignes for v in ligne if v is not None)


class ExportsBase(EducationBase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.t_fr = creer_compte(Role.ENSEIGNANT, "tfr@stat.test", cls.ecole)
        cls.maths = Affectation.objects.get(enseignant=cls.t1, classe=cls.c6a)
        cls.francais = Affectation.objects.create(enseignant=cls.t_fr, classe=cls.c6a, matiere="Français", coefficient=2)
        T1, T2 = Trimestre.T1, Trimestre.T2
        for eleve, affectation, trimestre, classe_, compo, prof in [
            (cls.e1, cls.maths, T1, 12, 24, cls.t1), (cls.e1, cls.francais, T1, 15, 30, cls.t_fr),
            (cls.e2, cls.maths, T1, 10, 20, cls.t1), (cls.e2, cls.francais, T1, 8, 16, cls.t_fr),
            (cls.e3, cls.maths, T1, 14, 28, cls.t1),
            (cls.e1, cls.maths, T2, 16, None, cls.t1), (cls.e2, cls.maths, T2, 6, None, cls.t1),
        ]:
            saisir_note(eleve=eleve, affectation=affectation, trimestre=trimestre, note_classe=classe_,
                        note_composition=compo, enseignant=prof)
        for eleve, jour, justifiee, motif in [
            (cls.e1, D(2026, 10, 5), False, ""), (cls.e1, D(2026, 10, 6), True, "Paludisme"),
            (cls.e2, D(2026, 10, 7), False, ""),
        ]:
            Absence.objects.create(eleve=eleve, classe=cls.c6a, date_absence=jour, justifiee=justifiee, motif=motif,
                                   enregistre_par=cls.t1)

    def liste(self, compte, classe=None):
        self.client.force_login(compte)
        return self.client.get(reverse("scolarite:exporter_liste_classe_xlsx", args=[(classe or self.c6a).id]))

    def releve(self, compte, classe=None):
        self.client.force_login(compte)
        return self.client.get(reverse("pedagogie:exporter_releve_notes_xlsx", args=[(classe or self.c6a).id]))

    def absences(self, compte, classe=None, **parametres):
        self.client.force_login(compte)
        return self.client.get(reverse("pedagogie:exporter_absences_xlsx", args=[(classe or self.c6a).id]), parametres)


class ListeDeClasseTests(ExportsBase):
    def test_contenu_pour_la_direction(self):
        reponse = self.liste(self.fondateur)
        self.assertEqual(reponse.status_code, 200)
        self.assertIn("spreadsheetml.sheet", reponse["Content-Type"])
        self.assertIn("liste_eleves_6A_2026-2027.xlsx", reponse["Content-Disposition"])
        self.assertIn("no-store", reponse["Cache-Control"])
        classeur = feuilles(reponse.content)
        self.assertEqual(list(classeur), ["Liste de classe"])
        lignes = classeur["Liste de classe"]
        self.assertIn(["Effectif", "3 (dont 2 filles et 1 garçon)"], lignes)
        self.assertIn(["Établissement", "École statistiques A"], lignes)
        entete = next(ligne for ligne in lignes if ligne and ligne[0] == "N°")
        self.assertEqual(entete, ["N°", "Matricule", "Nom", "Prénom", "Sexe", "Date de naissance", "Statut"])
        self.assertEqual(donnees(lignes, "N°"), [
            [1, self.e1.matricule, "eleve1", "Test", "Féminin", datetime.datetime(2014, 5, 1), "En cours"],
            [2, self.e2.matricule, "eleve2", "Test", "Masculin", datetime.datetime(2013, 3, 10), "En cours"],
            [3, self.e3.matricule, "eleve3", "Test", "Féminin", None, "En cours"],
        ])

    def test_la_date_de_naissance_n_est_pas_exportee_aux_roles_sans_etat_civil(self):
        for compte in (self.t1, self.comptable):
            with self.subTest(role=compte.role):
                reponse = self.liste(compte)
                if reponse.status_code != 200:  # selon la matrice de permissions, le comptable peut ne pas avoir le module
                    self.assertEqual(reponse.status_code, 403)
                    continue
                entete = next(ligne for ligne in feuilles(reponse.content)["Liste de classe"] if ligne and ligne[0] == "N°")
                self.assertNotIn("Date de naissance", entete)

    def test_l_enseignant_affecte_obtient_sa_classe_pas_les_autres(self):
        self.assertEqual(self.liste(self.t1).status_code, 200)
        self.assertEqual(self.liste(self.t_fr).status_code, 200)
        self.assertEqual(self.liste(self.t2).status_code, 403)  # t2 n'enseigne pas en 6A
        self.assertEqual(self.liste(self.t1, self.c6b).status_code, 403)

    def test_parents_et_eleves_refuses(self):
        for compte in (self.parent, self.e1):
            self.assertEqual(self.liste(compte).status_code, 403)

    def test_cloisonnement_ecole_et_cycle(self):
        # une classe d'une autre école est « introuvable », comme une classe qui n'existe pas
        self.assertEqual(self.liste(self.fondateur, self.classe_b).status_code, 404)
        self.client.force_login(self.fondateur)
        self.assertEqual(self.client.get(reverse("scolarite:exporter_liste_classe_xlsx", args=[999999])).status_code, 404)
        self.assertEqual(self.liste(self.fondateur_b, self.c6a).status_code, 404)
        # un directeur de cycle ne sort que les classes de son cycle
        self.assertEqual(self.liste(self.dir_1er).status_code, 200)
        self.assertEqual(self.liste(self.dir_1er, self.c9a).status_code, 404)

    def test_anonyme_redirige(self):
        self.client.logout()
        self.assertEqual(
            self.client.get(reverse("scolarite:exporter_liste_classe_xlsx", args=[self.c6a.id])).status_code, 302,
        )

    def test_un_nom_commencant_par_egal_reste_du_texte(self):
        Utilisateur.objects.filter(pk=self.e3.pk).update(nom="=HYPERLINK(\"http://x\")")
        contenu = self.liste(self.fondateur).content
        self.assertIn('=HYPERLINK("http://x")', texte_du_classeur(contenu))
        with zipfile.ZipFile(io.BytesIO(contenu)) as archive:
            for nom in archive.namelist():
                if nom.startswith("xl/worksheets/"):
                    self.assertNotIn("<f>", archive.read(nom).decode("utf-8"))

    def test_audit_sans_nom_d_eleve(self):
        self.liste(self.fondateur)
        entree = JournalAudit.objects.get(action="export_liste_classe")
        self.assertIn("6A", entree.cible)
        self.assertNotIn("eleve1", f"{entree.cible} {entree.details}")

    def test_les_pages_proposent_l_export(self):
        self.client.force_login(self.fondateur)
        adresse = reverse("scolarite:exporter_liste_classe_xlsx", args=[self.c6a.id])
        self.assertContains(self.client.get(reverse("scolarite:liste_classes")), adresse)
        self.assertContains(self.client.get(reverse("pedagogie:suivi_des_cours")), adresse)
        self.assertContains(self.client.get(reverse("scolarite:liste_eleves"), {"classe": self.c6a.id}), adresse)
        self.assertNotContains(self.client.get(reverse("scolarite:liste_eleves")), adresse)
        self.client.force_login(self.t1)
        self.assertContains(self.client.get(reverse("pedagogie:mes_classes")), adresse)


class ReleveDeNotesTests(ExportsBase):
    def test_vue_de_la_direction(self):
        reponse = self.releve(self.fondateur)
        self.assertEqual(reponse.status_code, 200)
        self.assertIn("releve_notes_6A_2026-2027.xlsx", reponse["Content-Disposition"])
        classeur = feuilles(reponse.content)
        self.assertEqual(list(classeur), ["Récapitulatif", "1er trimestre", "2ème trimestre", "Détail des notes"])

        premier = classeur["1er trimestre"]
        self.assertEqual(
            next(ligne for ligne in premier if ligne and ligne[0] == "N°"),
            ["N°", "Matricule", "Nom", "Prénom", "Français (coef 2)", "Maths (coef 1)", "Moyenne générale", "Rang"],
        )  # les matières sont classées par ordre alphabétique
        lignes = donnees(premier, "N°")
        self.assertEqual(lignes[0][4:], [15, 12, 14, 1])      # e1 : 14.00, rang 1
        self.assertEqual(lignes[1][4:], [8, 10, 8.67, 3])     # e2 : 8.67, rang 3 (e1 et e3 sont ex æquo premiers)
        self.assertEqual(lignes[2][4:], [None, 14, 14, 1])    # e3 : 14.00, ex æquo avec e1 -> rang 1
        self.assertEqual(lignes[3], [None, None, "Moyenne de la classe", None, 11.5, 12, 12.22])  # (14 + 8.67 + 14) / 3

        deuxieme = donnees(classeur["2ème trimestre"], "N°")
        self.assertEqual(deuxieme[0][4:], [None, 16, 16, 1])  # e1 : pas de Français ce trimestre
        self.assertEqual(deuxieme[1][4:], [None, 6, 6, 2])    # e2
        self.assertEqual(deuxieme[2][4:], [])  # e3 : aucune note ce trimestre (cellules vides de fin retirées)

        recap = donnees(classeur["Récapitulatif"], "N°")
        self.assertEqual(recap[0][4:], [14, 16, None, 14.5, 1])   # e1 : T1, T2, T3, annuelle, rang annuel
        self.assertEqual(recap[1][4:], [8.67, 6, None, 8, 3])
        self.assertEqual(recap[2][4:], [14, None, None, 14, 2])

        detail = donnees(classeur["Détail des notes"], "Matricule")
        self.assertEqual(len(detail), 7)
        self.assertIn([self.e1.matricule, "eleve1", "Test", "Maths", 1, "1er trimestre", 12, 24, 12], detail)
        self.assertIn([self.e1.matricule, "eleve1", "Test", "Maths", 1, "2ème trimestre", 16, None, 16], detail)

    def test_meme_resultat_que_le_bulletin(self):
        """Moyenne annuelle et rang du relevé = moyenne générale et rang affichés sur le bulletin de l'élève."""
        classeur = feuilles(self.releve(self.fondateur).content)
        recap = {ligne[1]: ligne for ligne in donnees(classeur["Récapitulatif"], "N°")}
        for eleve in (self.e1, self.e2, self.e3):
            inscription = Inscription.objects.get(eleve=eleve, classe=self.c6a)
            bulletin = _calculer_bulletin(eleve, inscription)
            rang, _ = _calculer_classement(inscription)
            with self.subTest(eleve=eleve.matricule):
                self.assertEqual(recap[eleve.matricule][7], float(bulletin["moyenne_generale"]))
                self.assertEqual(recap[eleve.matricule][8], rang)

    def test_vue_de_l_enseignant_limitee_a_sa_matiere(self):
        reponse = self.releve(self.t1)
        self.assertEqual(reponse.status_code, 200)
        classeur = feuilles(reponse.content)
        self.assertEqual(list(classeur), ["1er trimestre", "2ème trimestre", "Détail des notes"])  # pas de récapitulatif
        texte = texte_du_classeur(reponse.content)
        for interdit in ("Français", "Moyenne générale", "Rang", "Moyenne annuelle"):
            self.assertNotIn(interdit, texte)
        self.assertIn("Maths (coef 1)", texte)
        detail = donnees(classeur["Détail des notes"], "Matricule")
        self.assertEqual(len(detail), 5)  # ses 5 notes de Maths, aucune de Français
        self.assertEqual({ligne[3] for ligne in detail}, {"Maths"})

    def test_acces(self):
        self.assertEqual(self.releve(self.t2).status_code, 403)       # pas affecté à la 6A
        self.assertEqual(self.releve(self.t1, self.c6b).status_code, 403)
        for compte in (self.parent, self.e1):
            self.assertEqual(self.releve(compte).status_code, 403)
        self.assertEqual(self.releve(self.fondateur, self.classe_b).status_code, 404)
        self.assertEqual(self.releve(self.dir_1er, self.c9a).status_code, 404)
        self.assertEqual(self.releve(self.dir_1er).status_code, 200)
        self.client.logout()
        self.assertEqual(
            self.client.get(reverse("pedagogie:exporter_releve_notes_xlsx", args=[self.c6a.id])).status_code, 302,
        )

    def test_classe_sans_note(self):
        reponse = self.releve(self.fondateur, self.c6b)
        self.assertEqual(reponse.status_code, 200)
        self.assertEqual(list(feuilles(reponse.content)), ["Récapitulatif", "Détail des notes"])

    def test_audit(self):
        self.releve(self.t1)
        self.releve(self.fondateur)
        vues = sorted(e.details["vue"] for e in JournalAudit.objects.filter(action="export_releve_notes"))
        self.assertEqual(vues, ["complete", "enseignant"])


class AbsencesTests(ExportsBase):
    def test_synthese_et_detail(self):
        reponse = self.absences(self.fondateur)
        self.assertEqual(reponse.status_code, 200)
        classeur = feuilles(reponse.content)
        self.assertEqual(list(classeur), ["Synthèse", "Détail"])
        synthese = classeur["Synthèse"]
        self.assertIn(["Période", "Toute l'année"], synthese)
        lignes = donnees(synthese, "N°")
        self.assertEqual(lignes[0][4:], [2, 1, 1])  # e1 : 2 absences dont 1 justifiée
        self.assertEqual(lignes[1][4:], [1, 0, 1])
        self.assertEqual(lignes[2][4:], [0, 0, 0])
        self.assertEqual(lignes[3], [None, None, "Total", None, 3, 1, 2])
        detail = donnees(classeur["Détail"], "Date")
        self.assertEqual(detail, [
            [datetime.datetime(2026, 10, 5), self.e1.matricule, "eleve1", "Test", "non"],  # sans motif : cellule vide
            [datetime.datetime(2026, 10, 6), self.e1.matricule, "eleve1", "Test", "oui", "Paludisme"],
            [datetime.datetime(2026, 10, 7), self.e2.matricule, "eleve2", "Test", "non"],
        ])

    def test_periode(self):
        classeur = feuilles(self.absences(self.fondateur, du="2026-10-06", au="2026-10-07").content)
        self.assertEqual(len(donnees(classeur["Détail"], "Date")), 2)
        self.assertTrue(any(ligne[:1] == ["Période"] and "06/10/2026" in ligne[1] for ligne in classeur["Synthèse"]))
        self.assertEqual(donnees(classeur["Synthèse"], "N°")[3][4:], [2, 1, 1])

    def test_date_invalide_introuvable(self):
        for parametre in ({"du": "pas-une-date"}, {"au": "31/12/2026"}):
            self.assertEqual(self.absences(self.fondateur, **parametre).status_code, 404)

    def test_acces(self):
        self.assertEqual(self.absences(self.t1).status_code, 200)
        self.assertEqual(self.absences(self.t2).status_code, 403)
        for compte in (self.parent, self.e1):
            self.assertEqual(self.absences(compte).status_code, 403)
        self.assertEqual(self.absences(self.fondateur, self.classe_b).status_code, 404)
        self.assertEqual(self.absences(self.dir_1er, self.c9a).status_code, 404)
        self.client.logout()
        self.assertEqual(
            self.client.get(reverse("pedagogie:exporter_absences_xlsx", args=[self.c6a.id])).status_code, 302,
        )

    def test_absences_d_une_autre_classe_non_melangees(self):
        Absence.objects.create(eleve=self.e4, classe=self.c6b, date_absence=D(2026, 10, 8), enregistre_par=self.t2)
        texte = texte_du_classeur(self.absences(self.fondateur).content)
        self.assertNotIn(self.e4.matricule, texte)
