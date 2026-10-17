"""
Import en masse des élèves depuis un classeur Excel (.xlsx).

Chaque ligne est validée par le MÊME formulaire que l'inscription d'un élève (`InscrireEleveForm`) puis créée par la
même fonction (`creer_eleve_inscrit`) : mêmes règles, mêmes comptes, mêmes liens avec les parents.

Garanties :
- tout ou rien : s'il reste une seule erreur, rien n'est créé (l'utilisateur corrige son fichier et le recharge) ;
- aucune classe d'une autre école, ni d'un autre cycle pour un directeur de cycle (`classes_visibles_pour`) ;
- un fichier rechargé deux fois ne crée pas de doublons (élève déjà inscrit dans la classe = erreur) ;
- fichier borné : 2 Mo, 500 élèves, 40 colonnes lues.
"""
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from comptes.arabisation import chiffres_ascii, nom_de_classe_arabe
from comptes.transcription import transcrire
from scolarite.forms import InscrireEleveForm
from scolarite.inscription_eleve import creer_eleve_inscrit
from scolarite.models import Inscription, classes_visibles_pour

TAILLE_MAX_OCTETS = 2 * 1024 * 1024
LIGNES_MAX = 500
LIGNES_LUES_MAX = 5000  # lignes parcourues au total (un classeur truqué peut déclarer des millions de lignes vides)
COLONNES_LUES_MAX = 40
NOM_FEUILLE = "Élèves"

# (clé du formulaire, titre de colonne du modèle, obligatoire, aide)
COLONNES = [
    ("prenom", "Prénom", True, "Prénom de l'élève."),
    ("nom", "Nom", True, "Nom de famille de l'élève."),
    ("sexe", "Sexe", True, "M ou F."),
    ("classe", "Classe", True, "Nom exact de la classe, tel qu'il figure dans la feuille « Classes »."),
    ("date_naissance", "Date de naissance", False, "Au format jj/mm/aaaa (exemple : 25/03/2015)."),
    ("lieu_naissance", "Lieu de naissance", False, ""),
    ("numero_acte_naissance", "N° acte de naissance", False, "Numéro de l'extrait d'acte de naissance."),
    ("matricule", "Matricule", False, "Laissez vide : la plateforme attribue un matricule. À remplir seulement pour "
                                      "reprendre un matricule existant."),
    ("parent_email_1", "Email parent 1", True, "Obligatoire : c'est par cet email que le parent suivra son enfant."),
    ("parent_nom_1", "Nom parent 1", False, ""),
    ("parent_prenom_1", "Prénom parent 1", False, ""),
    ("parent_telephone_1", "Téléphone parent 1", False, ""),
    ("parent_email_2", "Email parent 2", False, "Facultatif, différent de l'email du parent 1."),
    ("parent_nom_2", "Nom parent 2", False, ""),
    ("parent_prenom_2", "Prénom parent 2", False, ""),
    ("parent_telephone_2", "Téléphone parent 2", False, ""),
    ("prenom_ar", "Prénom en arabe", False, "Facultatif : l'écriture arabe du prénom (en lettres arabes)."),
    ("nom_ar", "Nom en arabe", False, "Facultatif : l'écriture arabe du nom de famille (en lettres arabes)."),
]
CLES = [cle for cle, *_ in COLONNES]

# Titres du modèle en arabe (le fichier peut être rempli, et rechargé, dans l'une ou l'autre langue) et leur explication.
TITRES_ARABES = {
    "prenom": "الاسم الشخصي", "nom": "اسم العائلة", "sexe": "الجنس", "classe": "القسم", "date_naissance": "تاريخ الميلاد",
    "lieu_naissance": "مكان الميلاد", "numero_acte_naissance": "رقم عقد الميلاد", "matricule": "رقم التسجيل",
    "parent_email_1": "البريد الإلكتروني لولي الأمر 1", "parent_nom_1": "اسم عائلة ولي الأمر 1", "parent_prenom_1": "الاسم الشخصي لولي الأمر 1",
    "parent_telephone_1": "هاتف ولي الأمر 1", "parent_email_2": "البريد الإلكتروني لولي الأمر 2", "parent_nom_2": "اسم عائلة ولي الأمر 2",
    "parent_prenom_2": "الاسم الشخصي لولي الأمر 2", "parent_telephone_2": "هاتف ولي الأمر 2",
    "prenom_ar": "الاسم الشخصي بالعربية", "nom_ar": "اسم العائلة بالعربية",
}
AIDES_ARABES = {
    "prenom": "الاسم الشخصي للتلميذ.", "nom": "اسم عائلة التلميذ.", "sexe": "ذكر أو أنثى.",
    "classe": "اسم القسم كما يظهر في ورقة « الأقسام ».", "date_naissance": "بصيغة يوم/شهر/سنة (مثال : 25/03/2015).", "lieu_naissance": "",
    "numero_acte_naissance": "رقم مستخرج عقد الميلاد.",
    "matricule": "اتركه فارغًا : تمنح المنصة رقم تسجيل. لا تملأه إلا لاستعادة رقم تسجيل موجود.",
    "parent_email_1": "إجباري : بهذا البريد الإلكتروني يتابع ولي الأمر ابنه.", "parent_nom_1": "", "parent_prenom_1": "", "parent_telephone_1": "",
    "parent_email_2": "اختياري، ويختلف عن بريد ولي الأمر الأول.", "parent_nom_2": "", "parent_prenom_2": "", "parent_telephone_2": "",
    "prenom_ar": "اختياري : كتابة الاسم الشخصي بحروف عربية.", "nom_ar": "اختياري : كتابة اسم العائلة بحروف عربية.",
}
LIBELLES = {cle: libelle for cle, libelle, *_ in COLONNES}
OBLIGATOIRES = [cle for cle, _, obligatoire, _ in COLONNES if obligatoire]

SEXES = {
    "m": "M", "masculin": "M", "garcon": "M", "h": "M", "homme": "M",
    "f": "F", "feminin": "F", "fille": "F", "femme": "F",
}
FORMATS_DATE = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y")


class ErreurFichier(Exception):
    """Le fichier lui-même est inutilisable (format, taille, colonnes) : message affiché tel quel."""


class ImportInvalide(Exception):
    """Une ou plusieurs lignes sont refusées au moment de créer les élèves : rien n'a été créé."""

    def __init__(self, lignes):
        super().__init__("import refusé")
        self.lignes = lignes


@dataclass
class LigneValidee:
    numero: int
    donnees: dict
    erreurs: list = field(default_factory=list)
    classe: object = None
    donnees_validees: dict = None  # cleaned_data du formulaire, jamais stocké en session

    @property
    def resume(self):
        return f"{self.donnees.get('prenom', '')} {self.donnees.get('nom', '')}".strip()


@dataclass
class Lecture:
    lignes: list  # [(numéro de ligne Excel, {clé: texte})]
    colonnes_ignorees: list


def _normaliser(texte):
    sans_accents = unicodedata.normalize("NFKD", str(texte)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", sans_accents.lower()).strip("_")


def _normaliser_arabe(texte):
    """Titre ou valeur écrit en arabe : sans voyelles brèves ni allongement, alef et ya unifiés, chiffres occidentaux, espaces simples."""
    texte = chiffres_ascii(str(texte))
    texte = re.sub(r"[\u064b-\u065f\u0640]", "", texte)
    texte = texte.translate(str.maketrans("\u0623\u0625\u0622\u0671\u0649\u0629", "\u0627\u0627\u0627\u0627\u064a\u0647"))
    return " ".join(re.sub(r"[_\-:()/.]+", " ", texte).split()).casefold()


def _cle_arabe(texte):
    normalise = _normaliser_arabe(texte)
    return normalise if re.search(r"[\u0600-\u06ff]", normalise) else ""


# titre de colonne normalisé -> clé (le titre du modèle comme la clé technique sont acceptés)
ENTETES = {_normaliser(libelle): cle for cle, libelle, *_ in COLONNES}
ENTETES_ARABES = {_cle_arabe(titre): cle for cle, titre in TITRES_ARABES.items()}
ENTETES_ARABES.update({  # autres façons courantes de nommer la même colonne
    _cle_arabe("الاسم"): "nom", _cle_arabe("اللقب"): "nom", _cle_arabe("الاسم العائلي"): "nom", _cle_arabe("الاسم الاول"): "prenom",
    _cle_arabe("الصف"): "classe",
})
SEXES_ARABES = {
    "ذكر": "M", "ذ": "M", "ولد": "M", "صبي": "M", "رجل": "M", "م": "M",
    "انثي": "F", "انثى": "F", "ا": "F", "ب": "F", "بنت": "F", "فتاه": "F", "امراه": "F", "ث": "F", "اناث": "F",
}
ENTETES.update({_normaliser(cle): cle for cle in CLES})
ENTETES.update({
    "acte_de_naissance": "numero_acte_naissance", "n_acte_naissance": "numero_acte_naissance",
    "prenoms": "prenom", "email_du_parent_1": "parent_email_1", "email_du_parent_2": "parent_email_2",
})


def _texte(valeur):
    """Contenu d'une cellule en texte propre (nombres entiers sans « .0 », dates en ISO, espaces et contrôles retirés)."""
    if valeur is None:
        return ""
    if isinstance(valeur, datetime):
        return valeur.date().isoformat()
    if isinstance(valeur, date):
        return valeur.isoformat()
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    propre = "".join(c if c.isprintable() else " " for c in str(valeur))
    return " ".join(propre.split())


def lire_classeur(fichier):
    """Lit le premier classeur Excel envoyé ; lève ErreurFichier avec un message clair si on ne peut pas s'en servir."""
    nom = (getattr(fichier, "name", "") or "").lower()
    if not nom.endswith(".xlsx"):
        raise ErreurFichier("Le fichier doit être un classeur Excel au format .xlsx (pas .xls ni .csv).")
    if fichier.size > TAILLE_MAX_OCTETS:
        raise ErreurFichier("Le fichier dépasse 2 Mo : il contient trop de données ou des images.")
    contenu = fichier.read()
    if not zipfile.is_zipfile(io.BytesIO(contenu)):
        raise ErreurFichier("Ce fichier n'est pas un classeur Excel (.xlsx) valide.")
    try:
        classeur = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    except Exception as erreur:  # noqa: BLE001 - openpyxl lève des types variés selon l'altération du fichier
        raise ErreurFichier("Impossible de lire ce classeur Excel : le fichier est abîmé ou protégé.") from erreur
    try:
        feuille = classeur[NOM_FEUILLE] if NOM_FEUILLE in classeur.sheetnames else classeur.worksheets[0]
        rangees = []
        for numero, valeurs in enumerate(feuille.iter_rows(max_col=COLONNES_LUES_MAX, values_only=True), start=1):
            if numero > LIGNES_LUES_MAX:
                raise ErreurFichier("Le fichier contient trop de lignes : divisez-le en plusieurs fichiers.")
            if any(valeur not in (None, "") for valeur in valeurs):
                rangees.append((numero, valeurs))
    finally:
        classeur.close()

    if not rangees:
        raise ErreurFichier("Le fichier est vide.")
    (_, entete), *donnees = rangees
    colonnes, ignorees = {}, []
    for position, titre in enumerate(entete):
        texte = _texte(titre)
        if not texte:
            continue
        cle = ENTETES.get(_normaliser(texte)) or ENTETES_ARABES.get(_cle_arabe(texte))
        if cle is None:
            ignorees.append(texte)
        elif cle not in colonnes.values():
            colonnes[position] = cle
    manquantes = [LIBELLES[cle] for cle in OBLIGATOIRES if cle not in colonnes.values()]
    if manquantes:
        raise ErreurFichier(
            "Colonnes obligatoires absentes : " + ", ".join(manquantes)
            + ". Téléchargez le modèle et recopiez vos données dessous, sans changer les titres."
        )
    if not donnees:
        raise ErreurFichier("Le fichier ne contient aucun élève sous la ligne de titres.")
    if len(donnees) > LIGNES_MAX:
        raise ErreurFichier(f"Le fichier contient {len(donnees)} élèves : le maximum est {LIGNES_MAX} par fichier.")
    lignes = [
        (numero, {cle: _texte(valeurs[position]) if position < len(valeurs) else "" for position, cle in colonnes.items()})
        for numero, valeurs in donnees
    ]
    return Lecture(lignes=lignes, colonnes_ignorees=ignorees)


def _date_iso(texte):
    if not texte:
        return "", None
    for format_date in FORMATS_DATE:
        try:
            return datetime.strptime(texte, format_date).date().isoformat(), None
        except ValueError:
            continue
    return texte, "Date de naissance illisible : écrivez-la jj/mm/aaaa (exemple : 25/03/2015)."


def valider(lignes, utilisateur, annee):
    """
    Valide toutes les lignes avec les règles de l'inscription d'un élève. Ne crée rien.
    `lignes` : [(numéro de ligne Excel, {clé: texte})]. Renvoie une LigneValidee par ligne, dans l'ordre.
    """
    classes = classes_visibles_pour(utilisateur).filter(annee_scolaire=annee)
    classe_par_nom = {classe.nom.strip().lower(): classe for classe in classes}
    for classe in classes:  # le modèle en arabe liste les classes en arabe : on les reconnaît aussi
        ecrit = classe.nom_ar or nom_de_classe_arabe(classe.nom) or transcrire(classe.nom)
        classe_par_nom.setdefault(_normaliser_arabe(ecrit), classe)
    deja_inscrits = {
        (classe_id, (prenom or "").strip().lower(), (nom or "").strip().lower())
        for classe_id, prenom, nom in Inscription.objects.filter(classe__in=classes).values_list(
            "classe_id", "eleve__prenom", "eleve__nom",
        )
    }
    vus, matricules_vus = {}, {}
    resultats = []

    for numero, brut in lignes:
        donnees = {cle: brut.get(cle, "") for cle in CLES}
        resultat = LigneValidee(numero=numero, donnees=donnees)
        resultats.append(resultat)

        classe = None
        if donnees["classe"]:
            classe = classe_par_nom.get(donnees["classe"].strip().lower()) or classe_par_nom.get(_normaliser_arabe(donnees["classe"]))
        resultat.classe = classe
        sexe = (SEXES.get(_normaliser(donnees["sexe"])) or SEXES_ARABES.get(_cle_arabe(donnees["sexe"]), "")) if donnees["sexe"] else ""
        date_naissance, erreur_date = _date_iso(donnees["date_naissance"])

        formulaire_donnees = {
            **donnees, "sexe": sexe, "date_naissance": date_naissance, "classe": classe.pk if classe else "",
        }
        formulaire = InscrireEleveForm(
            formulaire_donnees, classes_disponibles=classes, etablissement=utilisateur.etablissement,
        )
        formulaire.is_valid()
        for nom_champ, erreurs in formulaire.errors.items():
            if nom_champ in {"classe", "sexe", "date_naissance"}:
                continue  # signalés plus bas avec un message adapté à Excel
            titre = LIBELLES.get(nom_champ) or formulaire.fields[nom_champ].label if nom_champ in formulaire.fields else ""
            resultat.erreurs.append(f"{titre} : {' '.join(erreurs)}" if titre else " ".join(erreurs))

        if not donnees["classe"]:
            resultat.erreurs.append("Classe : à renseigner.")
        elif classe is None:
            resultat.erreurs.append(
                f"Classe « {donnees['classe']} » introuvable pour l'année {annee.libelle} : recopiez un nom de la "
                "feuille « Classes » du modèle."
            )
        if not donnees["sexe"]:
            resultat.erreurs.append("Sexe : à renseigner (M ou F).")
        elif not sexe:
            resultat.erreurs.append(f"Sexe « {donnees['sexe']} » non reconnu : écrivez M ou F.")
        if erreur_date:
            resultat.erreurs.append(erreur_date)

        prenom, nom = donnees["prenom"].strip().lower(), donnees["nom"].strip().lower()
        if classe is not None and prenom and nom:
            cle_eleve = (classe.pk, prenom, nom)
            if cle_eleve in deja_inscrits:
                resultat.erreurs.append("Cet élève est déjà inscrit dans cette classe (fichier déjà importé ?).")
            elif cle_eleve in vus:
                resultat.erreurs.append(f"Même élève et même classe que la ligne {vus[cle_eleve]}.")
            else:
                vus[cle_eleve] = numero
        matricule = donnees["matricule"].strip()
        if matricule:
            if matricule in matricules_vus:
                resultat.erreurs.append(f"Matricule déjà utilisé à la ligne {matricules_vus[matricule]} de ce fichier.")
            else:
                matricules_vus[matricule] = numero

        if not resultat.erreurs:
            resultat.donnees_validees = formulaire.cleaned_data
    return resultats


def importer(lignes, utilisateur, annee):
    """
    Crée les élèves, tout ou rien. Revalide d'abord : entre l'aperçu et la confirmation, les données ont pu changer
    (classe supprimée, élève inscrit entre-temps). Renvoie [(élève, classe)] ; lève ImportInvalide sans rien créer.
    """
    resultats = valider(lignes, utilisateur, annee)
    refusees = [resultat for resultat in resultats if resultat.erreurs]
    if refusees:
        raise ImportInvalide(refusees)
    crees = []
    with transaction.atomic():  # une exception ci-dessous annule TOUTES les créations
        for resultat in resultats:
            try:
                eleve = creer_eleve_inscrit(resultat.donnees_validees, utilisateur.etablissement)
            except ValidationError as erreur:
                resultat.erreurs.append(" ".join(erreur.messages))
                raise ImportInvalide([resultat]) from erreur
            crees.append((eleve, resultat.classe))
    return crees


# ---------------------------------------------------------------------------
# Modèle à télécharger
# ---------------------------------------------------------------------------

_FOND_ENTETE = PatternFill("solid", fgColor="DDEBF7")
_FOND_OBLIGATOIRE = PatternFill("solid", fgColor="FCE4D6")
_GRAS = Font(bold=True)


def _texte_seul(feuille, ligne, colonne, valeur):
    """Cellule de texte, jamais une formule (un nom de classe commençant par « = » resterait du texte)."""
    cellule = feuille.cell(row=ligne, column=colonne, value=valeur)
    if isinstance(valeur, str):
        cellule.data_type = "s"
    return cellule


def modele_xlsx(utilisateur, annee, langue="fr") -> Workbook:
    """
    Classeur à remplir : feuille « Élèves » (titres), feuille « Classes » (noms exacts) et feuille « Mode d'emploi ».
    En arabe, tout le classeur est en arabe (titres, classes, aide) et se lit de droite à gauche ; le fichier rempli se recharge dans les deux langues.
    """
    arabe = langue == "ar"
    classeur = Workbook()
    classeur.properties.creator = getattr(settings, "NOM_PLATEFORME", "Nexora Éducation")
    eleves = classeur.active
    nom_eleves = "التلاميذ" if arabe else NOM_FEUILLE
    nom_classes = "الأقسام" if arabe else "Classes"
    eleves.title = nom_eleves
    classes_visibles = list(classes_visibles_pour(utilisateur).filter(annee_scolaire=annee).order_by("cycle", "nom"))
    if arabe:
        noms_classes = [classe.nom_ar or nom_de_classe_arabe(classe.nom) or transcrire(classe.nom) for classe in classes_visibles]
    else:
        noms_classes = [classe.nom for classe in classes_visibles]

    for colonne, (cle, libelle, obligatoire, _) in enumerate(COLONNES, start=1):
        titre = TITRES_ARABES[cle] if arabe else libelle
        cellule = _texte_seul(eleves, 1, colonne, titre)
        cellule.font = _GRAS
        cellule.fill = _FOND_OBLIGATOIRE if obligatoire else _FOND_ENTETE
        cellule.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        eleves.column_dimensions[cellule.column_letter].width = max(len(titre) + 4, 16)
    eleves.freeze_panes = "A2"

    if arabe:
        sexe_valide = DataValidation(type="list", formula1='"ذكر,أنثى"', allow_blank=True, showErrorMessage=True,
                                     errorTitle="الجنس", error="اختر ذكر أو أنثى.")
    else:
        sexe_valide = DataValidation(type="list", formula1='"M,F"', allow_blank=True, showErrorMessage=True,
                                     errorTitle="Sexe", error="Choisissez M ou F.")
    eleves.add_data_validation(sexe_valide)
    sexe_valide.add(f"C2:C{LIGNES_MAX + 1}")

    classes = classeur.create_sheet(nom_classes)
    titre_classes = f"أقسام السنة {annee.libelle}" if arabe else f"Classes de l'année {annee.libelle}"
    _texte_seul(classes, 1, 1, titre_classes).font = _GRAS
    for ligne, nom in enumerate(noms_classes, start=2):
        _texte_seul(classes, ligne, 1, nom)
    classes.column_dimensions["A"].width = 40
    if noms_classes:
        classe_valide = DataValidation(
            type="list", formula1=(f"'{nom_classes}'" if arabe else nom_classes) + f"!$A$2:$A${len(noms_classes) + 1}", allow_blank=True,
            showErrorMessage=True, errorTitle="القسم" if arabe else "Classe",
            error="اختر قسمًا من القائمة (ورقة « الأقسام »)." if arabe else "Choisissez une classe de la liste (feuille « Classes »).",
        )
        eleves.add_data_validation(classe_valide)
        classe_valide.add(f"D2:D{LIGNES_MAX + 1}")

    aide = classeur.create_sheet("دليل الاستعمال" if arabe else "Mode d'emploi")
    if arabe:
        lignes_aide = [
            ("كيف تملأ هذا الملف", ""),
            ("1.", f"اكتب تلميذًا واحدًا في كل سطر من ورقة « {nom_eleves} »، تحت سطر العناوين. لا تغيّر العناوين."),
            ("2.", "الأعمدة ذات الخلفية البرتقالية إجبارية."),
            ("3.", f"الحد الأقصى {LIGNES_MAX} تلميذ في الملف الواحد. احفظ الملف بصيغة إكسل ثم أرسله في صفحة الاستيراد."),
            ("4.", "تُظهر لك المعاينة الأخطاء سطرًا بسطر قبل أي إنشاء : لا يُسجَّل شيء ما لم يكن كل شيء صحيحًا، ثم تؤكد."),
            ("", ""),
            ("العمود", "الشرح"),
        ]
        lignes_aide += [
            (TITRES_ARABES[cle] + (" (إجباري)" if obligatoire else ""), AIDES_ARABES.get(cle, "")) for cle, _, obligatoire, _ in COLONNES
        ]
    else:
        lignes_aide = [
            ("Comment remplir ce fichier", ""),
            ("1.", f"Écrivez un élève par ligne dans la feuille « {NOM_FEUILLE} », sous la ligne de titres. Ne changez pas les titres."),
            ("2.", "Les colonnes sur fond orangé sont obligatoires."),
            ("3.", f"Maximum {LIGNES_MAX} élèves par fichier. Enregistrez au format .xlsx, puis envoyez-le sur la page d'import."),
            ("4.", "Un aperçu vous montre les erreurs ligne par ligne avant la moindre création : rien n'est inscrit tant que "
                   "tout n'est pas correct, et vous confirmez ensuite."),
            ("", ""),
            ("Colonne", "Explication"),
        ]
        lignes_aide += [(libelle + (" (obligatoire)" if obligatoire else ""), texte) for _, libelle, obligatoire, texte in COLONNES]
    for numero, (gauche, droite) in enumerate(lignes_aide, start=1):
        _texte_seul(aide, numero, 1, gauche).font = _GRAS if numero in (1, 7) else Font()
        _texte_seul(aide, numero, 2, droite)
    aide.column_dimensions["A"].width = 30
    aide.column_dimensions["B"].width = 100
    if arabe:
        for feuille in (eleves, classes, aide):
            feuille.sheet_view.rightToLeft = True
    return classeur
