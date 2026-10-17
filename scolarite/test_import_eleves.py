"""
Import Excel des élèves : modèle à télécharger, aperçu, confirmation, erreurs, doublons, cloisonnement entre écoles
et entre cycles, limites du fichier.

Les élèves de EducationBase s'appellent « Test eleveN » (6A : eleve1, eleve2, eleve3) ; les élèves importés ci-dessous
ont d'autres noms pour ne jamais se confondre avec eux.
"""
import datetime
import io
import json
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from comptes.models import JournalAudit, Utilisateur
from comptes.roles import Role
from scolarite import import_eleves
from scolarite.import_eleves import COLONNES, LIGNES_MAX, ErreurFichier, lire_classeur
from scolarite.models import Classe, Cycle, Inscription, ParentEnAttente
from statistiques.test_education import EducationBase, creer_compte

TITRES = [libelle for _, libelle, *_ in COLONNES]
CLE_SESSION = "import_eleves_en_attente"


def ligne(prenom="Awa", nom="Diarra", sexe="F", classe="6A", naissance="25/03/2015", lieu="Kati", acte="",
          matricule="", p1_email="parent-awa@parents.test", p1_nom="Diarra", p1_prenom="Oumar", p1_tel="70111222",
          p2_email="", p2_nom="", p2_prenom="", p2_tel="", prenom_ar="", nom_ar=""):
    return [prenom, nom, sexe, classe, naissance, lieu, acte, matricule, p1_email, p1_nom, p1_prenom, p1_tel,
            p2_email, p2_nom, p2_prenom, p2_tel, prenom_ar, nom_ar]


def fichier_excel(lignes, titres=None, nom="eleves.xlsx", feuille="Élèves"):
    classeur = Workbook()
    classeur.active.title = feuille
    classeur.active.append(TITRES if titres is None else titres)
    for valeurs in lignes:
        classeur.active.append(valeurs)
    for rangee in classeur.active.iter_rows():
        for cellule in rangee:
            if isinstance(cellule.value, str) and cellule.value.startswith("="):
                cellule.data_type = "s"  # du texte tapé par l'utilisateur, pas une formule
    tampon = io.BytesIO()
    classeur.save(tampon)
    return SimpleUploadedFile(nom, tampon.getvalue(), content_type="application/octet-stream")


class ImportBase(EducationBase):
    URL = "scolarite:importer_eleves"

    def eleves_ecole(self):
        return Utilisateur.objects.filter(role=Role.ELEVE, etablissement=self.ecole).count()

    def apercu(self, compte, fichier, annee=None):
        self.client.force_login(compte)
        return self.client.post(reverse(self.URL), {
            "action": "apercu", "annee": (annee or self.annee).id, "fichier": fichier,
        })

    def confirmer(self, jeton, annee=None):
        return self.client.post(reverse(self.URL), {
            "action": "confirmer", "annee": (annee or self.annee).id, "jeton": jeton,
        })

    def importer(self, compte, lignes, **kwargs):
        """Aperçu puis confirmation ; renvoie la réponse de la confirmation."""
        reponse = self.apercu(compte, fichier_excel(lignes, **kwargs))
        self.assertEqual(reponse.status_code, 200)
        self.assertIn("jeton", reponse.context, [r.erreurs for r in reponse.context.get("refusees", [])])
        return self.confirmer(reponse.context["jeton"])


class ModeleEtAccesTests(ImportBase):
    def test_modele_excel(self):
        self.client.force_login(self.fondateur)
        reponse = self.client.get(reverse("scolarite:modele_import_eleves"), {"annee": self.annee.id})
        self.assertEqual(reponse.status_code, 200)
        self.assertIn("spreadsheetml.sheet", reponse["Content-Type"])
        self.assertIn("modele_import_eleves.xlsx", reponse["Content-Disposition"])
        self.assertIn("no-store", reponse["Cache-Control"])
        classeur = load_workbook(io.BytesIO(reponse.content))
        self.assertEqual(classeur.sheetnames, ["Élèves", "Classes", "Mode d'emploi"])
        self.assertEqual([c.value for c in classeur["Élèves"][1]], TITRES)
        classes = [c.value for c in classeur["Classes"]["A"]][1:]
        self.assertEqual(sorted(classes), sorted(["6A", "6B", "9A", "TSE", "CAP-E"]))  # pas la 5A de 2025-2026
        # les listes déroulantes sont posées sur les colonnes Sexe et Classe
        validations = {str(v.sqref): v.formula1 for v in classeur["Élèves"].data_validations.dataValidation}
        self.assertEqual(validations[f"C2:C{LIGNES_MAX + 1}"], '"M,F"')
        self.assertIn("Classes!$A$2:$A$6", validations[f"D2:D{LIGNES_MAX + 1}"])

    def test_le_modele_du_directeur_de_cycle_ne_liste_que_son_cycle(self):
        self.client.force_login(self.dir_1er)
        classeur = load_workbook(io.BytesIO(
            self.client.get(reverse("scolarite:modele_import_eleves"), {"annee": self.annee.id}).content,
        ))
        classes = [c.value for c in classeur["Classes"]["A"]][1:]  # sans le titre
        self.assertEqual(sorted(classes), ["6A", "6B"])

    def test_pages_du_personnel_seulement(self):
        for compte in (self.secretaire, self.fondateur):
            self.client.force_login(compte)
            self.assertEqual(self.client.get(reverse(self.URL)).status_code, 200)
        for compte in (self.parent, self.t1, self.e1):
            self.client.force_login(compte)
            for nom in (self.URL, "scolarite:modele_import_eleves"):
                with self.subTest(role=compte.role, page=nom):
                    self.assertEqual(self.client.get(reverse(nom)).status_code, 403)
            self.assertEqual(self.client.post(reverse(self.URL), {"action": "apercu"}).status_code, 403)
        self.client.logout()
        for nom in (self.URL, "scolarite:modele_import_eleves"):
            self.assertEqual(self.client.get(reverse(nom)).status_code, 302)

    def test_annee_d_une_autre_ecole_ou_invalide_introuvable(self):
        self.client.force_login(self.fondateur)
        for annee in (self.annee_b.id, "abc"):
            for nom in (self.URL, "scolarite:modele_import_eleves"):
                with self.subTest(annee=annee, page=nom):
                    self.assertEqual(self.client.get(reverse(nom), {"annee": annee}).status_code, 404)

    def test_pas_de_cache_sur_la_page(self):
        self.client.force_login(self.fondateur)
        self.assertIn("no-store", self.client.get(reverse(self.URL))["Cache-Control"])


class ImportReussiTests(ImportBase):
    def test_apercu_puis_confirmation_creent_les_eleves(self):
        avant = self.eleves_ecole()
        reponse = self.apercu(self.secretaire, fichier_excel([
            ligne(),
            ligne(prenom="Moussa", nom="Keita", sexe="M", classe="6B", naissance="2014-07-01", lieu="Bamako",
                  acte="N°123", p1_email=self.parent.email, p2_email="parent2-moussa@parents.test", p2_nom="Keita",
                  p2_prenom="Aminata", p2_tel="70333444"),
            ligne(prenom="Fanta", nom="Sylla", sexe="Féminin", classe="9A", naissance="", lieu="", p1_nom="", p1_prenom="",
                  p1_tel="", p1_email="sylla@parents.test"),
        ]))
        self.assertEqual(reponse.status_code, 200)
        self.assertContains(reponse, "Rien n'est encore créé")
        self.assertEqual(self.eleves_ecole(), avant, "l'aperçu ne doit rien créer")
        self.assertContains(reponse, "Awa Diarra")

        confirmation = self.confirmer(reponse.context["jeton"])
        self.assertEqual(confirmation.status_code, 200)
        self.assertEqual(self.eleves_ecole(), avant + 3)
        self.assertContains(confirmation, "3 élèves inscrits")

        awa = Utilisateur.objects.get(prenom="Awa", nom="Diarra")
        self.assertEqual((awa.role, awa.etablissement, awa.sexe), (Role.ELEVE, self.ecole, "F"))
        self.assertEqual(awa.date_naissance, datetime.date(2015, 3, 25))
        self.assertEqual(awa.lieu_naissance, "Kati")
        self.assertTrue(awa.matricule)
        self.assertEqual(awa.email, f"{awa.matricule}@eleves.local")
        self.assertContains(confirmation, awa.matricule)
        self.assertTrue(Inscription.objects.filter(eleve=awa, classe=self.c6a).exists())
        # parent inconnu : son email est mémorisé, le lien se fera à son inscription
        self.assertTrue(ParentEnAttente.objects.filter(eleve=awa, email="parent-awa@parents.test", telephone="70111222").exists())

        moussa = Utilisateur.objects.get(prenom="Moussa", nom="Keita")
        self.assertEqual((moussa.sexe, moussa.date_naissance, moussa.numero_acte_naissance),
                         ("M", datetime.date(2014, 7, 1), "N°123"))
        self.assertIn(self.parent, moussa.parents_lies.all())  # parent déjà inscrit : lié tout de suite
        self.assertTrue(ParentEnAttente.objects.filter(eleve=moussa, email="parent2-moussa@parents.test").exists())
        self.assertTrue(Inscription.objects.filter(eleve=moussa, classe=self.c6b).exists())

        fanta = Utilisateur.objects.get(prenom="Fanta", nom="Sylla")
        self.assertEqual((fanta.sexe, fanta.date_naissance), ("F", None))
        self.assertTrue(Inscription.objects.filter(eleve=fanta, classe=self.c9a).exists())
        self.assertEqual(len({awa.matricule, moussa.matricule, fanta.matricule}), 3)

    def test_matricule_repris_et_cellules_de_types_excel(self):
        """Une vraie date Excel, un matricule saisi comme nombre, un téléphone numérique : tous lus correctement."""
        valeurs = ligne(prenom="Issa", nom="Camara", sexe="m", matricule=20269001, p1_tel=70123456,
                        naissance=datetime.datetime(2015, 1, 31))
        self.importer(self.fondateur, [valeurs])
        issa = Utilisateur.objects.get(prenom="Issa", nom="Camara")
        self.assertEqual((issa.matricule, issa.sexe, issa.date_naissance), ("20269001", "M", datetime.date(2015, 1, 31)))
        self.assertTrue(ParentEnAttente.objects.filter(eleve=issa, telephone="70123456").exists())

    def test_un_nom_commencant_par_egal_reste_du_texte(self):
        self.importer(self.fondateur, [ligne(prenom="=CMD()", nom="@x")])
        eleve = Utilisateur.objects.get(prenom="=CMD()")
        self.assertEqual(eleve.nom, "@x")

    def test_journal_d_audit_sans_nom_ni_email(self):
        self.importer(self.secretaire, [ligne(), ligne(prenom="Moussa", nom="Keita", sexe="M", p1_email="k@parents.test")])
        entree = JournalAudit.objects.get(action="import_eleves_excel")
        self.assertEqual(entree.cible, "2026-2027")
        self.assertEqual(entree.details["nombre"], 2)
        texte = json.dumps([entree.cible, entree.details]).lower()
        for interdit in ("awa", "diarra", "moussa", "keita", "parents.test", "70111222"):
            self.assertNotIn(interdit, texte)

    def test_aucune_donnee_personnelle_ne_reste_en_session_apres_confirmation(self):
        reponse = self.apercu(self.fondateur, fichier_excel([ligne()]))
        self.assertIn(CLE_SESSION, self.client.session)  # le temps de l'aperçu seulement
        self.confirmer(reponse.context["jeton"])
        self.assertNotIn(CLE_SESSION, self.client.session)

    def test_colonnes_inconnues_signalees_et_ignorees(self):
        titres = TITRES + ["Commentaire"]
        reponse = self.apercu(self.fondateur, fichier_excel([ligne() + ["à ne pas lire"]], titres=titres))
        self.assertContains(reponse, "Colonnes ignorées")
        self.assertContains(reponse, "Commentaire")
        self.assertIn("jeton", reponse.context)

    def test_titres_techniques_et_ordre_libre_acceptes(self):
        titres = ["classe", "sexe", "nom", "prenom", "parent_email_1"]
        reponse = self.apercu(self.fondateur, fichier_excel([["6A", "F", "Traore", "Kadi", "kadi@parents.test"]], titres=titres))
        self.assertIn("jeton", reponse.context)
        self.confirmer(reponse.context["jeton"])
        self.assertTrue(Utilisateur.objects.filter(prenom="Kadi", nom="Traore", role=Role.ELEVE).exists())

    def test_feuille_nommee_autrement_et_lignes_vides_ignorees(self):
        fichier = fichier_excel([ligne(), [None] * 16, ligne(prenom="Moussa", nom="Keita", sexe="M")], feuille="Feuil1")
        reponse = self.apercu(self.fondateur, fichier)
        self.assertEqual(reponse.context["nombre_lignes"], 2)


class ImportRefuseTests(ImportBase):
    def refuse(self, compte, lignes, attendu, **kwargs):
        avant = self.eleves_ecole()
        reponse = self.apercu(compte, fichier_excel(lignes, **kwargs))
        self.assertEqual(reponse.status_code, 200)
        self.assertNotIn("jeton", reponse.context)
        self.assertGreaterEqual(reponse.context["nombre_refusees"], 1)
        for texte in ([attendu] if isinstance(attendu, str) else attendu):
            self.assertContains(reponse, texte)
        self.assertEqual(self.eleves_ecole(), avant, "une erreur ne doit rien créer")
        self.assertNotIn(CLE_SESSION, self.client.session)
        return reponse

    def test_toutes_les_erreurs_de_ligne_sont_signalees_et_rien_n_est_cree(self):
        reponse = self.refuse(self.fondateur, [
            ligne(classe="Z9"),                                     # ligne 2 : classe inconnue
            ligne(prenom="Bakary", sexe="X"),                       # ligne 3 : sexe
            ligne(prenom="Cheick", naissance="31/02/2015"),         # ligne 4 : date impossible
            ligne(prenom="Djeneba", p1_email=""),                   # ligne 5 : email du parent manquant
            ligne(prenom="Etienne", p1_email="pas-un-email"),       # ligne 6 : email invalide
            ligne(prenom="", nom="Fofana"),                         # ligne 7 : prénom manquant
            ligne(prenom="Gaoussou", classe=""),                    # ligne 8 : classe manquante
        ], ["Classe « Z9 » introuvable", "Sexe « X » non reconnu", "Date de naissance illisible", "Email parent 1",
            "Prénom", "Classe : à renseigner"])
        self.assertEqual(reponse.context["nombre_refusees"], 7)
        self.assertEqual([r.numero for r in reponse.context["refusees"]], [2, 3, 4, 5, 6, 7, 8])

    def test_une_seule_erreur_suffit_a_tout_refuser(self):
        self.refuse(self.fondateur, [ligne(), ligne(prenom="Moussa", nom="Keita", sexe="M"), ligne(prenom="Zoe", classe="inconnue")],
                    "introuvable")
        self.assertFalse(Utilisateur.objects.filter(prenom__in=["Awa", "Moussa"]).exists())

    def test_doublons_dans_le_fichier(self):
        self.refuse(self.fondateur, [ligne(), ligne()], "Même élève et même classe que la ligne 2")
        self.refuse(self.fondateur, [ligne(matricule="MAT-1"), ligne(prenom="Moussa", nom="Keita", matricule="MAT-1")],
                    "Matricule déjà utilisé à la ligne 2")

    def test_matricule_deja_attribue(self):
        self.refuse(self.fondateur, [ligne(matricule=self.e1.matricule)], "déjà utilisé")

    def test_fichier_deja_importe_ne_cree_pas_de_doublons(self):
        self.importer(self.fondateur, [ligne()])
        avant = self.eleves_ecole()
        self.refuse(self.fondateur, [ligne()], "déjà inscrit dans cette classe")
        self.assertEqual(self.eleves_ecole(), avant)
        # un élève existant de l'école (même prénom et nom, même classe) est reconnu aussi
        self.refuse(self.fondateur, [ligne(prenom="Test", nom="eleve1")], "déjà inscrit dans cette classe")

    def test_un_homonyme_dans_une_autre_classe_est_accepte(self):
        reponse = self.apercu(self.fondateur, fichier_excel([ligne(prenom="Test", nom="eleve1", classe="9A")]))
        self.assertIn("jeton", reponse.context)

    def test_parent_dont_le_compte_n_a_pas_le_role_parent(self):
        self.refuse(self.fondateur, [ligne(p1_email=self.t1.email)], "pas le rôle")

    def test_le_deuxieme_parent_doit_differer_du_premier(self):
        self.refuse(self.fondateur, [ligne(p2_email="parent-awa@parents.test")], "différents")


class ImportCloisonnementTests(ImportBase):
    def test_une_classe_d_une_autre_ecole_est_introuvable(self):
        Classe.objects.create(nom="CLASSE-B-SEULE", cycle=Cycle.PREMIER_CYCLE, niveau="annee_6", annee_scolaire=self.annee_b)
        ImportRefuseTests.refuse(self, self.fondateur, [ligne(classe="CLASSE-B-SEULE")], "introuvable")

    def test_une_classe_d_une_autre_annee_est_introuvable(self):
        ImportRefuseTests.refuse(self, self.fondateur, [ligne(classe="5A")], "introuvable")  # la 5A est de 2025-2026

    def test_le_directeur_de_cycle_n_importe_que_dans_son_cycle(self):
        reponse = self.apercu(self.dir_1er, fichier_excel([ligne(classe="6A"), ligne(prenom="Moussa", nom="Keita", sexe="M", classe="6B")]))
        self.assertIn("jeton", reponse.context)
        ImportRefuseTests.refuse(self, self.dir_1er, [ligne(classe="9A")], "introuvable")  # 2ème cycle : pas le sien
        ImportRefuseTests.refuse(self, self.dir_1er, [ligne(classe="TSE")], "introuvable")

    def test_l_import_confirme_ne_cree_que_dans_l_ecole_de_l_utilisateur(self):
        self.importer(self.fondateur, [ligne()])
        awa = Utilisateur.objects.get(prenom="Awa", nom="Diarra")
        self.assertEqual(awa.etablissement, self.ecole)
        self.assertFalse(Utilisateur.objects.filter(etablissement=self.ecole_b, role=Role.ELEVE).exists())

    def test_l_import_de_l_ecole_b_ne_voit_pas_les_classes_de_l_ecole_a(self):
        secretaire_b = creer_compte(Role.SECRETAIRE, "secretaire-b@stat.test", self.ecole_b)
        reponse = self.apercu(secretaire_b, fichier_excel([ligne(classe="9A")]), annee=self.annee_b)
        self.assertNotIn("jeton", reponse.context)
        self.assertContains(reponse, "introuvable")
        # et l'année de l'école A lui est interdite
        self.assertEqual(self.client.get(reverse(self.URL), {"annee": self.annee.id}).status_code, 404)


class ConfirmationTests(ImportBase):
    def test_la_confirmation_ne_sert_qu_une_fois(self):
        reponse = self.apercu(self.fondateur, fichier_excel([ligne()]))
        jeton = reponse.context["jeton"]
        self.confirmer(jeton)
        apres_premier = self.eleves_ecole()
        second = self.confirmer(jeton)  # rechargement de la page / double clic
        self.assertEqual(second.status_code, 302)
        self.assertEqual(self.eleves_ecole(), apres_premier)
        self.assertEqual(Utilisateur.objects.filter(prenom="Awa", nom="Diarra").count(), 1)

    def test_mauvais_jeton_refuse(self):
        avant = self.eleves_ecole()
        for jeton in ("", "faux", "é" * 22):  # un jeton non ASCII ne doit pas provoquer d'erreur serveur
            with self.subTest(jeton=jeton):
                self.apercu(self.fondateur, fichier_excel([ligne()]))  # un refus consomme l'import : aperçu neuf
                suite = self.confirmer(jeton)
                self.assertEqual(suite.status_code, 302)
                self.assertEqual(self.eleves_ecole(), avant)

    def test_confirmation_sans_apercu_refusee(self):
        self.client.force_login(self.fondateur)
        avant = self.eleves_ecole()
        self.assertEqual(self.confirmer("nimporte").status_code, 302)
        self.assertEqual(self.eleves_ecole(), avant)

    def test_le_jeton_d_un_autre_utilisateur_ne_marche_pas(self):
        reponse = self.apercu(self.fondateur, fichier_excel([ligne()]))
        jeton = reponse.context["jeton"]
        self.client.force_login(self.secretaire)  # autre session : aucun import en attente
        avant = self.eleves_ecole()
        self.assertEqual(self.confirmer(jeton).status_code, 302)
        self.assertEqual(self.eleves_ecole(), avant)

    def test_annuler_efface_l_import_en_attente(self):
        reponse = self.apercu(self.fondateur, fichier_excel([ligne()]))
        jeton = reponse.context["jeton"]
        self.client.post(reverse(self.URL), {"action": "annuler", "annee": self.annee.id})
        self.assertNotIn(CLE_SESSION, self.client.session)
        avant = self.eleves_ecole()
        self.confirmer(jeton)
        self.assertEqual(self.eleves_ecole(), avant)

    def test_un_nouvel_apercu_remplace_le_precedent(self):
        premier = self.apercu(self.fondateur, fichier_excel([ligne()]))
        self.apercu(self.fondateur, fichier_excel([ligne(prenom="Moussa", nom="Keita", sexe="M")]))
        avant = self.eleves_ecole()
        self.confirmer(premier.context["jeton"])  # le jeton du premier aperçu n'est plus valable
        self.assertEqual(self.eleves_ecole(), avant)

    def test_les_donnees_changees_entre_l_apercu_et_la_confirmation_font_tout_refuser(self):
        reponse = self.apercu(self.fondateur, fichier_excel([ligne(), ligne(prenom="Moussa", nom="Keita", sexe="M")]))
        # entre-temps, quelqu'un inscrit Awa Diarra dans la 6A
        intrus = creer_compte(Role.ELEVE, "awa-intrus@stat.test", self.ecole)
        Utilisateur.objects.filter(pk=intrus.pk).update(prenom="Awa", nom="Diarra")
        Inscription.objects.create(eleve=intrus, classe=self.c6a)
        avant = self.eleves_ecole()
        suite = self.confirmer(reponse.context["jeton"])
        self.assertEqual(suite.status_code, 200)
        self.assertContains(suite, "déjà inscrit dans cette classe")
        self.assertContains(suite, "a été refusé")
        self.assertEqual(self.eleves_ecole(), avant)
        self.assertFalse(Utilisateur.objects.filter(prenom="Moussa", nom="Keita").exists())

    def test_tout_ou_rien_meme_si_la_creation_echoue_en_cours_de_route(self):
        reponse = self.apercu(self.fondateur, fichier_excel([
            ligne(), ligne(prenom="Moussa", nom="Keita", sexe="M"), ligne(prenom="Fanta", nom="Sylla"),
        ]))
        vraie_creation = import_eleves.creer_eleve_inscrit
        appels = []

        def creation_qui_echoue_a_la_deuxieme(donnees, etablissement):
            appels.append(donnees["prenom"])
            if len(appels) == 2:
                raise ValidationError("Valeur interdite.")
            return vraie_creation(donnees, etablissement)

        avant = self.eleves_ecole()
        with patch("scolarite.import_eleves.creer_eleve_inscrit", side_effect=creation_qui_echoue_a_la_deuxieme):
            suite = self.confirmer(reponse.context["jeton"])
        self.assertEqual(len(appels), 2)
        self.assertContains(suite, "a été refusé")
        self.assertContains(suite, "Valeur interdite.")
        self.assertEqual(self.eleves_ecole(), avant, "le premier élève créé doit être annulé aussi")
        self.assertFalse(Utilisateur.objects.filter(prenom="Awa", nom="Diarra").exists())
        self.assertFalse(JournalAudit.objects.filter(action="import_eleves_excel").exists())


class FichierInutilisableTests(ImportBase):
    def refuse(self, fichier, attendu):
        avant = self.eleves_ecole()
        reponse = self.apercu(self.fondateur, fichier)
        self.assertEqual(reponse.status_code, 200)
        self.assertNotIn("jeton", reponse.context)
        self.assertContains(reponse, attendu)
        self.assertEqual(self.eleves_ecole(), avant)

    def test_pas_de_fichier(self):
        self.client.force_login(self.fondateur)
        reponse = self.client.post(reverse(self.URL), {"action": "apercu", "annee": self.annee.id})
        self.assertContains(reponse, "le fichier Excel à importer")

    def test_extension_refusee(self):
        self.refuse(SimpleUploadedFile("eleves.csv", b"prenom;nom\nAwa;Diarra\n"), "format .xlsx")
        self.refuse(SimpleUploadedFile("eleves.xls", b"\xd0\xcf\x11\xe0"), "format .xlsx")

    def test_faux_classeur(self):
        self.refuse(SimpleUploadedFile("eleves.xlsx", b"ceci n'est pas un classeur"), "classeur Excel (.xlsx) valide")

    def test_classeur_abime(self):
        faux_zip = b"PK\x03\x04" + b"\x00" * 40
        self.refuse(SimpleUploadedFile("eleves.xlsx", faux_zip), "classeur Excel")

    def test_fichier_trop_gros(self):
        self.refuse(SimpleUploadedFile("eleves.xlsx", b"PK" + b"x" * (2 * 1024 * 1024)), "dépasse 2 Mo")

    def test_classeur_vide_ou_sans_eleve(self):
        classeur = Workbook()
        tampon = io.BytesIO()
        classeur.save(tampon)
        self.refuse(SimpleUploadedFile("vide.xlsx", tampon.getvalue()), "Le fichier est vide")
        self.refuse(fichier_excel([]), "aucun élève")

    def test_colonnes_obligatoires_absentes(self):
        titres = [t for t in TITRES if t not in ("Sexe", "Email parent 1")]
        reponse = self.apercu(self.fondateur, fichier_excel([["Awa", "Diarra", "6A"]], titres=titres))
        self.assertContains(reponse, "Colonnes obligatoires absentes : Sexe, Email parent 1")

    def test_trop_d_eleves(self):
        self.refuse(
            fichier_excel([ligne(prenom=f"Eleve{n}", nom="Test") for n in range(LIGNES_MAX + 1)]),
            f"le maximum est {LIGNES_MAX} par fichier",
        )

    def test_lecture_directe_d_un_classeur_valide(self):
        lecture = lire_classeur(fichier_excel([ligne(), ligne(prenom="Moussa", nom="Keita", sexe="M")]))
        self.assertEqual([numero for numero, _ in lecture.lignes], [2, 3])
        self.assertEqual(lecture.lignes[0][1]["prenom"], "Awa")
        self.assertEqual(lecture.lignes[0][1]["parent_email_1"], "parent-awa@parents.test")
        with self.assertRaises(ErreurFichier):
            lire_classeur(SimpleUploadedFile("x.xlsx", b"non"))


class ModeleRoundTripTests(ImportBase):
    def test_le_modele_rempli_s_importe_sans_changer_les_titres(self):
        self.client.force_login(self.fondateur)
        modele = self.client.get(reverse("scolarite:modele_import_eleves"), {"annee": self.annee.id}).content
        classeur = load_workbook(io.BytesIO(modele))
        classeur["Élèves"].append(ligne(prenom="Hawa", nom="Maiga", classe="6B"))
        tampon = io.BytesIO()
        classeur.save(tampon)
        reponse = self.client.post(reverse(self.URL), {
            "action": "apercu", "annee": self.annee.id,
            "fichier": SimpleUploadedFile("rempli.xlsx", tampon.getvalue()),
        })
        self.assertIn("jeton", reponse.context)
        self.assertEqual(reponse.context["nombre_lignes"], 1)
        self.confirmer(reponse.context["jeton"])
        self.assertTrue(Inscription.objects.filter(eleve__prenom="Hawa", classe=self.c6b).exists())


class InscriptionUniqueInchangeeTests(ImportBase):
    """Le formulaire d'inscription d'un seul élève utilise maintenant la même fonction de création que l'import."""

    def test_inscription_d_un_eleve_toujours_possible(self):
        self.client.force_login(self.secretaire)
        reponse = self.client.post(reverse("scolarite:inscrire_eleve"), {
            "prenom": "Solo", "nom": "Unique", "sexe": "M", "classe": self.c6a.id,
            "parent_email_1": "solo@parents.test", "handicap": "visuel",
        })
        self.assertEqual(reponse.status_code, 302)
        solo = Utilisateur.objects.get(prenom="Solo", nom="Unique")
        self.assertEqual((solo.role, solo.etablissement, solo.handicap), (Role.ELEVE, self.ecole, "visuel"))
        self.assertTrue(Inscription.objects.filter(eleve=solo, classe=self.c6a).exists())
        self.assertTrue(ParentEnAttente.objects.filter(eleve=solo, email="solo@parents.test").exists())
        self.assertTrue(JournalAudit.objects.filter(action="inscription_eleve").exists())

    def test_les_pages_proposent_l_import(self):
        self.client.force_login(self.fondateur)
        for nom in ("scolarite:liste_eleves", "scolarite:inscrire_eleve"):
            with self.subTest(page=nom):
                self.assertContains(self.client.get(reverse(nom)), reverse(self.URL))


class ImportEnArabeTests(ImportBase):
    """Une école arabe télécharge un modèle entièrement en arabe, le remplit en arabe et le recharge : mêmes règles, mêmes comptes."""

    TITRES_AR = [import_eleves.TITRES_ARABES[cle] for cle, *_ in COLONNES]

    def modele_en_arabe(self):
        self.client.force_login(self.fondateur)
        self.client.cookies["nexora_langue"] = "ar"
        reponse = self.client.get(reverse("scolarite:modele_import_eleves"), {"annee": self.annee.id})
        self.assertEqual(reponse.status_code, 200)
        return load_workbook(io.BytesIO(reponse.content))

    def test_le_modele_est_entierement_en_arabe(self):
        classeur = self.modele_en_arabe()
        self.assertEqual(classeur.sheetnames, ["التلاميذ", "الأقسام", "دليل الاستعمال"])
        self.assertEqual([c.value for c in classeur["التلاميذ"][1]], self.TITRES_AR)
        classes = [c.value for c in classeur["الأقسام"]["A"]][1:]
        self.assertEqual(len(classes), 5)
        self.assertTrue(all(not any("a" <= lettre.lower() <= "z" for lettre in nom) or nom in {"6A", "6B", "9A", "TSE", "CAP-E"} for nom in classes))
        validations = {str(v.sqref): v.formula1 for v in classeur["التلاميذ"].data_validations.dataValidation}
        self.assertEqual(validations[f"C2:C{LIGNES_MAX + 1}"], '"ذكر,أنثى"')
        self.assertTrue(classeur["التلاميذ"].sheet_view.rightToLeft)

    def test_un_fichier_rempli_en_arabe_est_importe(self):
        Classe.objects.filter(pk=self.c6a.pk).update(nom_ar="السادسة أ")
        lignes = [[
            "سالم", "كيتا", "ذكر", "السادسة أ", "25/03/2015", "كاتي", "", "", "parent-salem@parents.test", "كيتا", "عمر", "٧٠١١١٢٢٢", "", "", "", "",
            "سالم", "كيتا",
        ]]
        reponse = self.apercu(self.fondateur, fichier_excel(lignes, titres=self.TITRES_AR))
        self.assertIn("jeton", reponse.context, [r.erreurs for r in reponse.context.get("refusees", [])])
        self.confirmer(reponse.context["jeton"])
        eleve = Utilisateur.objects.get(role=Role.ELEVE, prenom="سالم", nom="كيتا")
        self.assertEqual(eleve.sexe, "M")
        self.assertEqual((eleve.prenom_ar, eleve.nom_ar), ("سالم", "كيتا"))
        self.assertTrue(Inscription.objects.filter(eleve=eleve, classe=self.c6a).exists())

    def test_les_colonnes_en_arabe_facultatives_sont_enregistrees_depuis_un_fichier_francais(self):
        self.importer(self.fondateur, [ligne(prenom="Awa", nom="Traore", prenom_ar="عوا", nom_ar="تراوري")])
        eleve = Utilisateur.objects.get(role=Role.ELEVE, prenom="Awa", nom="Traore")
        self.assertEqual((eleve.prenom_ar, eleve.nom_ar), ("عوا", "تراوري"))

    def test_une_ecriture_arabe_avec_des_lettres_latines_est_refusee(self):
        reponse = self.apercu(self.fondateur, fichier_excel([ligne(prenom_ar="Awa")]))
        self.assertNotIn("jeton", reponse.context)
